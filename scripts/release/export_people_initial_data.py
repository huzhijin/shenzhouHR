#!/usr/bin/env python3
"""Export one company's approved people baseline into production import workbooks.

The exporter deliberately reads only current organization, employee, and employment
records. It never exports accounts, credentials, sessions, audit history, source
attachments, attendance facts, or report projections.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import subprocess
import tarfile
from dataclasses import dataclass
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Iterable

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font
except ImportError as exc:  # pragma: no cover - exercised by the CLI environment
    raise SystemExit(
        "openpyxl is required. Install it only on the trusted export workstation: "
        "python3 -m pip install openpyxl"
    ) from exc


SAFE_IDENTIFIER = re.compile(r"^[A-Za-z0-9_]+$")
SAFE_COMPANY_CODE = re.compile(r"^[A-Za-z0-9_.-]{1,64}$")
XLSX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)


@dataclass(frozen=True)
class Field:
    key: str
    label: str
    required: bool
    value_type: str
    match_key: bool
    description: str
    allowed_values: tuple[str, ...] = ()


WORKBOOKS: dict[str, tuple[str, tuple[Field, ...]]] = {
    "organizations": (
        "01-组织期初.xlsx",
        (
            Field("organizationCode", "组织编码", True, "TEXT", True, "公司内稳定且唯一"),
            Field("name", "组织名称", True, "TEXT", False, "组织显示名称"),
            Field("parentOrganizationCode", "上级组织编码", False, "TEXT", False, "空值表示根组织"),
            Field(
                "organizationType",
                "组织类型",
                True,
                "ENUM",
                False,
                "仅允许正式模板枚举值",
                ("COMPANY", "DEPARTMENT", "TEAM"),
            ),
            Field("effectiveFrom", "生效日期", True, "DATE", False, "YYYY-MM-DD"),
        ),
    ),
    "employees": (
        "02-员工期初.xlsx",
        (
            Field("employeeNumber", "员工编号", True, "TEXT", True, "公司内稳定且唯一"),
            Field("externalEmployeeId", "外部精确员工ID", False, "PRECISE_ID", True, "可选精确匹配键"),
            Field("displayName", "姓名", True, "TEXT", False, "员工显示姓名，不作为唯一键"),
            Field("effectiveFrom", "生效日期", True, "DATE", False, "YYYY-MM-DD"),
        ),
    ),
    "employments": (
        "03-任职期初.xlsx",
        (
            Field("employeeNumber", "员工编号", True, "TEXT", True, "通过员工编号精确匹配"),
            Field("organizationCode", "组织编码", True, "TEXT", True, "通过组织编码精确匹配"),
            Field("startDate", "任职开始日", True, "DATE", False, "YYYY-MM-DD"),
            Field("terminationDate", "业务离职日", False, "DATE", False, "在职人员留空"),
        ),
    ),
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Export a sanitized, company-scoped people initial-data package."
    )
    parser.add_argument("--mysql-bin", required=True, type=Path)
    parser.add_argument("--defaults-extra-file", required=True, type=Path)
    parser.add_argument("--database", required=True)
    parser.add_argument("--source-company-code", required=True)
    parser.add_argument("--target-company-code", required=True)
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--expected-organizations", required=True, type=int)
    parser.add_argument("--expected-employees", required=True, type=int)
    parser.add_argument("--expected-employments", required=True, type=int)
    return parser.parse_args()


def validate_args(args: argparse.Namespace) -> None:
    if not args.mysql_bin.is_file() or not os.access(args.mysql_bin, os.X_OK):
        raise ValueError(f"MySQL client is not executable: {args.mysql_bin}")
    if not args.defaults_extra_file.is_file():
        raise ValueError(f"MySQL defaults file is missing: {args.defaults_extra_file}")
    if args.defaults_extra_file.stat().st_mode & 0o077:
        raise ValueError("MySQL defaults file must not be group/world accessible")
    if not SAFE_IDENTIFIER.fullmatch(args.database):
        raise ValueError("database must contain only letters, digits, and underscore")
    for name in ("source_company_code", "target_company_code"):
        value = getattr(args, name)
        if not SAFE_COMPANY_CODE.fullmatch(value):
            raise ValueError(f"{name} contains unsafe characters")
    for name in (
        "expected_organizations",
        "expected_employees",
        "expected_employments",
    ):
        value = getattr(args, name)
        if value <= 0 or value > 50_000:
            raise ValueError(f"{name} must be between 1 and 50000")
    if args.output_dir.exists():
        raise ValueError(f"output directory already exists: {args.output_dir}")
    archive = args.output_dir.with_suffix(".tar.gz")
    if archive.exists():
        raise ValueError(f"output archive already exists: {archive}")
    checksum = Path(f"{archive}.sha256")
    if checksum.exists():
        raise ValueError(f"output archive checksum already exists: {checksum}")


def sql_literal(value: str) -> str:
    if not SAFE_COMPANY_CODE.fullmatch(value):
        raise ValueError("unsafe SQL literal")
    return "'" + value + "'"


def mysql_json_rows(args: argparse.Namespace, query: str) -> list[dict[str, Any]]:
    command = [
        str(args.mysql_bin),
        f"--defaults-extra-file={args.defaults_extra_file}",
        "--default-character-set=utf8mb4",
        "--connect-timeout=10",
        "--batch",
        "--raw",
        "--skip-column-names",
        args.database,
        "--execute",
        query,
    ]
    environment = os.environ.copy()
    environment["MYSQL_HISTFILE"] = "/dev/null"
    result = subprocess.run(
        command,
        check=False,
        capture_output=True,
        text=True,
        encoding="utf-8",
        env=environment,
    )
    if result.returncode != 0:
        message = result.stderr.strip().splitlines()
        tail = message[-1] if message else "unknown MySQL error"
        raise RuntimeError(f"MySQL query failed: {tail}")
    rows: list[dict[str, Any]] = []
    for line_number, line in enumerate(result.stdout.splitlines(), start=1):
        if not line.strip():
            continue
        try:
            value = json.loads(line)
        except json.JSONDecodeError as exc:
            raise RuntimeError(
                f"MySQL returned invalid JSON on output line {line_number}"
            ) from exc
        if not isinstance(value, dict):
            raise RuntimeError("MySQL JSON row must be an object")
        rows.append(value)
    return rows


def load_company(args: argparse.Namespace) -> dict[str, Any]:
    company_code = sql_literal(args.source_company_code)
    rows = mysql_json_rows(
        args,
        "SELECT JSON_OBJECT("
        "'companyCode', code, 'companyName', name, 'status', status) "
        "FROM company WHERE BINARY code = BINARY " + company_code,
    )
    if len(rows) != 1:
        raise ValueError(
            f"expected one source company {args.source_company_code}, found {len(rows)}"
        )
    if rows[0].get("status") != "ACTIVE":
        raise ValueError("source company must be ACTIVE")
    return rows[0]


def load_organizations(args: argparse.Namespace) -> list[dict[str, Any]]:
    code = sql_literal(args.source_company_code)
    return mysql_json_rows(
        args,
        "SELECT JSON_OBJECT("
        "'organizationCode', ov.code, "
        "'name', ov.name, "
        "'parentOrganizationCode', COALESCE(pov.code, ''), "
        "'organizationType', ov.org_type, "
        "'effectiveFrom', DATE_FORMAT(ov.effective_from, '%Y-%m-%d')) "
        "FROM company c "
        "JOIN organization_identity oi ON oi.company_id = c.company_id "
        "JOIN organization_current_projection ocp ON ocp.organization_id = oi.organization_id "
        "JOIN organization_version ov ON ov.organization_version_id = ocp.current_version_id "
        "LEFT JOIN organization_identity poi "
        "ON poi.organization_id = ov.parent_organization_id "
        "AND poi.company_id = oi.company_id AND poi.identity_status = 'ACTIVE' "
        "LEFT JOIN organization_current_projection pcp "
        "ON pcp.organization_id = poi.organization_id "
        "LEFT JOIN organization_version pov "
        "ON pov.organization_version_id = pcp.current_version_id "
        "WHERE BINARY c.code = BINARY " + code + " "
        "AND c.status = 'ACTIVE' AND oi.identity_status = 'ACTIVE' "
        "AND ov.status = 'ACTIVE' "
        "AND (ov.parent_organization_id IS NULL OR poi.organization_id IS NOT NULL) "
        "ORDER BY ov.code",
    )


def load_employees(args: argparse.Namespace) -> list[dict[str, Any]]:
    code = sql_literal(args.source_company_code)
    return mysql_json_rows(
        args,
        "SELECT JSON_OBJECT("
        "'employeeNumber', ev.employee_number, "
        "'externalEmployeeId', COALESCE(ev.external_employee_id, ''), "
        "'displayName', ev.display_name, "
        "'effectiveFrom', DATE_FORMAT(ev.effective_from, '%Y-%m-%d')) "
        "FROM company c "
        "JOIN employee e ON e.company_id = c.company_id "
        "JOIN employee_current_projection ecp ON ecp.employee_id = e.employee_id "
        "JOIN employee_version ev ON ev.employee_version_id = ecp.current_version_id "
        "WHERE BINARY c.code = BINARY " + code + " "
        "AND c.status = 'ACTIVE' AND ev.status = 'ACTIVE' "
        "ORDER BY ev.employee_number",
    )


def load_employments(args: argparse.Namespace) -> list[dict[str, Any]]:
    code = sql_literal(args.source_company_code)
    return mysql_json_rows(
        args,
        "SELECT JSON_OBJECT("
        "'employeeNumber', ev.employee_number, "
        "'organizationCode', ov.code, "
        "'startDate', DATE_FORMAT(ea.effective_from, '%Y-%m-%d'), "
        "'terminationDate', COALESCE(DATE_FORMAT(ea.termination_date, '%Y-%m-%d'), '')) "
        "FROM company c "
        "JOIN employee e ON e.company_id = c.company_id "
        "JOIN employee_current_projection ecp ON ecp.employee_id = e.employee_id "
        "JOIN employee_version ev ON ev.employee_version_id = ecp.current_version_id "
        "JOIN employment_assignment ea ON ea.employee_id = e.employee_id "
        "AND ea.current_version_marker = 1 AND ea.record_status = 'ACTIVE' "
        "JOIN organization_identity oi ON oi.organization_id = ea.organization_id "
        "AND oi.company_id = e.company_id AND oi.identity_status = 'ACTIVE' "
        "JOIN organization_current_projection ocp "
        "ON ocp.organization_id = oi.organization_id "
        "JOIN organization_version ov ON ov.organization_version_id = ocp.current_version_id "
        "WHERE BINARY c.code = BINARY " + code + " "
        "AND c.status = 'ACTIVE' AND ev.status = 'ACTIVE' AND ov.status = 'ACTIVE' "
        "ORDER BY ev.employee_number",
    )


def required_text(row: dict[str, Any], key: str) -> str:
    value = row.get(key)
    if value is None:
        return ""
    return str(value).strip()


def require_length(value: str, maximum: int, field_name: str) -> None:
    if not value or len(value) > maximum:
        raise ValueError(f"{field_name} length must be between 1 and {maximum}")


def require_optional_length(value: str, maximum: int, field_name: str) -> None:
    if len(value) > maximum:
        raise ValueError(f"{field_name} length must not exceed {maximum}")


def parse_iso_date(value: str, field_name: str) -> date:
    try:
        return date.fromisoformat(value)
    except ValueError as exc:
        raise ValueError(f"{field_name} must use YYYY-MM-DD") from exc


def validate_dataset(
    organizations: list[dict[str, Any]],
    employees: list[dict[str, Any]],
    employments: list[dict[str, Any]],
    expected: tuple[int, int, int],
) -> None:
    actual = (len(organizations), len(employees), len(employments))
    if actual != expected:
        raise ValueError(f"row counts do not match approval: expected {expected}, got {actual}")

    organization_codes = [required_text(row, "organizationCode") for row in organizations]
    if any(not value for value in organization_codes):
        raise ValueError("organization code cannot be empty")
    if len(organization_codes) != len(set(organization_codes)):
        raise ValueError("organization codes are not unique")
    organization_code_set = set(organization_codes)
    roots = 0
    parents: dict[str, str] = {}
    for row in organizations:
        code = required_text(row, "organizationCode")
        require_length(code, 128, "organizationCode")
        parent = required_text(row, "parentOrganizationCode")
        if parent:
            require_optional_length(parent, 128, "parentOrganizationCode")
            if parent not in organization_code_set:
                raise ValueError(f"organization parent is absent from export: {parent}")
        else:
            roots += 1
        if required_text(row, "organizationType") not in {
            "COMPANY",
            "DEPARTMENT",
            "TEAM",
        }:
            raise ValueError("organization type is outside the production template")
        require_length(required_text(row, "name"), 200, "organization name")
        parse_iso_date(required_text(row, "effectiveFrom"), "organization effectiveFrom")
        parents[code] = parent
    if roots != 1:
        raise ValueError(f"expected one organization root, found {roots}")
    root_rows = [row for row in organizations if not required_text(row, "parentOrganizationCode")]
    if required_text(root_rows[0], "organizationType") != "COMPANY":
        raise ValueError("the organization root must use type COMPANY")
    for start in organization_code_set:
        visited: set[str] = set()
        current = start
        while current:
            if current in visited:
                raise ValueError(f"organization hierarchy contains a cycle at: {current}")
            visited.add(current)
            current = parents[current]

    employee_numbers = [required_text(row, "employeeNumber") for row in employees]
    if any(not value for value in employee_numbers):
        raise ValueError("employee number cannot be empty")
    if len(employee_numbers) != len(set(employee_numbers)):
        raise ValueError("employee numbers are not unique")
    employee_number_set = set(employee_numbers)
    external_ids = [required_text(row, "externalEmployeeId") for row in employees]
    populated_external_ids = [value for value in external_ids if value]
    if len(populated_external_ids) != len(set(populated_external_ids)):
        raise ValueError("populated external employee IDs are not unique")
    for row in employees:
        require_length(required_text(row, "employeeNumber"), 128, "employeeNumber")
        require_optional_length(
            required_text(row, "externalEmployeeId"), 128, "externalEmployeeId"
        )
        require_length(required_text(row, "displayName"), 100, "displayName")
        parse_iso_date(required_text(row, "effectiveFrom"), "employee effectiveFrom")

    assigned_numbers: set[str] = set()
    for row in employments:
        number = required_text(row, "employeeNumber")
        organization_code = required_text(row, "organizationCode")
        if number not in employee_number_set:
            raise ValueError(f"employment references unknown employee number: {number}")
        if organization_code not in organization_code_set:
            raise ValueError(
                f"employment references unknown organization code: {organization_code}"
            )
        if number in assigned_numbers:
            raise ValueError(f"multiple current employments exported for employee: {number}")
        assigned_numbers.add(number)
        start = parse_iso_date(required_text(row, "startDate"), "employment startDate")
        termination_value = required_text(row, "terminationDate")
        if termination_value:
            termination = parse_iso_date(termination_value, "employment terminationDate")
            if termination < start:
                raise ValueError("employment terminationDate cannot precede startDate")
    if assigned_numbers != employee_number_set:
        raise ValueError("every exported employee must have exactly one current employment")


def write_workbook(
    output: Path,
    title: str,
    fields: tuple[Field, ...],
    rows: Iterable[dict[str, Any]],
) -> None:
    workbook = Workbook()
    workbook.properties.creator = "ShenzhouHR initial-data exporter"
    workbook.properties.title = title
    workbook.properties.created = datetime(2026, 8, 6, tzinfo=timezone.utc)
    workbook.properties.modified = datetime(2026, 8, 6, tzinfo=timezone.utc)
    data = workbook.active
    data.title = "Data"
    data.freeze_panes = "A2"
    headers = [field.label for field in fields]
    data.append(headers)
    for cell in data[1]:
        cell.font = Font(bold=True)
    data.auto_filter.ref = f"A1:{column_letter(len(fields))}1"
    for index in range(1, len(fields) + 1):
        data.column_dimensions[column_letter(index)].width = 22
    for row in rows:
        data.append([required_text(row, field.key) for field in fields])
        for cell in data[data.max_row]:
            cell.number_format = "@"

    instructions = workbook.create_sheet("Field Instructions")
    instructions.append(
        [
            "Key",
            "Label",
            "Required",
            "Value Type",
            "Match Key",
            "Description",
            "Allowed Values",
        ]
    )
    for cell in instructions[1]:
        cell.font = Font(bold=True)
    for field in fields:
        instructions.append(
            [
                field.key,
                field.label,
                field.required,
                field.value_type,
                field.match_key,
                field.description,
                ",".join(field.allowed_values),
            ]
        )
    for index, width in enumerate((24, 24, 12, 16, 12, 52, 42), start=1):
        instructions.column_dimensions[column_letter(index)].width = width
    workbook.save(output)


def column_letter(index: int) -> str:
    value = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        value = chr(65 + remainder) + value
    return value


def validate_workbook(path: Path, fields: tuple[Field, ...], expected_rows: int) -> None:
    with path.open("rb") as source:
        prefix = source.read(4)
    if prefix != b"PK\x03\x04":
        raise ValueError(f"workbook is not an OOXML ZIP file: {path.name}")
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if workbook.sheetnames != ["Data", "Field Instructions"]:
            raise ValueError(f"unexpected worksheets in {path.name}")
        sheet = workbook["Data"]
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        if headers != [field.label for field in fields]:
            raise ValueError(f"unexpected headers in {path.name}")
        if sheet.max_row - 1 != expected_rows:
            raise ValueError(f"unexpected row count in {path.name}")
        for row in workbook.worksheets:
            for cells in row.iter_rows():
                if any(cell.data_type == "f" for cell in cells):
                    raise ValueError(f"formula cell is forbidden in {path.name}")
    finally:
        workbook.close()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def write_supporting_files(
    args: argparse.Namespace,
    company: dict[str, Any],
    workbook_paths: list[Path],
    counts: dict[str, int],
) -> None:
    generated_at = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
    manifest = {
        "formatVersion": 1,
        "generatedAtUtc": generated_at,
        "sourceDatabase": args.database,
        "sourceCompanyCode": company["companyCode"],
        "sourceCompanyName": company["companyName"],
        "targetCompanyCode": args.target_company_code,
        "scope": ["ORGANIZATION", "EMPLOYEE", "EMPLOYMENT"],
        "excluded": [
            "accounts",
            "credentials",
            "sessions",
            "auditHistory",
            "sourceAttachments",
            "attendanceConfiguration",
            "punches",
            "oaDocuments",
            "reportProjections",
        ],
        "counts": counts,
        "files": [
            {
                "name": path.name,
                "mediaType": XLSX_MEDIA_TYPE,
                "sha256": sha256(path),
                "sizeBytes": path.stat().st_size,
            }
            for path in workbook_paths
        ],
    }
    (args.output_dir / "MANIFEST.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    readme = f"""神州 HR 客户专用期初人员数据补充包

目标公司代码：{args.target_company_code}
来源公司：{company['companyName']}（{company['companyCode']}）
批准计数：组织 {counts['organizations']}、员工 {counts['employees']}、当前任职 {counts['employments']}

注意：本目录含员工个人信息，只能通过批准的安全通道交给客户，不得提交 Git、发到公共群或长期留在桌面。

导入顺序（必须逐个完成“创建批次 → 上传 → 保存字段映射 → 预检 → 发布”）：
导入前先对客户数据库做完整备份。每一个批次发布前都要截图保存预检结果。

1. 01-组织期初.xlsx，类型选“组织”。预检应为新增 {counts['organizations']}，其余分类和阻断均为 0。
2. 02-员工期初.xlsx，类型选“员工”。预检应为新增 {counts['employees']}，其余分类和阻断均为 0。
3. 03-任职期初.xlsx，类型选“任职”。预检应为新增 {counts['employments']}，其余分类和阻断均为 0。

三个批次都必须选择客户服务器上已经由 szadmin 绑定的公司“{args.target_company_code}”，不能新建 SZSC，也不能运行本机 four-company-finalization 脚本。

必须直接上传本目录中的原文件，不要用 Excel/WPS 打开后另存。相同原文件可由系统按哈希识别重复；任职文件如果被另存成不同二进制后再次发布，会因为任职周期重叠而被预检阻断。

本包不包含账号密码、登录会话、审计记录、考勤地点/班次/日历/考勤组/规则、打卡、OA 单据或报表结果。导入完成仅表示人员主数据恢复，不表示考勤生产上线。
"""
    (args.output_dir / "README.txt").write_text(readme, encoding="utf-8")
    verification_sql = f"""-- Read-only verification after all three people imports are published.
SELECT c.code AS company_code,
       (SELECT COUNT(*)
          FROM organization_identity oi
          JOIN organization_current_projection ocp
            ON ocp.organization_id = oi.organization_id
         WHERE oi.company_id = c.company_id) AS organizations,
       (SELECT COUNT(*)
          FROM employee e
          JOIN employee_current_projection ecp
            ON ecp.employee_id = e.employee_id
         WHERE e.company_id = c.company_id) AS employees,
       (SELECT COUNT(*)
          FROM employment_assignment ea
          JOIN employee e ON e.employee_id = ea.employee_id
         WHERE e.company_id = c.company_id
           AND ea.current_version_marker = 1
           AND ea.record_status = 'ACTIVE') AS current_employments
FROM company c
WHERE UPPER(c.code) = UPPER('{args.target_company_code}')
ORDER BY c.code;
-- Expected: {args.target_company_code} | {counts['organizations']} | {counts['employees']} | {counts['employments']}
"""
    (args.output_dir / "VERIFY-COUNTS.sql").write_text(
        verification_sql, encoding="utf-8"
    )


def create_archive(output_dir: Path) -> Path:
    archive = output_dir.with_suffix(".tar.gz")
    with tarfile.open(archive, "w:gz", format=tarfile.PAX_FORMAT) as bundle:
        bundle.add(output_dir, arcname=output_dir.name, recursive=True)
    Path(f"{archive}.sha256").write_text(
        f"{sha256(archive)}  {archive.name}\n", encoding="utf-8"
    )
    return archive


def main() -> None:
    os.umask(0o077)
    args = parse_args()
    validate_args(args)
    company = load_company(args)
    organizations = load_organizations(args)
    employees = load_employees(args)
    employments = load_employments(args)
    expected = (
        args.expected_organizations,
        args.expected_employees,
        args.expected_employments,
    )
    validate_dataset(organizations, employees, employments, expected)

    args.output_dir.mkdir(parents=True, mode=0o700)
    os.chmod(args.output_dir, 0o700)
    datasets = {
        "organizations": organizations,
        "employees": employees,
        "employments": employments,
    }
    workbook_paths: list[Path] = []
    for key, rows in datasets.items():
        file_name, fields = WORKBOOKS[key]
        path = args.output_dir / file_name
        write_workbook(path, file_name, fields, rows)
        os.chmod(path, 0o600)
        validate_workbook(path, fields, len(rows))
        workbook_paths.append(path)
    counts = {key: len(rows) for key, rows in datasets.items()}
    write_supporting_files(args, company, workbook_paths, counts)
    for path in args.output_dir.iterdir():
        os.chmod(path, 0o600)
    archive = create_archive(args.output_dir)
    os.chmod(archive, 0o600)
    os.chmod(Path(f"{archive}.sha256"), 0o600)
    print(
        json.dumps(
            {
                "outputDirectory": str(args.output_dir),
                "archive": str(archive),
                "archiveSha256": sha256(archive),
                "counts": counts,
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
