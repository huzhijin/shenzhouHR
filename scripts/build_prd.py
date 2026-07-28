#!/usr/bin/env python3
from __future__ import annotations

import os
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs"
QA_DIR = ROOT / ".docx-qa"
OUT_PATH = OUT_DIR / "神州HR考勤与薪资核算系统_PRD_V1.7.docx"
DEV_FLOW_PATH = QA_DIR / "development-flow.png"
DATA_FLOW_PATH = QA_DIR / "system-data-flow.png"
def optional_path(environment_name: str) -> Path | None:
    value = os.environ.get(environment_name)
    return Path(value) if value else None


SOURCE_IMAGE = optional_path("SHENZHOUHR_PRD_SOURCE_IMAGE")
SOURCE_RULE_IMAGE_1 = optional_path("SHENZHOUHR_PRD_RULE_IMAGE_1")
SOURCE_RULE_IMAGE_2 = optional_path("SHENZHOUHR_PRD_RULE_IMAGE_2")

TABLE_HELPER_DIR = Path(
    os.environ.get(
        "SHENZHOUHR_DOCUMENT_TOOLS",
        "/Users/huzhijin/.codex/plugins/cache/openai-primary-runtime/"
        "documents/26.723.12215/skills/documents/scripts",
    )
)
sys.path.insert(0, str(TABLE_HELPER_DIR))
from table_geometry import apply_table_geometry, column_widths_from_weights  # noqa: E402


# Use a single Unicode font in every OOXML font slot. This keeps Chinese text
# stable across Word and the headless LibreOffice renderer used for visual QA.
FONT_LATIN = "Arial Unicode MS"
FONT_CJK = "Arial Unicode MS"
PIL_FONT = "/System/Library/Fonts/PingFang.ttc"
BLUE = "2E74B5"
DARK_BLUE = "17365D"
INK = "1F2937"
MUTED = "667085"
LIGHT_BLUE = "EAF2F8"
LIGHTER_BLUE = "F5F9FC"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "D0D5DD"
GREEN = "277A53"
LIGHT_GREEN = "EAF6EF"
ORANGE = "A15C00"
LIGHT_ORANGE = "FFF4E5"
RED = "9B1C1C"
LIGHT_RED = "FDECEC"
WHITE = "FFFFFF"
CONTENT_WIDTH_DXA = 9360


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **kwargs) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in kwargs:
            continue
        edge_data = kwargs[edge]
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        for key in ("val", "sz", "space", "color"):
            if key in edge_data:
                element.set(qn(f"w:{key}"), str(edge_data[key]))


def set_run_font(run, size: float | None = None, color: str | None = None,
                 bold: bool | None = None, italic: bool | None = None,
                 latin: str = FONT_LATIN, cjk: str = FONT_CJK) -> None:
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), cjk)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_keep(paragraph, keep_with_next: bool = False,
                       keep_together: bool = False) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    if keep_with_next:
        tag = p_pr.find(qn("w:keepNext"))
        if tag is None:
            p_pr.append(OxmlElement("w:keepNext"))
    if keep_together:
        tag = p_pr.find(qn("w:keepLines"))
        if tag is None:
            p_pr.append(OxmlElement("w:keepLines"))


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row) -> None:
    """Keep a logical table row on one page for readable PRD pagination."""
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_cell_text(cell, text: str, *, bold: bool = False, color: str = INK,
                  size: float = 9.3, align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(str(text))
    set_run_font(r, size=size, color=color, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc: Document, headers: list[str], rows: list[list[str]],
              weights: list[float], *, font_size: float = 9.1,
              header_fill: str = LIGHT_GRAY,
              alignments: list | None = None,
              first_col_bold: bool = False) -> object:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = 0
    table.autofit = False
    widths = column_widths_from_weights(weights, CONTENT_WIDTH_DXA)
    header = table.rows[0]
    set_repeat_table_header(header)
    set_row_cant_split(header)
    for i, text in enumerate(headers):
        set_cell_text(header.cells[i], text, bold=True, color=DARK_BLUE, size=9.2,
                      align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(header.cells[i], header_fill)
    for row_data in rows:
        row = table.add_row()
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        set_row_cant_split(row)
        for i, text in enumerate(row_data):
            alignment = alignments[i] if alignments else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(row.cells[i], text, bold=(first_col_bold and i == 0),
                          size=font_size, align=alignment)
    apply_table_geometry(table, widths, indent_dxa=120,
                         cell_margins_dxa={"top": 80, "bottom": 80, "start": 120, "end": 120})
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    return table


def add_body(doc: Document, text: str, *, bold_prefix: str | None = None,
             italic: bool = False, after: float = 6, color: str = INK) -> object:
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(after)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, size=11, color=color, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, size=11, color=color, italic=italic)
    else:
        r = p.add_run(text)
        set_run_font(r, size=11, color=color, italic=italic)
    return p


def add_bullet(doc: Document, text: str, level: int = 0, *, color: str = INK) -> object:
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.167
    r = p.add_run(text)
    set_run_font(r, size=11, color=color)
    return p


def add_number(doc: Document, text: str, level: int = 0) -> object:
    style = "List Number" if level == 0 else "List Number 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.167
    r = p.add_run(text)
    set_run_font(r, size=11, color=INK)
    return p


def add_callout(doc: Document, label: str, text: str, *, fill: str = LIGHT_BLUE,
                label_color: str = BLUE) -> object:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.1
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), label_color)
    borders.append(left)
    p_pr.append(borders)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "160")
    ind.set(qn("w:right"), "160")
    p_pr.append(ind)
    spacing = p_pr.find(qn("w:spacing"))
    if spacing is not None:
        spacing.set(qn("w:before"), "100")
        spacing.set(qn("w:after"), "160")
    r = p.add_run(f"{label}  ")
    set_run_font(r, size=10.5, color=label_color, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=INK)
    set_paragraph_keep(p, keep_together=True)
    return p


def add_heading(doc: Document, text: str, level: int = 1) -> object:
    p = doc.add_paragraph(text, style=f"Heading {level}")
    set_paragraph_keep(p, keep_with_next=True)
    return p


def add_figure(doc: Document, path: Path, caption: str, *, width: float = 6.2,
               alt_text: str | None = None) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    shape = run.add_picture(str(path), width=Inches(width))
    if alt_text:
        doc_pr = shape._inline.docPr
        doc_pr.set("descr", alt_text)
        doc_pr.set("title", caption)
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_before = Pt(0)
    c.paragraph_format.space_after = Pt(8)
    r = c.add_run(caption)
    set_run_font(r, size=9.2, color=MUTED, italic=True)
    set_paragraph_keep(c, keep_together=True)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT_LATIN)
    fonts.set(qn("w:hAnsi"), FONT_LATIN)
    fonts.set(qn("w:eastAsia"), FONT_CJK)
    r_pr.extend([fonts, color, underline])
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def rgb(hex_color: str) -> tuple[int, int, int]:
    return tuple(int(hex_color[i:i + 2], 16) for i in (0, 2, 4))


def pil_font(size: int, index: int = 0) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(PIL_FONT, size=size, index=index)


def rounded_box(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int],
                text: str, fill: str, outline: str, *, font_size: int = 30,
                radius: int = 24, text_color: str = INK) -> None:
    draw.rounded_rectangle(box, radius=radius, fill=rgb(fill), outline=rgb(outline), width=3)
    font = pil_font(font_size)
    x1, y1, x2, y2 = box
    lines = text.split("\n")
    heights = []
    widths = []
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font)
        widths.append(bbox[2] - bbox[0])
        heights.append(bbox[3] - bbox[1])
    total_h = sum(heights) + (len(lines) - 1) * 10
    y = y1 + (y2 - y1 - total_h) / 2
    for line, w, h in zip(lines, widths, heights):
        draw.text((x1 + (x2 - x1 - w) / 2, y), line, font=font, fill=rgb(text_color))
        y += h + 10


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int],
          *, color: str = MUTED, width: int = 5) -> None:
    draw.line([start, end], fill=rgb(color), width=width)
    import math
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    length = 22
    spread = 0.55
    p1 = (end[0] - length * math.cos(angle - spread), end[1] - length * math.sin(angle - spread))
    p2 = (end[0] - length * math.cos(angle + spread), end[1] - length * math.sin(angle + spread))
    draw.polygon([end, p1, p2], fill=rgb(color))


def draw_development_flow() -> None:
    img = Image.new("RGB", (2400, 1520), rgb(WHITE))
    d = ImageDraw.Draw(img)
    title_font = pil_font(44)
    d.text((80, 42), "HR考勤与薪资核算系统开发流程", font=title_font, fill=rgb(DARK_BLUE))
    small = pil_font(24)
    d.text((82, 105), "先锁定数据契约、考勤/工资/法定规则，再并行推进多源接入、核算、安全与验收。", font=small, fill=rgb(MUTED))

    boxes = {
        "A": (120, 210, 600, 360),
        "B": (770, 210, 1510, 360),
        "C": (1730, 210, 2250, 360),
        "D1": (120, 500, 720, 680),
        "D2": (930, 500, 1530, 680),
        "E": (1740, 500, 2250, 680),
        "F": (770, 810, 1530, 980),
        "G1": (120, 1110, 720, 1280),
        "G2": (930, 1110, 1530, 1280),
        "H": (1740, 1110, 2250, 1280),
        "I": (1740, 1360, 2250, 1480),
    }
    rounded_box(d, boxes["A"], "需求与范围确认", LIGHT_BLUE, BLUE, font_size=32)
    rounded_box(d, boxes["B"], "数据契约与考勤/工资/法定规则评审", LIGHT_BLUE, BLUE, font_size=28)
    rounded_box(d, boxes["C"], "交互原型与\n报表口径设计", LIGHT_BLUE, BLUE, font_size=30)
    rounded_box(d, boxes["D1"], "得力 API 接入\n初始化 · next_id · 幂等", LIGHTER_BLUE, BLUE, font_size=29)
    rounded_box(d, boxes["D2"], "致远 OA 与 HR Excel\n字段/附件 · 受控导入 · 冲突", LIGHTER_BLUE, BLUE, font_size=27)
    rounded_box(d, boxes["E"], "统一数据底座\n标准化与工号关联", LIGHT_GREEN, GREEN, font_size=29)
    rounded_box(d, boxes["F"], "考勤与薪资核算引擎\n分段 · 税社保公积金 · 双月结", LIGHT_GREEN, GREEN, font_size=27)
    rounded_box(d, boxes["G1"], "管理后台\n同步 · 考勤 · 薪资 · 权限", LIGHT_ORANGE, ORANGE, font_size=29)
    rounded_box(d, boxes["G2"], "响应式员工端\n考勤 · 工资条 · 假期 · 反馈", LIGHT_ORANGE, ORANGE, font_size=29)
    rounded_box(d, boxes["H"], "联调与金标准验收", LIGHT_RED, RED, font_size=30)
    rounded_box(d, boxes["I"], "试运行 · 月结演练 · 上线", LIGHT_GREEN, GREEN, font_size=27)

    arrow(d, (600, 285), (770, 285))
    arrow(d, (1510, 285), (1730, 285))
    arrow(d, (1140, 360), (420, 500))
    arrow(d, (1140, 360), (1230, 500))
    arrow(d, (720, 590), (1740, 590))
    arrow(d, (1530, 590), (1740, 590))
    arrow(d, (1995, 680), (1300, 810))
    arrow(d, (1150, 980), (420, 1110))
    arrow(d, (1150, 980), (1230, 1110))
    arrow(d, (720, 1195), (1740, 1195))
    arrow(d, (1530, 1195), (1740, 1195))
    arrow(d, (1995, 1280), (1995, 1360))
    img.save(DEV_FLOW_PATH, quality=95)


def draw_system_data_flow() -> None:
    img = Image.new("RGB", (2400, 1260), rgb(WHITE))
    d = ImageDraw.Draw(img)
    d.text((80, 42), "系统数据与考勤薪资产品流程", font=pil_font(44), fill=rgb(DARK_BLUE))
    d.text((82, 105), "得力、致远与 HR Excel 先暂存校验；来源冲突关闭后，考勤月结进入分段与法定核算。",
           font=pil_font(24), fill=rgb(MUTED))

    deli = (80, 210, 540, 360)
    oa = (80, 500, 540, 650)
    excel = (80, 790, 540, 940)
    raw1 = (720, 190, 1250, 370)
    raw2 = (720, 440, 1250, 700)
    raw3 = (720, 770, 1250, 960)
    core = (1410, 350, 1990, 830)
    admin = (2100, 220, 2330, 500)
    portal = (2100, 700, 2330, 980)
    rounded_box(d, deli, "得力云\n考勤 API", LIGHT_BLUE, BLUE, font_size=32)
    rounded_box(d, oa, "致远 OA\n只读数据库", LIGHT_BLUE, BLUE, font_size=32)
    rounded_box(d, excel, "HR 薪资 Excel\n受控导入", LIGHT_BLUE, BLUE, font_size=30)
    rounded_box(d, raw1, "原始打卡层\n游标 · 批次 · 原始载荷", LIGHTER_BLUE, BLUE, font_size=28)
    rounded_box(d, raw2, "OA 暂存与版本层\n组织 · 岗位 · 任职 · 权限\n考勤/调动/薪资字段与附件", LIGHTER_BLUE, BLUE, font_size=23)
    rounded_box(d, raw3, "Excel 暂存与冲突层\n模板 · Sheet/行 · 文件哈希\n来源优先级 · 有效版本", LIGHTER_BLUE, BLUE, font_size=23)
    rounded_box(d, core, "统一 HR 核算服务\n组织/岗位身份 · 考勤月结 · 工资分段\n实际工作日/工时 · 社保公积金 · 个税\n复核发布 · 工资条 · 成本分摊", LIGHT_GREEN, GREEN, font_size=24)
    rounded_box(d, admin, "管理后台\n考勤 · 薪资\n权限 · 审计", LIGHT_ORANGE, ORANGE, font_size=25)
    rounded_box(d, portal, "员工前端\n考勤 · 工资条\n假期 · 反馈", LIGHT_ORANGE, ORANGE, font_size=25)
    arrow(d, (540, 285), (720, 280))
    arrow(d, (540, 575), (720, 570))
    arrow(d, (540, 865), (720, 865))
    arrow(d, (1250, 280), (1410, 470))
    arrow(d, (1250, 570), (1410, 590))
    arrow(d, (1250, 865), (1410, 710))
    arrow(d, (1990, 490), (2100, 360))
    arrow(d, (1990, 690), (2100, 840))
    img.save(DATA_FLOW_PATH, quality=95)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_LATIN
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    specs = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in specs.items():
        s = styles[name]
        s.font.name = FONT_LATIN
        s._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
        s._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
        s._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor.from_string(color)
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True
        s.paragraph_format.keep_together = True

    for name in ("List Bullet", "List Bullet 2", "List Number", "List Number 2"):
        s = styles[name]
        s.font.name = FONT_LATIN
        s._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
        s.font.size = Pt(11)
        s.paragraph_format.space_after = Pt(4)
        s.paragraph_format.line_spacing = 1.167

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("神州HR · 考勤与薪资核算系统 PRD")
    set_run_font(r, size=8.5, color=MUTED, bold=True)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run("内部评审稿  |  ")
    set_run_font(r, size=8.5, color=MUTED)
    add_page_number(p)


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("产品需求文档  /  PRD")
    set_run_font(r, size=11, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("神州HR考勤与薪资核算系统")
    set_run_font(r, size=26, color=DARK_BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("双源直连 · 异动分段计薪 · 薪资复核 · 权限隔离 · 员工自助")
    set_run_font(r, size=14, color=MUTED)

    add_callout(
        doc,
        "产品定位",
        "将得力云原始打卡与致远 OA 考勤、人事调动、组织任职及薪资资料统一到可解释的考勤与薪资账本，支持月内跨部门/岗位分段计薪、考勤工资项、社保公积金、个税、其他增减项、独立复核发布、工资条和成本分摊，并以独立高敏权限域保护薪资数据。",
        fill=LIGHTER_BLUE,
    )

    metadata = [
        ["文档版本", "V1.7", "文档状态", "架构设计输入稿"],
        ["项目优先级", "P0", "创建日期", "2026-07-19"],
        ["需求来源", "业务访谈、流程图与制度补充", "适用范围", "首版产品与验收基线"],
        ["编制", "Codex（基于需求讨论整理）", "保密级别", "内部使用"],
    ]
    add_table(doc, ["项目", "内容", "项目", "内容"], metadata,
              [1.0, 2.1, 1.0, 2.4], font_size=9.0,
              alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT,
                          WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
              first_col_bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("版本记录")
    set_run_font(r, size=11, color=DARK_BLUE, bold=True)
    add_table(doc, ["版本", "日期", "状态", "说明"],
              [["V1.0", "2026-07-17", "评审稿", "双源直连、字段统一及首版考勤闭环"],
               ["V1.1", "2026-07-17", "评审稿", "补签、跨夜加班、地区班次、假别与销假规则"],
               ["V1.2", "2026-07-17", "架构准入评审稿", "补充领域、状态、数据、容量与架构准入"],
               ["V1.3", "2026-07-18", "架构设计输入稿", "确认 OA 主数据权限；补充地图、三端与防越权"],
               ["V1.4", "2026-07-18", "架构设计输入稿", "组织自动/手动同步、有效期版本与历史归属"],
               ["V1.5", "2026-07-19", "工资模块扩展评审稿", "OA 调动分段计薪、工资项目/成本、审批工资条及 PAYROLL 独立权限"],
               ["V1.6", "2026-07-19", "架构设计输入稿", "确认税社保公积金、OA+HR Excel 薪资来源、工作日/工时折算、调动日归新部门、无 MFA 与独立复核发布"],
               ["V1.7", "2026-07-19", "集成校准版", "核对 V80 数据字典及致远/得力官方文档；自建表采用版本化逻辑映射，修正得力人员绑定、附件与接口边界"]],
              [0.8, 1.2, 1.0, 3.5], font_size=7.6)
    # The next heading owns the cover-to-content page break. A standalone
    # page-break paragraph after a table can be pushed to the following page
    # by LibreOffice and create a blank page.


def add_contents(doc: Document) -> None:
    heading = add_heading(doc, "文档目录与阅读说明", 1)
    heading.paragraph_format.page_break_before = True
    sections = [
        "1. 需求概述", "2. 用户与权限", "3. 产品范围与优先级", "4. 产品及开发流程",
        "5. 信息架构与用户场景", "6. 功能需求", "7. 核算与报表口径", "8. 数据契约与外部依赖",
        "9. 交互与响应式要求", "10. 非功能与安全要求", "11. 数据指标与埋点", "12. 验收标准",
        "13. 项目计划与风险", "14. 架构决策确认与剩余数据契约", "15. 架构准入评审", "附录A. 字段字典与参考资料",
    ]
    for item in sections:
        add_bullet(doc, item)
    add_callout(doc, "阅读指引", "标记为“默认方案”的内容可用于首版设计与估算；标记为“上线前确认”的内容若未澄清，开发可以搭建框架，但 QA 无法确认最终核算结果。")
    add_body(doc, "本文档针对内部定制系统，不做市场规模与商业化论证；“方案对比”重点比较现有 Excel 手工对账、得力云单源能力、致远 OA 审批能力与本系统双源核算闭环。")
    doc.add_page_break()


def build_prd_body(doc: Document) -> None:
    add_heading(doc, "1. 需求概述", 1)
    add_heading(doc, "1.1 业务背景", 2)
    add_body(doc, "当前考勤人员需要分别从得力考勤机云端和致远 OA 自建单据中取得数据。两侧 Excel 导出格式不一，考勤人员仍需人工整合字段，再判断请假、补签、加班、出差/外出与打卡记录的覆盖关系并完成汇总。两个来源各自记录一部分事实：得力保存实际打卡，致远保存审批通过的例外，但缺少统一的按人按日核算和可追溯结果。")
    add_body(doc, "本项目建设一套内部考勤核算与员工查询平台，直接通过得力官方 API 拉取上线后的增量打卡，通过致远数据库只读查询请假/销假、加班、出差/外出、补签等自建单据，先经过版本化查询映射和人员来源绑定，再形成日明细、异常、专项清单、个人/部门/公司汇总及月结快照。本 PRD 只定义自建表的逻辑契约，不固化物理表名、字段名或 SQL。")
    add_body(doc, "致远组织架构还会发生改名、改编码、调整上级、停用、拆分合并或来源 ID 复用。如果本系统只保存致远当前值并按来源 ID 聚合，历史采购部可能被显示成后来的销售部，当前负责人也可能错误继承旧部门数据。因此本系统必须同步当前组织副本，同时保存组织身份、来源绑定和有效期历史版本。")
    add_body(doc, "工资模块需要继续消费致远 OA 的人事调动、岗位/职级和生效时间，以及本系统已月结的考勤结果。员工若在月中由销售部调到采购部，且两个部门或岗位适用的工资标准不同，系统必须形成两个互不重叠的工资分段，分别展示部门工资、岗位工资、考勤产生的工资项和其他增减项，再汇总为当月工资；员工应发工资与部门成本承担必须分别核算、分别勾稽。")
    add_callout(doc, "已确认边界", "不处理上线前历史打卡；得力打卡与致远考勤/人事事实均直接取数，致远自建业务表使用只读数据库接入，物理表名与字段本次忽略。OA 工号与得力 employee_num 一致只用于生成候选匹配，不能替代对得力 ext_id 的唯一绑定校验。薪资资料当前来自致远 OA 记录/附件及 HR 受控 Excel，最终均沉淀为本系统可追溯版本。P0 计算社保、公积金和个人所得税，但不生成银行付款文件；银企直连和税务直接申报列入 P2。")

    add_heading(doc, "1.2 核心问题", 2)
    problems = [
        ["P1", "双源割裂", "打卡和 OA 审批单据分散，需人工按工号、日期、时间段反复匹配。"],
        ["P2", "规则计算繁琐", "迟到、早退、缺卡、补签时效、销假、跨夜加班、外出与假期扣减依赖人工公式和经验。"],
        ["P3", "异常定位慢", "数据缺失、迟报、撤销单据和冲突单据难以及时发现，月底集中处理。"],
        ["P4", "结果不可解释", "员工或管理者提出异议时，需要回查多个系统和 Excel，缺少统一证据链。"],
        ["P5", "缺少自助查询", "员工无法集中查看个人考勤、年假余额、当日签到排行、留言互动和反馈处理进度。"],
        ["P6", "权限和审计不足", "公司、部门、个人、位置等数据需要不同授权范围，人工文件难以持续控制。"],
        ["P7", "组织变化污染历史", "名称、编码、上级和 OA ID 均可能变化或复用；直接覆盖会混淆历史统计并造成跨组织权限继承。"],
        ["P8", "工资分段与保密复杂", "月内调动、调岗、调薪及考勤变化会改变工资项目和成本归属；人工拆段难核对，普通 HR、负责人、系统管理员或抓包用户又不得越权取得薪资。"],
    ]
    add_table(doc, ["编号", "问题", "具体表现"], problems, [0.65, 1.25, 4.6], first_col_bold=True)

    add_heading(doc, "1.3 产品目标与成功指标", 2)
    goals = [
        ["G1 数据自动化", "双源正常时，新数据自动进入统一考勤账本。", "测试样本拉取完整率 100%；同源幂等率 100%；不重不漏。"],
        ["G2 降低手工工作", "考勤员只处理系统识别出的异常和冲突。", "正常场景自动核算率目标 >=90%［试运行验证］；月度人工耗时下降目标 >=70%［试运行验证］。"],
        ["G3 结果准确可解释", "任一结果可追溯至班次、打卡、OA 单据、规则和调整。", "金标准案例计算一致率 100%；汇总与下钻明细差异为 0。"],
        ["G4 员工与管理自助", "按权限查看个人、部门和公司数据。", "常用查询 P95 <=2 秒；员工可在 360px 及以上宽度完成查询和反馈。"],
        ["G5 月结可审计", "冻结正式结果，同时管理月结后迟报变化。", "人工调整、位置查看、导出、月结和反月结操作留痕率 100%。"],
        ["G6 分段算薪与保密", "按人事异动有效期和考勤事实自动形成可解释工资。", "金标准分段与工资项目一致率 100%；工资/成本勾稽差异为 0；薪资越权成功数为 0。"],
    ]
    add_table(doc, ["目标", "价值", "首版验收指标"], goals, [1.35, 2.25, 2.9], font_size=8.9, first_col_bold=True)

    add_heading(doc, "1.4 方案对比", 2)
    compare = [
        ["现有 Excel 手工对账", "灵活、短期无需开发", "重复劳动、公式易错、权限与审计弱、员工无法自助", "淘汰为主；保留导出格式兼容"],
        ["仅使用得力云", "稳定获取设备与原始打卡", "无法覆盖公司在致远中的自建请假、加班、外出规则", "作为打卡事实来源"],
        ["仅使用致远 OA", "审批流程和单据完整", "缺少独立考勤机的实际打卡事实及统一核算", "作为审批事实来源"],
        ["本系统双源核算", "统一数据、自动重算、异常复核、月结、响应式自助", "需明确班次、规则、年假额度和权限边界", "推荐方案"],
    ]
    add_table(doc, ["方案", "优势", "不足", "结论"], compare, [1.35, 1.55, 2.45, 1.15], font_size=8.6, first_col_bold=True)

    add_heading(doc, "2. 用户与权限", 1)
    add_heading(doc, "2.1 用户角色", 2)
    roles = [
        ["系统管理员", "维护数据源、账号、角色和运行参数", "默认不拥有员工考勤和定位明细权限"],
        ["考勤主管", "授权范围异常复核、人工调整、月结、反月结、报表", "致远 OA 授权组织范围"],
        ["考勤专员", "处理授权范围内异常、发起重算、导出", "授权组织"],
        ["HR/人事管理员", "授权范围考勤查询；班次日历、假别、规则、月结和制度", "致远 OA 授权组织范围"],
        ["薪酬制表员", "维护授权薪资组输入、执行算薪、处理异常并提交复核", "独立 PAYROLL 授权的法人/薪资组；不能复核或发布本人批次"],
        ["薪酬复核发布人", "只读复核差异与勾稽，退回或发布工资批次与工资条", "独立 PAYROLL 授权范围；不能修改源事实、公式、金额或自授权限"],
        ["薪资权限管理员", "按预设角色模板配置薪资授权、有效期与范围", "不能查看薪资数据、不能给自己授权、不能兼任薪资数据角色"],
        ["财务成本/税费人员", "查看法定申报与会计成本所需字段", "独立授权的已发布批次；默认不看假勤原因、公式或无关分项"],
        ["部门负责人", "查看本部门看板、下属明细和反馈进度", "致远 OA 当前组织及可穿透下级范围"],
        ["老板/公司管理层", "查看授权范围考勤明细、汇总和看板；按授权导出", "致远 OA 授权范围；敏感字段另授权"],
        ["普通员工", "查看本人考勤、假期、制度和反馈；查看受控签到排行并留言", "本人数据 + 排行公开字段"],
        ["审计员", "查看月结版本、调整、导出和访问日志", "授权范围内只读"],
    ]
    add_table(doc, ["角色", "核心职责", "默认数据范围"], roles, [1.35, 3.0, 2.15], font_size=8.9, first_col_bold=True)

    add_heading(doc, "2.2 权限模型", 2)
    add_body(doc, "致远 OA 是员工、组织层级、员工任职关系和组织穿透数据范围的权威来源。本系统不自行扩大组织范围；最终权限取“数据域 × 功能动作 × 本系统角色 × 致远 OA 当前组织范围 × 对象关系 × 字段分级 × 业务目的 × 记录/月状态”的交集。默认拒绝；菜单可见、知道 URL、获得对象 ID 或前端隐藏字段均不构成授权。")
    for text in [
        "功能权限：查看、配置、复核、调整、重算、月结、反月结、导出、发布。",
        "数据范围：本人、本部门、可穿透下级部门、指定组织或全公司，具体节点和向下穿透关系均使用致远 OA 最新有效授权快照；OA 组织 ID、编码和名称须先按有效期解析为本系统内部组织身份，不得直接作为永久授权主键。",
        "字段权限：请假类型、手机号、精确坐标、地点文字、外勤照片等。",
        "月份状态：开放月份允许核算与调整；已月结月份只读，除非授权反月结。",
        "对象关系：员工本人、当前下属、授权组织成员、操作人创建对象等，由服务端根据会话和权威关系计算，不接受客户端自报 employee_id/user_id 决定权限。",
        "数据域：ATTENDANCE、LEAVE、LOCATION、PAYROLL、AUDIT 相互隔离；角色在一个域的权限不得自动继承到另一个域。致远组织范围只是薪资授权上限，不能直接授予薪资权限。",
    ]:
        add_bullet(doc, text)
    permission_rows = [
        ["账号/角色/数据源配置", "管理", "查看", "查看状态", "否", "否", "只读"],
        ["公司/部门汇总", "默认无业务权限", "OA 授权范围", "OA 授权范围", "OA 可穿透范围", "仅本人", "授权范围"],
        ["个人明细与 OA 单据", "默认无", "OA 授权范围", "OA 授权范围", "OA 授权下属", "仅本人", "授权范围"],
        ["精确位置查看", "默认无", "单独授权", "单独授权", "默认无", "本人该次记录", "默认无"],
        ["人工调整/重算", "否", "否", "全公司", "提交意见", "提交反馈", "否"],
        ["月结/反月结", "否", "只读", "均可，需理由", "否", "否", "否"],
        ["普通报表导出", "默认无", "OA 授权范围", "OA 授权范围", "OA 范围汇总", "本人", "授权范围"],
        ["签到排行/留言", "内容治理", "查看", "查看/治理", "查看", "查看/留言", "审计"],
    ]
    add_table(doc, ["能力", "系统管理员", "老板/管理层", "人事/考勤", "部门负责人", "员工", "审计"],
              permission_rows, [1.35, 0.82, 0.9, 0.92, 0.92, 0.77, 0.82], font_size=7.45,
              alignments=[WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.CENTER] * 6,
              first_col_bold=True)
    add_callout(doc, "权限解释", "老板、人事和考勤角色能够查看的组织范围以致远 OA 当前授权为上限；只有致远授予全公司范围时才可查看全公司。P0 的位置权限统一为 attendance.location.map，控制 GPS/外勤地点文字与单点地图；不提供精确位置导出。病假原因、孕检/哺乳原因等敏感字段仍须本系统独立授权。员工只看本人明细。", fill=LIGHT_RED, label_color=RED)

    add_heading(doc, "2.2.1 薪资独立权限矩阵", 3)
    payroll_permission_rows = [
        ["员工本人", "本人已发布工资条、历史工资条及本人分段组成", "草稿/待复核结果、他人工资和批量接口；P0 不采集银行账号"],
        ["部门负责人", "默认无工资明细；显式授权后仅看满足小样本保护的本成本中心汇总", "个人总工资、另一部门分段、税社保明细和个人导出"],
        ["HR 普通/考勤", "只看调动事实、考勤结果和资料是否齐备，不看金额", "工资档案、金额、薪级、税社保明细和工资条"],
        ["薪酬制表", "授权法人/薪资组的输入、计算、异常和明细", "复核/发布本人创建批次、修改自身权限或绕过独立复核"],
        ["薪酬复核发布", "授权范围的只读差异与明细、退回或发布", "修改源事实/公式/金额、发布本人制表批次、配置自身权限"],
        ["薪资权限管理", "使用预设模板配置账号、角色、范围、字段、用途和有效期", "读取工资明文、任意通配授权、自授权或兼任薪资数据角色"],
        ["财务", "已发布批次的法定申报汇总与会计成本分摊", "未发布结果、假勤原因、工资公式和无关薪资项目"],
        ["老板/高管", "经显式授权且满足小样本保护的人工成本汇总", "因职级直接查看个人工资、税社保明细或工资排行"],
        ["系统管理员", "技术配置、任务状态和授权流程配置", "工资明文、解密密钥、生产直库查询和自行授予薪资角色"],
        ["审计员", "版本、复核发布链、访问/导出日志；默认金额脱敏", "修改、发布；无审计工单的全量明文查询"],
    ]
    add_table(doc, ["薪资角色", "默认允许", "强制禁止"], payroll_permission_rows,
              [1.25, 2.7, 2.55], font_size=8.1, first_col_bold=True)
    add_callout(doc, "薪资授权公式", "账号有效 × PAYROLL 角色 × 功能动作 × 独立薪资数据范围 × 对象关系 × 字段白名单 × 使用目的 × 工资批次状态 × 最近登录/短会话状态 × 风险策略。工资归属部门或 OA 穿透范围不等于查看工资的授权；OA 组织同步不得自动新增薪资授权。", fill=LIGHT_RED, label_color=RED)

    add_heading(doc, "2.3 数据分级与授权落点", 2)
    classification_rows = [
        ["A 级：受控发布", "签到排行名次、展示名、允许的时间粒度", "仅通过独立发布投影输出；按榜单范围访问，不可反查原始卡或 OA 单据"],
        ["B 级：内部汇总", "公司/部门出勤率、异常数、趋势", "按组织范围查看；无个人明细权限时禁止下钻或返回可识别人员字段"],
        ["C 级：个人业务", "个人打卡、OA 单据、假期余额、反馈、日结果", "本人或具备组织明细权限的角色可见；普通导出继续受字段权限约束"],
        ["D 级：敏感信息", "GPS/外勤地点名称与地址、精确坐标、外勤照片、病假原因、孕检/哺乳原因", "独立权限、最小展示、必要时二次确认；位置 P0 不导出，查看均审计"],
        ["E 级：高度敏感", "工资条、薪资明细、奖金、个税、社保公积金基数与身份证件", "PAYROLL 独立数据域和授权对象；字段加密、本人/薪酬专岗白名单、短会话/敏感操作重新登录、水印、禁止大屏和默认批量导出"],
    ]
    add_table(doc, ["分级", "典型数据", "强制控制"], classification_rows,
              [1.25, 2.05, 3.2], font_size=8.35, first_col_bold=True)
    for text in [
        "所有业务表和访问令牌均携带 company_id；P0 虽为单公司，数据隔离键不得省略，避免后续多公司扩展时重构主键与权限链路。",
        "服务端授权统一计算“功能 × 当前 OA 授权快照已解析的内部组织身份 × 字段分级 × 记录/月份状态”；历史报表按发生日 OrganizationVersion 归属，访问仍按当前致远授权判断。",
        "组织单纯改名、改编码或调整上级并确认业务身份连续时，可沿用同一 organization_identity_id；部门撤销、OA ID/编码复用或业务性质改变时必须创建新身份。当前新组织负责人不得因来源 ID 或编码相同继承旧组织的详情、数量、导出、缓存、地图或薪资访问权。",
        "致远组织或权限同步失败时不得生成新的本地授权；账号在致远被停用、删除或移出范围后，本系统下一次成功同步即撤销对应访问，权限快照版本须进入审计日志。",
        "敏感导出与普通导出使用不同权限点和审计事件；禁止通过列表缓存、搜索索引、导出任务或前端隐藏绕过字段分级。",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "2.4 服务端防越权与权限可维护性", 2)
    for text in [
        "建立统一授权策略服务/中间件作为 Policy Enforcement Point。列表、详情、统计数量、聚合、搜索建议、批量接口、附件、地图、导出、WebSocket/推送和后台任务在访问数据前逐次授权，禁止各页面自行拼接权限条件。",
        "查询必须先把 company_id、当前致远授权组织、对象关系和字段策略下推到数据库/查询层，再分页、聚合和序列化；禁止先查全量再在前端或内存中过滤。",
        "URL、路径、查询参数、请求体和自定义 Header 中的 employee_id、organization_id、role、scope 均视为不可信输入。本人接口从登录会话推导员工 ID；管理接口即使 ID 合法，也必须重新验证对象级和字段级授权。",
        "对外对象标识使用不可预测的 UUID/ULID 或等效 opaque ID，内部连续主键不暴露；但不可预测 ID 只降低枚举效率，不能代替服务端授权。无权对象与不存在对象对外采用不可枚举的一致响应，真实原因只写安全审计。",
        "响应 DTO 采用字段白名单，按数据域和字段策略组装，禁止直接序列化数据库实体或依赖前端删字段；写接口同样采用可写字段白名单，防止批量赋值越权。",
        "缓存键至少包含 company_id、subject_id、OA 授权快照版本、组织绑定/版本号、权限策略版本和查询摘要；组织绑定变化或权限收缩后主动失效。位置和薪资等 D/E 级响应设置 private/no-store，不进入共享 CDN、浏览器持久缓存、Service Worker 或普通离线包。",
        "致远授权快照或 PAYROLL 授权快照超过可配置新鲜度阈值时，组织明细、批量导出、位置、薪资详情、复核和发布均 fail closed；可以继续查看不含敏感明细的最近成功聚合，但必须标记数据过期。同步恢复后只按新快照开放。",
        "异步导出在创建、执行和下载三个时点重新授权；冻结申请时的组织/字段范围，下载使用绑定申请人的短时一次性令牌。文件名、对象存储路径和 URL 不可枚举，过期或权限已撤销立即拒绝。",
        "若使用 WebSocket/SSE，握手、订阅主题和每条消息均按当前会话、company_id、致远范围和字段策略授权；禁止先广播全公司数据再由前端过滤，注销或撤权后立即断开。",
        "权限策略配置须版本化，支持预览“某用户为何可/不可访问某记录/字段”、影响范围分析、定时生效与回滚。薪资权限管理员只可使用预设角色模板，授权人与被授权人相同、权限管理员兼任薪资数据角色或通配范围时原子拒绝；变更立即审计并通知固定 HR/安全联系人，但不形成审批流程。",
    ]:
        add_bullet(doc, text)
    add_callout(doc, "薪资安全边界", "PAYROLL 为本次正式独立数据域，不能复用考勤角色、接口、DTO、缓存、导出或密钥。默认仅员工本人查看已发布工资条，明确授权的薪酬岗位按职责处理工资；系统管理员、部门负责人、普通 HR、考勤角色和高管均不因现有身份自动获得个人薪资。", fill=LIGHT_RED, label_color=RED)

    add_heading(doc, "3. 产品范围与优先级", 1)
    add_heading(doc, "3.1 P0：首版闭环", 2)
    p0 = [
        "得力 API 初始化、next_id 增量同步、幂等入库和同步监控；不迁移上线前历史打卡。",
        "致远数据库只读查询请假/销假、加班、出差/外出登记、补签等自建单据；物理表名/字段通过版本化 OAQueryMappingProfile 适配，本 PRD 不固化。",
        "以 EmployeeSourceBinding 绑定致远人员与得力 user_id/ext_id；OA 工号与得力 employee_num 一致仅生成候选，唯一校验通过后才可正式入账。",
        "从致远 OA 首次全量、周期自动和手动全量同步员工、组织节点/边、编码、任职关系和组织数据范围；保存合同字段完整源副本、内部组织身份、来源绑定、不可变有效期版本、当前投影、变更复核与历史路径。",
        "总部季节班次、地区班次、个人有效期班次、工作日历、特殊工时和规则版本配置；数据模型预留跨日、轮班、弹性和临时排班扩展字段，P0 不实现其计算界面与算法。",
        "按人按日核算正常、迟到、早退、缺卡、缺勤、补签、请假/销假、出差/外出和跨夜加班。",
        "异常工作台、人工调整、重算、月结、反月结和月结后差异。",
        "原始打卡、OA 单据、日明细、异常、专项清单、个人/部门/公司汇总和 Excel 导出。",
        "管理后台、响应式员工端与 1920×1080 管理驾驶舱大屏；个人、部门和公司三级看板按权限展示；当日签到排行与留言互动。",
        "年假、丧假、护理假、病假、事假规则与台账；哺乳/孕检特殊时段；制度与反馈。",
        "功能角色、致远组织数据范围、敏感字段权限、GPS/外勤打卡地图查看和全链路审计。",
        "致远人事调动、岗位/职级及调薪事实只读同步；工资方案、项目、标准、公式、折算和取整规则按有效期版本化。",
        "工资数据准备：只读同步致远 OA 薪资记录与关联文件，受控导入 HR Excel；保存来源、文件哈希、导入批次、版本和冲突处置，最终统一沉淀在本系统。",
        "月内跨部门/岗位分段计薪、按实际工作日/计薪工时折算、考勤工资项、其他收入/扣减、社保、公积金、个人所得税、差异复核发布、工资月结、追溯补发追扣、员工工资条和部门/成本中心分摊。",
        "PAYROLL 独立授权、制表与复核发布职责分离、短会话/敏感操作重新登录、可信网络、字段加密、水印、三次导出鉴权、运维隔离和全链路薪资审计；不建设 MFA。",
    ]
    for item in p0:
        add_bullet(doc, item)

    add_heading(doc, "3.2 P1/P2 后续范围", 2)
    future = [
        ["P1", "夜班、轮班、弹性工时与排班导入；更多 OA 单据；按当前组织重述的非正式分析口径；跨地区/跨法人等高级假期策略；互动内容增强；制度阅读确认；单点登录；定时报表和站内通知。"],
        ["P2", "税社保政策自动更新与多地区完整政策库、银企直连、税务/社保直接申报、多公司多账套、完整预算激励/佣金、自定义报表、趋势预警、企业微信/钉钉入口或独立移动端。"],
        ["明确排除", "P0 银行付款文件生成、实时位置跟踪、持续轨迹采集、人脸/指纹特征管理、向得力或致远写回、招聘和完整绩效模块。"],
    ]
    add_table(doc, ["阶段", "范围"], future, [1.0, 5.5], first_col_bold=True)
    add_callout(doc, "范围说明", "首批已确认没有夜班、轮班、弹性工时或临时排班；P0 只实现现有行政班与行政班结束后的跨午夜加班。ShiftDefinition、ShiftSegment、ShiftAssignment 和考勤日归属策略仍保留扩展能力，后续增加班型时新增策略版本，不改写历史结构。员工在 OA 发起审批，本系统继续只读，不向致远写回。", fill=LIGHT_ORANGE, label_color=ORANGE)

    add_heading(doc, "4. 产品及开发流程", 1)
    add_heading(doc, "4.1 系统数据流程", 2)
    add_figure(doc, DATA_FLOW_PATH, "图1  系统数据与产品流程（基于用户原始流程图重绘）", width=6.25,
               alt_text="得力云API、致远OA只读数据库及HR薪资Excel分别进入暂存与版本层，统一服务完成考勤、工资分段、税社保公积金和个税核算，并向管理后台和员工前端提供最小授权视图。")
    add_body(doc, "原始数据不可由业务用户覆盖。致远组织、岗位、任职、调动、薪资字段/附件和授权，以及 HR Excel 均先进入暂存；结构校验、去重和来源冲突关闭后，发布内部身份、有效期版本与工资有效数据。考勤月结、法定政策、个税年度累计和全部来源版本进入同一工资输入快照。前端和后台只通过统一服务访问，不直连考勤或薪资数据库。")

    add_heading(doc, "4.2 开发流程", 2)
    add_figure(doc, DEV_FLOW_PATH, "图2  开发流程与质量门", width=6.25,
               alt_text="需求和考勤、工资、法定规则评审后，并行开发得力接口、致远与Excel接入、考勤薪资法定核算、管理后台和员工前端，最后通过金标准与双月结演练上线。")
    add_body(doc, "数据契约、OA/Excel 来源优先级、考勤规则、工资项目折算、税社保政策与个税期初累计是第一个质量门；金标准准确性、多源完整性、工资/法定金额勾稽、制表与复核发布互斥、权限隔离和多端适配是第二个质量门。任何质量门未通过，不进入下一阶段。")

    add_heading(doc, "4.3 日常业务闭环", 2)
    daily_steps = [
        "定时拉取得力增量打卡，并查询致远开放考勤期间内的请假/销假、加班、出差/外出、补签等单据。",
        "按计划自动同步致远组织架构；管理员也可点击“同步组织架构”立即全量比对。普通变化自动版本化，疑似复用、拆分合并或连续性未知进入复核。",
        "按工号匹配员工，完成去重、格式校验、状态映射和标准化。",
        "按受影响员工与日期重算日考勤，生成结果、解释和异常。",
        "考勤专员在异常工作台复核，必要时发起人工调整并填写原因。",
        "月末执行预检查、全量重算和月结，冻结正式版本并生成报表。",
        "月结后迟报或单据变化进入差异清单；授权反月结后生成新版本，旧版本保留。",
        "工资期间收集 OA 调动/调薪、发生时组织岗位、OA 薪资记录/附件、受控 HR Excel、工资配置、本次考勤月结版本和其他输入；冲突关闭后锁定快照并按有效期生成工资分段。",
        "薪酬制表员完成计算和异常处理；不同账号的复核发布人只读确认差异、税社保公积金、个税与勾稽，可退回或发布，不设置额外业务审批节点。",
        "工资发布/关闭后迟到的人事、考勤或配置变化只形成工资差异；正式版本不得破坏性覆盖，只能在新运行中补发、追扣或冲销。",
    ]
    for step in daily_steps:
        add_number(doc, step)
    state_rows = [
        ["开放 OPEN", "双源持续同步，结果动态更新，可处理异常", "允许进入核算；禁止直接跳到已月结"],
        ["核算中 CALCULATING", "执行受影响范围或全月核算", "任务成功后回到开放或进入待月结；失败不改正式版本"],
        ["待月结 READY_TO_CLOSE", "全量重算完成，等待预检查通过", "硬阻断项存在时退回开放，禁止月结"],
        ["已月结 CLOSED", "正式结果和报表冻结", "新数据只形成月结后差异；不可原地修改"],
        ["已反月结 REOPENED", "考勤主管填写原因后重新开放", "旧版本永久保留；下一步只能进入核算中"],
    ]
    add_table(doc, ["月份状态", "含义", "允许操作"], state_rows, [1.2, 3.0, 2.3], first_col_bold=True)

    add_heading(doc, "5. 信息架构与用户场景", 1)
    add_heading(doc, "5.1 信息架构", 2)
    ia_rows = [
        ["管理后台", "工作台；同步中心；组织/岗位镜像；历史版本；变更复核；考勤核算；异常/月结/报表；工资数据准备/导入；调动影响；工资方案/项目/标准；税社保公积金；分段核算；差异复核发布/月结；工资条；成本分摊；权限；位置；审计"],
        ["员工前端", "首页；我的考勤；我的工资；个人看板；公司/部门看板（按权限）；当日签到排行与留言；我的假期；制度中心；反馈中心；个人设置"],
        ["管理驾驶舱", "授权范围考勤总览；组织对比；趋势；异常结构；数据新鲜度；月结状态；只展示考勤聚合，不展示个人位置、工资或可推算个人工资的成本小组"],
        ["公共能力", "登录认证；消息提示；筛选搜索；详情下钻；Excel 导出；数据新鲜度；附件查看；错误与空状态"],
    ]
    add_table(doc, ["产品区", "一级模块"], ia_rows, [1.4, 5.1], font_size=9.1, first_col_bold=True)

    add_heading(doc, "5.2 核心用户场景", 2)
    scenarios = [
        ["S1 月度核算", "考勤专员", "月末打开工作台，确认双源同步正常，处理未匹配和异常，执行全量重算、预检查和月结，导出正式汇总。"],
        ["S2 员工异议", "普通员工", "从个人月历打开异常日期，查看班次、打卡、OA 单据与计算依据，提交关联反馈并跟踪处理。"],
        ["S3 部门管理", "部门负责人", "查看本部门出勤、迟到、请假和加班趋势，按权限下钻到员工日期，但默认不能查看精确位置。"],
        ["S4 迟报处理", "考勤主管", "得力设备补传已月结日期的打卡，系统产生月结后差异；主管决定保留原结果或授权反月结。"],
        ["S5 位置核验", "授权考勤员", "打开一条 GPS/外勤打卡，二次确认查看原因后读取源记录已有地点；系统记录访问日志。"],
        ["S6 同步故障", "系统管理员", "同步中心显示签名、限流或数据库查询失败；管理员修复后从原水位安全重试，既有结果仍可查询。"],
        ["S7 跨夜加班", "普通员工/考勤员", "员工凌晨下班，系统把凌晨卡归入前一考勤日；48 小时内提交并最终获批后重算，超期则加班认定为 0。"],
        ["S8 提前返岗", "普通员工/HR", "批准假期内形成完整有效打卡区间时，系统先恢复实际出勤并生成 DETECTED；余额返还按当期 AUTO_FULL_PUNCH、OA_CANCEL_ONLY、HR_CONFIRM 或 HYBRID 策略处理，零散单卡提示复核。"],
        ["S9 地图核验", "员工/授权管理者", "在一条 GPS 或外勤打卡详情中点击“查看地图”，系统校验本人或致远组织范围及位置权限后，在地图弹层标记该次打卡位置并记录审计。"],
        ["S10 组织变更", "HR/组织数据管理员", "点击“同步组织架构”执行全量比对，查看改名、改编码、调父级及疑似 ID 复用；普通变化自动版本化，无法判断的变更确认延续或新组织，并预览受影响员工、月份和权限。"],
        ["S11 月中调动算薪", "薪酬制表/复核/员工", "员工 7 月 1-15 日在销售部、16-31 日在采购部；调动生效日归新部门。系统生成两个工资分段，按项目配置使用实际工作日或计薪工时折算，并归集对应日期考勤工资项；复核发布后员工查看两段明细与月度合计。"],
        ["S12 工资追溯", "薪酬复核/财务", "工资关闭后收到历史生效的调动、调薪或考勤差异，系统不覆盖原工资，生成引用原月份/分段/项目的补发追扣建议，在新工资运行中复核发布。"],
    ]
    add_table(doc, ["场景", "主要角色", "用户旅程"], scenarios, [1.15, 1.25, 4.1], font_size=8.8, first_col_bold=True)

    add_heading(doc, "6. 功能需求", 1)
    add_heading(doc, "6.1 数据源配置与同步中心（FR-01，P0）", 2)
    add_body(doc, "输入：得力 App-Key/App-Secret、source_instance_id、初始化状态和 next_id；致远只读数据库连接、经 DBA 审核并发布的查询视图/模板、OAQueryMappingProfile 和调度周期。生产环境不允许业务管理员录入任意 SQL；密钥只允许录入或替换，页面始终掩码显示。")
    for item in [
        "得力正式上线时执行一次初始化；初始化前历史记录不采集。普通用户不能再次初始化。",
        "得力从持久化 next_id 开始，每页最多 500 条，持续拉取至空集合；只有该页数据全部入库后才能推进游标。",
        "以来源记录 ID 去重；设备离线补传按 check_time 归入实际考勤日并触发重算。",
        "致远使用独立只读账号读取自建业务表或 DBA 提供的只读视图；员工、标准组织、岗位、任职和附件元数据可由同一可替换接入适配层读取，具体通道在详细设计冻结。本次不要求提供自建表名或字段。",
        "每类自建对象配置独立 OAQueryMappingProfile：查询模板/视图、逻辑对象、业务键、人员和时间映射、原始状态到标准状态映射、增量与删除/撤销策略、单据关联、附件定位、时区/空值规则及能力标志。",
        "自建单据有可靠更新时间时使用复合游标并设置重叠窗口；没有时周期扫描开放期间并比较内容指纹。首次提交、审批完成、业务生效、关联单号等逻辑值允许由单字段、组合字段或受控推导取得；无法可靠取得时按对象能力降级或进入人工复核。",
        "组织自动同步先进入暂存区并计算字段指纹：无变化只更新 last_seen_at；有变化则自动关闭旧 OrganizationVersion、创建新版本并更新当前组织投影，禁止覆盖或删除历史版本。",
        "同步中心提供“同步组织架构”手动按钮，具备 organization.sync.execute 权限的用户可立即发起一次全量比对；按钮展示新增、停用、改名、改编码、调整上级和疑似复用数量，重复点击受运行锁和幂等键保护。",
        "手动同步没有发现变化时只生成 OrganizationSyncBatch 审计记录，不生成空组织版本；发现普通变化时按同一版本规则发布，疑似 ID/编码复用、拆分、合并或无法判断的变化进入组织变更复核。",
        "组织批次发布前校验节点、父子边、任职和授权范围数量，检查重复来源 ID/编码、孤儿节点、层级环、悬空任职/授权引用、完整快照标志及总哈希；组织图、任职和授权快照必须同批原子发布。",
        "组织在单次查询中缺失时先标记 MISSING_CANDIDATE，不得直接删除或关闭；仅在来源提供可靠停用/删除语义，或已确认完整的全量快照连续缺失并通过复核后，才关闭版本与绑定。",
        "查询模板、映射配置和数据源连接均版本化，记录字段指纹、能力清单、脱敏样本试运行结果与回滚点；字段缺失、类型变化或未知状态时整批隔离，不得静默发布。",
        "同步中心展示最近成功时间、水位、批次数量、新增/更新/失败数、错误码、重试次数和数据新鲜度。",
    ]:
        add_bullet(doc, item)
    add_callout(doc, "异常规则", "任何一页写入失败时不得推进该页得力游标，之前已原子提交的页面保留；OA 暂存批次校验失败时不得发布为正式版本，查询为空也不得推断源数据已删除；连续失败 3 次或超过 15 分钟未成功时触发告警。", fill=LIGHT_RED, label_color=RED)

    add_heading(doc, "6.2 人员、组织与工号映射（FR-02，P0）", 2)
    add_body(doc, "致远 OA 是员工、组织、任职、账号状态和组织数据范围的权威来源，本系统保存可查询的当前副本和不可变历史。本系统分别以内部 employee_id 与 organization_identity_id 作为永久身份主键；致远人员 ID、组织 ID、组织编码、名称和工号均作为带来源与有效期的外部属性或绑定，不得直接充当永久业务身份。得力官方以 ext_id 作为外部员工关联标识，employee_num 是独立工号字段；两侧工号一致只能生成候选绑定，不能直接认定 ext_id 一致。")
    for item in [
        "EmployeeSourceBinding 保存 source_instance_id、deli_user_id、deli_ext_id、deli_employee_num、oa_employee_source_id、oa_employee_no、internal employee_id、effective_from/to、binding_status、match_method 和确认审计。",
        "绑定状态为 CANDIDATE、CONFIRMED、CONFLICT、CLOSED；仅 CONFIRMED 可参与打卡核算。ext_id 为空、重复、冲突或人员工号不唯一时进入异常队列。",
        "OA 工号与得力 employee_num 相同可自动生成 CANDIDATE，但须校验两侧人员唯一、在职范围和 ext_id 后才能确认；不得按姓名、手机号自动合并人员。",
        "不得按姓名、手机号自动合并人员；人工映射需记录操作人、原因、生效时间和旧值。",
        "入职、离职、调岗、组织名称/编码/上级/状态和账号停用以致远同步结果及其生效时间为准；本系统只保存副本、版本和复核结论，不得反向修改或覆盖致远。",
        "OrganizationIdentity 表示一个连续存在的业务组织，内部 ID 永久不变；OrganizationVersion 保存该组织某一有效期内的致远 ID、编码、名称、父级、完整路径、状态和来源指纹。",
        "SourceOrganizationBinding 按 source_instance_id + source_org_id + 有效期绑定内部组织身份；同一来源 ID 在同一时点只能有一个有效绑定，但允许不同时间绑定不同内部身份。组织编码同样按版本保存，允许变更和复用，不能单独决定连续性。",
        "单纯改名、改编码或调整上级默认保留同一内部身份并新建版本；部门撤销、来源 ID/编码复用或业务性质改变默认创建新身份；拆分、合并使用 OrganizationLineage 记录前后关系，不强行合并历史。",
        "名称、编码、来源 ID 同时变化或系统无法判断连续性时生成 OrganizationChangeReview。待复核范围在授权层 fail closed，不继承旧组织历史权限；具备 organization.change.review 权限的人员确认“同一组织延续”或“新组织”后原子发布。",
        "系统禁止仅凭名称相似度、人员重叠比例、父级相同或编码相近自动认定组织连续；这些信号只能作为复核参考，可靠 OA 变更日志或有权限人员结论才可建立跨来源 ID 的连续绑定。",
        "报表部门归属使用考勤发生日的 OrganizationVersion；正式月结冻结 organization_identity_id、organization_version_id、当时编码、名称、父级路径和版本校验值。",
        "访问权限使用当前致远组织授权快照解析出的内部组织身份及向下穿透关系。来源 ID/编码被新组织复用时，新组织负责人不能访问旧组织数据；只有被致远当前范围明确覆盖且通过本系统字段权限的角色可以查看。",
        "本系统只允许维护功能角色、敏感字段权限和例外授权；任何例外授权不得突破致远提供的组织数据范围。",
    ]:
        add_bullet(doc, item)
    org_change_rows = [
        ["无变化", "仅更新 last_seen_at 与同步批次", "不生成版本"],
        ["改名 / 改编码 / 调整上级", "同一 organization_identity_id 下关闭旧版本并创建新版本", "自动发布；保留旧名称、编码和路径"],
        ["停用 / 恢复", "新建状态版本并同步任职、授权影响", "停用立即撤权；恢复按当前绑定重新授权"],
        ["来源 ID 变化但明确为同一组织", "关闭旧绑定，新增来源绑定，内部身份不变", "需可靠变更标识或人工复核"],
        ["ID/编码复用、业务性质改变", "关闭旧绑定并创建新内部身份", "不得继承旧历史或权限"],
        ["拆分 / 合并 / 无法判断", "进入复核并记录 lineage 或连续性结论", "受影响范围默认拒绝，复核后发布"],
    ]
    add_table(doc, ["变更类型", "版本处理", "权限与发布"], org_change_rows,
              [1.65, 2.75, 2.1], font_size=8.1, first_col_bold=True)
    add_callout(doc, "“覆盖”的准确含义", "自动同步可以 upsert 当前组织投影，供页面和权限快速查询；但 OrganizationVersion、SourceOrganizationBinding、变更批次和审计记录只能追加或关闭有效期，绝不能物理覆盖。手动按钮是立即全量比对和复核入口，不是唯一的版本生成入口。", fill=LIGHT_ORANGE, label_color=ORANGE)

    add_heading(doc, "6.3 班次、工作日历与考勤规则（FR-03，P0）", 2)
    rule_rows = [
        ["班次定义", "名称、一个或多个工作/休息段、打卡窗口、宽限、是否跨日", "不在员工资料中写死时间；按版本引用"],
        ["工作日历", "工作日、休息日、法定节假日、调休工作日", "按年度版本化，允许公司级覆盖"],
        ["适用规则", "个人/地点/总部季节、班次、生效起止日、优先级", "同层级有效期不得重叠；缺少规则时进入异常"],
        ["规则版本", "迟到、早退、缺卡、缺勤、请假、外出、加班及取整", "变更须指定生效日，不静默改写已月结月份"],
        ["排班扩展", "schedule_type、跨日锚点、周期组、弹性窗口、临时分配", "P0 只保存可扩展结构；NIGHT/ROTATION/FLEX/TEMP 计算策略在 P1 启用"],
    ]
    add_table(doc, ["配置", "核心字段", "规则"], rule_rows, [1.1, 3.0, 2.4], font_size=8.8, first_col_bold=True)
    add_callout(doc, "班次匹配", "初始优先级为“个人有效期规则 > 考勤地点规则 > 总部季节规则”。员工考勤地点来自带有效期的人事档案，不按某次设备或 GPS 自动改变；命中规则及版本须在日详情中可追溯。")

    shift_rows = [
        ["总部夏令", "总部默认", "08:30-18:00", "每年 05-01 至 09-30", "午休时段待确认"],
        ["总部冬令 A", "总部默认", "08:30-17:30", "每年 01-01 至 04-30", "午休时段待确认"],
        ["总部冬令 B", "总部默认", "08:30-17:30", "每年 10-01 至 12-30", "12-31 与午休待确认"],
        ["大连", "地点规则", "07:30-12:00；13:00-16:30", "后台有效期", "午休不计工时；晚餐扣减可配置"],
        ["成都指定组", "地点/人员规则", "09:00-12:00；13:00-18:00", "后台有效期", "适用于已确认的成都指定人员"],
        ["个人总部规则", "个人覆盖", "引用总部夏/冬令", "后台有效期", "用于考勤地与执行制度不同的员工"],
    ]
    add_table(doc, ["规则", "层级", "工作时段", "有效期", "备注"], shift_rows,
              [1.15, 1.0, 1.75, 1.35, 1.25], font_size=7.9, first_col_bold=True)

    initial_assignments = [
        ["周文武、周彦沛、彭帆、唐浩", "成都", "成都指定组 09:00-12:00；13:00-18:00"],
        ["赵俊君、时晨、王颂雅", "上海", "个人覆盖：执行总部季节班次"],
        ["张静（总部销售中心）", "上海", "个人覆盖：执行总部季节班次"],
        ["大连办公员工", "大连", "地点规则：大连班次"],
    ]
    add_table(doc, ["首批配置对象", "考勤地点", "执行规则"], initial_assignments,
              [2.25, 1.1, 3.15], font_size=8.4, first_col_bold=True)
    add_body(doc, "上述人员仅作为 V1.7 首批配置输入，不得硬编码在程序中。调岗、异地办公或个人规则到期后，应按有效期恢复下一优先级班次；规则变更影响开放月份时重算，影响已月结月份时只生成差异。")

    add_heading(doc, "6.4 双源标准化与日考勤核算（FR-04，P0）", 2)
    add_body(doc, "输入：员工当日应出勤区间、得力原始打卡、已审批且有效的请假/销假、补签、加班、出差/外出单、规则版本和人工调整。工作时段统一以分钟核算；自然日类假期以连续日期区间核算；业务时区默认为 Asia/Shanghai。")
    calculation_rows = [
        ["正常出勤", "在有效窗口内取得上下班有效卡，未触发迟到、早退、缺卡或冲突。"],
        ["迟到/早退", "分别按实际有效卡与班次起止时间计算分钟数，并应用配置的宽限值。"],
        ["缺卡/缺勤", "只有一张有效卡时标明缺上班卡或缺下班卡；无有效卡且无有效 OA 覆盖时按配置判断缺勤。"],
        ["补签", "漏卡后须在“一周”时效内提交并最终审批通过；补签仅替代对应卡点，不伪造成得力原始卡。时效按可配置边界校验。"],
        ["请假", "仅审批通过且未撤回/作废的单据生效；有效时长为单据时间与应工作时段交集，休息段不计。"],
        ["销假/提前返岗", "批准假期内形成完整有效打卡区间时，日考勤先按实际证据计算出勤，并创建提前返岗事件；正式假期余额仅在该事件按当期生效策略转为 CONFIRMED 后记账。策略可选择自动确认、仅 OA 销假确认、HR 确认或混合确认；零散单卡必须人工复核。"],
        ["出差/外出", "须同时具备提前提交的出差申请和外出登记；目的地、事由、起止时间及关联单号完整后，才覆盖相交工作时段。"],
        ["跨夜下班卡", "次日凌晨卡命中可配置跨日窗口时，可作为前一考勤日下班卡并保留真实时间；同一卡不得再作为次日上班卡，冲突进入复核。"],
        ["加班", "首次提交时间须不晚于实际加班结束后 48 个自然小时，且单据最终审批有效；默认以审批区间与实际打卡佐证区间交集扣除配置休息后认定。"],
        ["超时未申报", "48 小时内未提交加班单时，最终加班为 0；原始凌晨/晚间卡仍保留，正常出勤结果按公司班次结束规则封顶，不改写原始记录。"],
        ["冲突", "不同类型 OA 单据重叠、工号无法匹配、班次缺失或时间异常时，不自动猜测，进入异常工作台。"],
        ["结果标签", "一个员工日可以同时包含迟到、请假、外出等多个标签，不强制压缩为单一状态。"],
    ]
    add_table(doc, ["结果/事件", "核算规则"], calculation_rows, [1.25, 5.25], font_size=8.8, first_col_bold=True)
    add_callout(doc, "跨夜处理", "例如员工当日早晨签到、次日凌晨签退：在跨日窗口内，凌晨卡归属前一考勤日，从而不误判缺下班卡。若没有凌晨卡，则仍是缺下班卡，可按一周时效办理补签；跨日窗口截止时刻待业务确认。")
    add_callout(doc, "证据链", "每条日结果必须可展开查看适用班次与日历、全部原始打卡、OA 首次提交时间和审批状态、规则版本、餐时扣减、实际休假/出勤区间及人工调整记录。")

    add_heading(doc, "6.5 异常工作台与人工调整（FR-05，P0）", 2)
    for item in [
        "异常类型至少包括：人员未匹配、数据格式错误、规则空档、重复/冲突、无班次、缺卡、补签超时、跨日卡归属冲突、外出手续不完整、加班超时申报、假期内打卡、单据重叠、审批状态异常、月结后变化，以及 ORG_ID_REUSE_SUSPECTED、ORG_CONTINUITY_UNKNOWN、ORG_EFFECTIVE_TIME_UNKNOWN、ORG_HIERARCHY_CYCLE、ORG_HIERARCHY_ORPHAN、ORG_SCOPE_UNRESOLVED。",
        "支持按月份、部门、工号、姓名、异常类型、状态和来源筛选，并可批量指派、备注和重算。",
        "人工调整不得修改原始打卡或 OA 单据；必须填写原因，记录调整前值、调整后值、操作人和时间。",
        "员工反馈可以关联异常，但反馈本身不等同于 OA 审批，也不能直接修改考勤结果。",
        "排行留言的举报和屏蔽属于内容治理，不得作为补签、请假或考勤申诉的正式依据。",
        "组织类异常只能由具备 organization.change.review 权限的 HR/组织数据管理员处理；普通考勤调整不能修改组织身份、来源绑定、有效期或层级快照。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "6.6 月结、反月结与版本（FR-06，P0）", 2)
    month_rows = [
        ["硬阻断", "双源同步失败或过期、人员未匹配、班次/日历缺失、核算任务失败、影响当月归属/权限的组织变更待复核，以及按策略要求确认但仍未确认的提前返岗事件", "禁止月结"],
        ["软提醒", "待审批 OA 单据、普通考勤异常、员工反馈未关闭", "主管确认并留痕后可继续"],
        ["月结动作", "取得期间锁与 close_cutoff，等待截止点前同步/重算完成，冻结输入清单、日/月结果、源水位、规则/台账版本，以及组织同步批次、身份、版本、发生时来源 ID/编码/名称/路径和校验值", "快照原子发布；失败回到原状态"],
        ["月结后变化", "close_cutoff 之后到达的迟报打卡、补单、撤单或规则变化", "统一生成差异，不覆盖快照"],
        ["反月结", "仅考勤主管，必须填写原因；原则上只允许最近一个已结月份", "旧版本永久保留，新版本递增"],
    ]
    add_table(doc, ["阶段", "检查/处理", "系统行为"], month_rows, [1.1, 3.1, 2.3], font_size=8.7, first_col_bold=True)

    add_heading(doc, "6.7 清单、汇总与看板（FR-07，P0）", 2)
    report_rows = [
        ["源数据清单", "得力原始打卡；OA 请假/销假、补签、加班、出差/外出登记；同步批次和错误"],
        ["核算清单", "员工日明细；异常清单；单据冲突；月结后差异；人工调整"],
        ["专项清单", "补签/超时补签、请假/销假、加班/超时申报、出差/外出、跨日卡、迟到、早退、缺卡、缺勤"],
        ["个人月汇总", "应出勤、实出勤、迟到/早退、缺卡、缺勤、各类假期、外出、审批/认定加班、餐扣、年假余额"],
        ["部门/公司汇总", "以上个人指标按发生时 OrganizationVersion 聚合；即使来源 ID 相同，不同内部组织身份仍分行，支持公司→部门→员工→日期下钻"],
        ["看板", "数据新鲜度、出勤率、异常趋势、部门对比、请假/加班/外出趋势；按权限控制下钻"],
    ]
    add_table(doc, ["输出", "内容"], report_rows, [1.4, 5.1], font_size=8.9, first_col_bold=True)
    add_body(doc, "所有列表、看板和导出使用同一筛选条件、同一指标口径和同一数据版本。导出文件应包含生成时间、操作者、数据范围、月份状态和月结版本。")
    add_callout(doc, "组织统计口径", "P0 正式报表仅使用“业务发生时组织”口径，跨组织版本可连续展示但不得覆盖当时编码、名称和路径。P1 如增加“按当前组织重述”，必须明确标记为非正式分析口径，单独生成结果且不得改写月结快照。", fill=LIGHTER_BLUE, label_color=BLUE)

    add_heading(doc, "6.8 多角色、三端与员工前端（FR-08，P0）", 2)
    frontend_rows = [
        ["角色工作台", "员工、考勤、HR、管理者、系统管理员各自展示最相关的状态、待办、风险和入口；页面显隐不替代服务端授权"],
        ["首页", "今日状态、本月摘要、待处理异常/反馈、年假余额、当日签到排行入口、数据截至时间"],
        ["我的考勤", "月历、每日详情、有效/原始打卡、OA 单据、计算解释、反馈入口"],
        ["数据看板", "个人看板；授权用户可切换公司或部门看板并按权限下钻"],
        ["签到排行/留言", "按考勤日和发布范围查看名次、展示名及公开状态；发布、回复/举报或删除本人留言"],
        ["我的假期", "年度额度、结转、调整、已用、剩余、到期日和变动流水"],
        ["制度中心", "分类、搜索、有效版本、正文/附件、发布时间和适用范围"],
        ["反馈中心", "新建反馈、关联考勤日期、双方回复、处理进度和最终结论"],
        ["管理驾驶舱", "公司/授权组织的 KPI、趋势、组织对比、异常结构、新鲜度与月结状态；小样本保护，不提供个人敏感下钻"],
    ]
    add_table(doc, ["页面", "核心内容"], frontend_rows, [1.25, 5.25], first_col_bold=True)

    add_heading(doc, "6.9 假期、特殊工时、制度与反馈（FR-09，P0）", 2)
    add_body(doc, "V1.7 以假别规则、员工适用班次和 OA 有效单据共同核算。最小申请单位用于校验，不满足时提示修改或进入异常，不得静默向上取整；工作日假期只计算与应工作时段的交集。假期规则、发放与返岗结算均采用可配置的受控策略和不可变台账，不把“5 天”等当前制度值硬编码在员工表。")
    leave_rows = [
        ["年假", "入职满 1 年后享有 5 天/年；最小 0.5 天", "0.5 天按当日所选实际半日时段计算，不统一换算为 4 小时；发放周期、折算、结转和失效均按版本化策略配置"],
        ["丧假", "配偶、子女、父母 3 天；兄弟姐妹、祖父母、外祖父母 1 天", "申请页以亲属关系单选项校验上限；自然日/工作日及分段规则待确认"],
        ["护理假", "15 个自然日", "连续包含周末和节假日；适用对象、起算事件及是否必须连续待确认"],
        ["病假", "最小 0.5 天", "按员工当日实际半日时段折算；是否必须按 0.5 天递增待确认"],
        ["事假", "最小 0.5 小时", "以 30 分钟为最小申请值；是否同时作为递增单位待确认"],
        ["哺乳/孕检", "按员工配置独立免考勤时段", "在基础班次上叠加生效日、星期和一个或多个时段；与请假/外出重叠不重复扣减"],
    ]
    add_table(doc, ["假别/规则", "当前业务口径", "系统处理"], leave_rows,
              [1.2, 2.0, 3.3], font_size=8.2, first_col_bold=True)
    leave_strategy_rows = [
        ["额度来源", "AUTO_RULE / IMPORT / MANUAL", "自动规则、HR 导入或人工调整均形成台账分录，禁止直接改余额"],
        ["发放策略", "自然年 / 入职周年 / 自定义计划", "配置适用人群、司龄门槛、额度公式、按比例折算、发放日和生效区间"],
        ["结转与失效", "结转上限、到期日、使用顺序、余额不足处理", "按规则版本执行；发布后不可原地修改，变更创建新版本"],
        ["申请口径", "分钟 / 半天 / 工作日 / 自然日", "配置最小单位、递增单位、日历基准及显示单位"],
        ["提前返岗", "AUTO_FULL_PUNCH / OA_CANCEL_ONLY / HR_CONFIRM / HYBRID", "完整有效打卡、OA 销假和 HR 决定按模式驱动确认；零散单卡始终复核"],
        ["作用范围", "公司 / 假别 / 员工组 / 生效期", "同一范围同一时段只能有一个已发布版本；已月结结果不回写"],
    ]
    add_table(doc, ["策略维度", "可配置选项", "约束"], leave_strategy_rows,
              [1.25, 2.15, 3.1], font_size=7.9, first_col_bold=True)
    add_body(doc, "假别目录本身也可由授权 HR 配置：编码、名称、工作日/自然日口径、最小与递增单位、必填字段、亲属关系选项集、敏感级别、适用范围和生效期。已经被 OA 单据或台账引用的假别/选项只能失效，不能物理删除；编码一经发布不可复用。")
    add_callout(doc, "配置发布流程", "假期策略仅允许从受控参数和策略枚举中选择，不开放任意代码、表达式、SQL 或脚本。版本状态为 DRAFT → VALIDATED → PUBLISHED → EXPIRED；发布前必须完成冲突校验、指定员工/月份预览和金标准回归。已发布版本不可编辑，变更以新版本及生效日发布，不重写已月结月份。", fill=LIGHT_ORANGE, label_color=ORANGE)
    add_body(doc, "工作日类假期台账以分钟作为内部核算单位，同时保留来源单位、展示单位和转换规则版本；“天/半天”只按员工当日班次换算展示。护理假等自然日类假期以连续日期区间记录，不与工作分钟直接相加。")
    add_body(doc, "年假余额口径：已确认发放/导入额度（当前制度基线为入职满 1 年 5 天）+ 结转 + 人工调整 - 最终认定使用量 + 已确认的销假/撤回返还。批准假期内形成完整有效打卡区间时，出勤结果可先按事实计算；返还额度必须生成可追溯台账分录，并由当期 AUTO_FULL_PUNCH、OA_CANCEL_ONLY、HR_CONFIRM 或 HYBRID 策略决定确认方式。年假期间提前出差时，以有效出差区间或 OA 销假/变更确定实际休假终点，重叠时段不得同时计为出勤/出差和年假。")
    add_callout(doc, "半天换算示例", "大连上午 07:30-12:00 为 4.5 小时，下午 13:00-16:30 为 3.5 小时，因此“0.5 天”应记录所选半日实际时段，而不是固定扣 4 小时。总部午休未确认前，不能冻结总部半天假分钟数。")
    feedback_states = [
        ["待处理", "员工提交成功，系统带入关联日期、工号和异常类型"],
        ["处理中", "考勤员受理、回复或要求补充材料"],
        ["已解决", "给出结论；如需调整，关联独立人工调整记录"],
        ["已驳回", "说明理由；员工可新建反馈或按权限重新打开"],
        ["已关闭", "双方确认或超过配置期限关闭，历史回复只读"],
    ]
    add_table(doc, ["反馈状态", "规则"], feedback_states, [1.25, 5.25], first_col_bold=True)
    add_body(doc, "制度管理支持草稿、发布、失效和版本管理。已发布版本不得直接覆盖；新版本需设置生效时间和适用范围。员工只查看当前对其有效的版本。")

    add_heading(doc, "6.10 当日签到排行与留言（FR-10，P0）", 2)
    ranking_rows = [
        ["排行分组", "默认按同一考勤日 + 同一考勤地点 + 同一班次生成；公司/部门/地点范围可配置，禁止不同班次按绝对时刻直接混排。"],
        ["排名依据", "取当天第一条有效现场上班卡；补签、人工调整、外出卡默认不参与；迟报原始卡可更新榜单并显示数据截至时间。"],
        ["公开字段", "仅发布名次、展示名及经确认的时间粒度；不发布未签到/请假名单、原始卡、OA 单据、设备、坐标或敏感假别。"],
        ["留言", "员工在所属榜单发布纯文本留言；支持回复、删除本人内容、举报；HR 可屏蔽并填写原因，操作均留痕。"],
        ["业务边界", "排行与留言是员工互动，不是补签、请假或申诉凭证；后台可关闭，内容访问范围与榜单范围一致。"],
    ]
    add_table(doc, ["能力", "规则"], ranking_rows, [1.25, 5.25], font_size=8.7, first_col_bold=True)
    add_callout(doc, "隐私边界", "员工“仅看本人考勤明细”与“查看排行”不冲突：排行是单独发布的最小公开字段集，接口不得返回他人的精确签到时间（除非业务明确授权）、原始记录、位置、设备或请假原因。", fill=LIGHT_ORANGE, label_color=ORANGE)

    add_heading(doc, "6.11 定位与地图查看（FR-11，P0）", 2)
    add_callout(doc, "产品边界", "定位查看仅用于核验一次具体 GPS/外勤打卡地点，不提供实时定位、持续后台定位、轨迹拼接、批量人员地图或轨迹导出。普通人脸、指纹、刷卡等没有有效经纬度的记录仅展示已有设备或地点文字，不提供地图入口。", fill=LIGHT_RED, label_color=RED)
    for item in [
        "仅当得力 check_type 为 gps 或 out_work，且 check_data.lat 与 check_data.lgt 可解析且纬度在 [-90,90]、经度在 [-180,180] 内时，详情显示“查看地图”；字段名按官方接口为 lgt，不得误写为 lng。",
        "入库时规范化为 latitude、longitude、location_name、location_address、source_coordinate_system 和 coordinate_status；缺失或非法坐标不展示地图入口，有 name/location 时仍可在通过同一位置权限后展示地点文字。",
        "员工可查看本人该次打卡地图；管理者在服务端必须同时通过当前致远组织穿透范围和 attendance.location.map 独立权限。查看时二次点击，管理者需选择/填写原因。",
        "点击后才调用地图适配器，在弹层/抽屉显示单点标记、地点名称/地址、打卡时间与方式；默认不直接显示数字坐标。地图服务失败时保留地点文字并提示“地图暂不可用”。",
        "得力文档未声明坐标系；联调必须用真实样本确认 WGS84、GCJ-02 或 BD-09。坐标系未确认时不得在生产地图上绘制未经验证的点位，只向已通过位置权限的用户展示来源地点文字。",
        "地图供应商、应用 Key、坐标转换方式和版本均后台配置，密钥不下发日志；只允许受控供应商适配器，不开放任意地图 URL 或脚本。坐标只在用户点击时发送给地图服务。",
        "系统记录查看人、对象、时间、原因、来源记录、地图供应商、坐标系和成功/失败结果；公司和部门看板及普通导出不得包含个人位置或精确坐标。",
        "首版不保存外勤照片、手机型号、Wi-Fi MAC/SSID；规范化坐标保存在业务数据库的敏感列，最小化来源 JSON 不重复保存坐标。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "6.12 工资配置、分段核算与工资条（FR-12，P0）", 2)
    add_callout(doc, "核心口径", "员工工资明细与部门/成本中心承担明细是两个对象。工资分段决定某段采用哪个组织、岗位和工资规则；成本分摊决定金额由哪个责任组织、成本中心或项目承担。两者默认相同，也可依据已配置的分摊规则不同；任何一方都不得反向授予工资查看权限。", fill=LIGHT_BLUE, label_color=BLUE)
    add_body(doc, "输入包括：工资期间、员工有效任职、OA 人事调动/调薪、OA 薪资字段及附件、HR 受控 Excel、本系统工资档案、发生时组织与岗位版本、工资方案/项目/标准/公式版本、考勤月结版本、其他工资输入和人工补发追扣。所有外部输入先暂存、校验、去重和解决冲突，再发布为本系统有效版本；禁止直接对 OA 附件或临时 Excel 计算。输出包括工资分段、项目明细、社保公积金个人/单位行、个税计算行、计算轨迹、异常、差异、复核发布记录、工资条版本和成本分摊。")
    payroll_steps = [
        "创建工资期间并选择考勤月结版本；从 OA 字段/关联附件、HR Excel 与本系统收集工资资料，保存来源单据、附件/文件、Sheet/行号、内容哈希、导入批次和字段映射版本。",
        "在暂存区完成模板、员工、格式、有效期、重复和跨来源冲突校验；未命中已发布来源权威策略的冲突由授权薪酬人员选择有效候选并留痕，未解决时不得锁定。",
        "锁定输入快照后，合并期间起止、入离职、任职、组织、岗位/职级、调动/调薪、工资方案和个人定薪的全部有效期边界，生成互不重叠的半开区间 [valid_from, valid_to)。",
        "每个工资分段必须唯一命中一个任职版本、内部组织/岗位身份与版本、工资方案、员工定薪和折算策略；零个或多个匹配均进入硬阻断异常，不自动猜测。",
        "工资引擎按项目策略选择实际计薪工作日或实际计薪工时，生成部门工资、岗位工资、基本工资、考勤收入/扣减和其他项目行；每行保存数量、分母、单价、基数、公式/取整版本、来源证据和结果。",
        "按员工参保/缴存地、基数和生效政策计算社保与公积金个人承担额、单位承担额；个人承担进入员工扣减，单位承担只进入企业成本。",
        "按中国居民个人累计预扣计算框架和已冻结年度累计台账计算个税，保存本期/累计收入、扣除、已预扣、本期税额及政策版本；正式参数与期初累计数据在政策契约中冻结。",
        "系统另行生成成本分摊行，校验责任组织/成本中心/项目比例和金额勾稽；分摊不会改变员工应发金额，也不赋予责任部门负责人个人工资权限。",
        "薪酬制表员处理异常并查看本期/上期、变更前/后的差异；提交后由不同账号的复核发布人只读核对并退回或发布。没有额外业务审批状态，制表人不能复核发布本人批次。",
        "发布后生成不可变正式版本；员工通过当前登录会话查看本人已发布工资条，工资条展示月内部门/岗位分段、考勤工资项、社保公积金、个税、其他项目和实发合计。",
        "关闭后到达的 OA 文件、Excel、调动、调薪、考勤或政策变化只生成 PayrollPostCloseDifference；正式版本不反向覆盖，只能在新运行中补发、追扣或冲销。",
    ]
    for item in payroll_steps:
        add_number(doc, item)

    component_rows = [
        ["分段固定项", "基本工资、部门工资、岗位工资、随组织/岗位变化的固定津贴", "按工资分段和配置折算；分别展示，不合并成不可解释金额"],
        ["考勤事件项", "加班工资、事假/缺勤/迟到早退扣款、餐扣等", "按考勤事实发生日归入当天有效工资分段；引用考勤月结结果和规则版本"],
        ["月度公共项", "月度奖金、固定补贴、一次性增减", "按项目策略仅计算一次；录入主体可配置为薪酬 HR、授权部门或两者，需要分摊时另生成分摊行"],
        ["法定扣缴项", "社保、公积金、个人所得税", "P0 必算；个人与单位承担分开，个税保留年度累计计算轨迹及政策版本"],
        ["补发追扣项", "历史调动/调薪/考勤变化的差额", "引用原工资期间、原分段、原项目和原责任组织；正负金额与原因可追溯"],
        ["仅分摊项", "不改变员工实发、仅用于成本责任的调整", "只进入 PayrollCostAllocationLine，不进入工资条应发合计"],
    ]
    add_table(doc, ["项目类别", "典型项目", "归属规则"], component_rows,
              [1.35, 2.2, 2.95], font_size=8.35, first_col_bold=True)

    proration_rows = [
        ["按实际计薪工作日", "分段实际计薪工作日 / 期间实际计薪工作日", "P0；数量来源与缺勤扣款去重规则须在项目策略冻结"],
        ["按实际计薪工时", "分段实际计薪工时 / 期间实际计薪工时", "P0；适配跨地区、半日或工时计薪，引用考勤月结/班次日历版本"],
        ["不折算", "按月一次性全额或条件金额", "仅用于明确不随分段折算的奖金、补贴等项目"],
    ]
    add_table(doc, ["折算策略", "计算基线", "约束"], proration_rows,
              [1.35, 2.45, 2.7], font_size=8.35, first_col_bold=True)
    add_callout(doc, "已确认折算边界", "P0 同时支持按实际计薪工作日和按实际计薪工时，具体策略由每个工资项目绑定，不使用全局统一算法。工作日/工时的事实来源、缺勤是否已在数量中扣除及防止重复扣款的规则仍须在工资规则确认单冻结；中间值保留高精度，按版本化阶段统一取整到分。", fill=LIGHT_ORANGE, label_color=ORANGE)

    transfer_example = [
        ["7月1日-15日", "销售部 / 销售岗位", "销售部工资标准 + 销售岗位标准", "归入 1-15 日的加班、请假、缺勤等", "该段适用的其他项目", "销售责任组织/成本中心"],
        ["7月16日-31日", "采购部 / 采购岗位", "采购部工资标准 + 采购岗位标准", "归入 16-31 日的加班、请假、缺勤等", "该段适用的其他项目", "采购责任组织/成本中心"],
        ["月度公共项", "不强制拆段", "—", "—", "月度奖金、补贴及其他配置项", "按项目分摊策略"],
    ]
    add_table(doc, ["工资分段", "发生时组织/岗位", "固定工资来源", "考勤工资项", "其他项目", "成本承担"],
              transfer_example, [1.0, 1.35, 1.45, 1.25, 1.25, 1.2], font_size=7.45,
              first_col_bold=True)
    add_body(doc, "调动生效日已确认归新部门和新岗位：旧分段为 [期间开始, 生效日 00:00)，新分段从生效日 00:00 开始；当日全部计薪工作量进入新分段。若 OA 明确提供并要求采用精确业务生效时刻，则按来源时刻切段。部门变化但工资标准相同也保留两个分段，以保证历史展示和成本归属可解释。")

    payroll_exception_rows = [
        ["任职/调动缺失", "OA 调动单缺生效时间、审批状态、新旧部门/岗位，或任职区间重叠/空档", "禁止计算/关闭"],
        ["工资配置冲突", "分段无法唯一命中方案、工资标准、定薪或折算规则；公式循环、除零或未发布", "禁止计算/关闭"],
        ["考勤输入不稳定", "工资选择的考勤月份未月结，或锁定后考勤快照变化", "禁止正式关闭；产生输入变化提示"],
        ["来源/输入冲突", "OA 字段、OA 附件、HR Excel 或本系统同一业务键多值、重复导入、模板不识别或文件哈希变化", "隔离异常；按来源策略或人工处置后发布有效版本"],
        ["法定计算缺失", "参保/缴存地、基数、费率上下限、个税规则或年度累计台账缺失/多匹配", "禁止关闭与发布，不得以 0 伪装已计算"],
        ["成本不勾稽", "分摊比例不等于 100%，或分摊金额与可分摊工资金额不一致", "禁止关闭与发布"],
        ["权限/授权过期", "PAYROLL 策略未发布、授权快照过期、组织/岗位绑定待复核", "薪资访问与批次操作 fail closed"],
    ]
    add_table(doc, ["异常", "判断", "处理"], payroll_exception_rows,
              [1.3, 3.55, 1.65], font_size=8.1, first_col_bold=True)
    add_callout(doc, "北森参考边界", "参考北森官方公开的人员异动/考勤联动、分段计薪、类 Excel 多 Sheet 核对、差异提示和 PC/移动工资条的信息组织思路；本系统不设置业务审批与 MFA，采用自己的数据模型、复核发布和安全补偿控制，不推断或复制北森未公开的算法和底层实现。", fill=LIGHTER_BLUE, label_color=BLUE)

    build_prd_body_part2(doc)


def build_prd_body_part2(doc: Document) -> None:
    add_heading(doc, "7. 核算与报表口径", 1)
    add_heading(doc, "7.1 数据优先关系", 2)
    priority_rows = [
        ["班次/日历", "定义员工本来应该怎样出勤", "应出勤基线"],
        ["得力打卡", "记录实际发生的打卡事实", "原始事实，不可人工覆盖"],
        ["致远 OA 单据", "定义审批通过的请假/销假、补签、加班、出差/外出例外", "保留首次提交时间和状态；仅符合时效且最终有效的单据参与核算"],
        ["致远人事异动", "定义组织、岗位、职级、调动/调薪及业务生效时间", "只使用最终有效版本切分工资段；撤回/修改触发重算或工资差异"],
        ["人工调整", "记录考勤员最终复核决定", "独立记录，必须有原因和审计"],
        ["月结快照", "冻结正式结果版本", "后续变化进入差异或反月结版本"],
        ["工资来源与输入", "OA 结构化字段/附件、HR Excel、本系统工资档案与其他输入", "先暂存校验并按来源策略解决冲突，发布为有效期版本后参与快照；不得按最后更新时间静默覆盖"],
        ["工资配置", "工资方案、项目、标准、员工定薪、公式、税社保公积金政策", "均按有效期发布版本参与工资输入快照；不得原地覆盖"],
        ["工资关闭版本", "冻结分段、项目、工资条和成本分摊", "后续变化进入补发追扣或经授权反结的新版本"],
    ]
    add_table(doc, ["数据层", "含义", "优先原则"], priority_rows, [1.3, 2.8, 2.4], first_col_bold=True)
    source_truth_rows = [
        ["得力云", "原始打卡、设备上报事实", "本系统只读同步并保存不可变原文；不得用人工调整回写或覆盖来源事实"],
        ["致远 OA", "员工、组织、岗位、任职、调动/调薪、账号状态、组织数据范围、考勤申请，以及现有薪资记录和关联附件", "人员、组织、岗位与异动的权威来源；薪资字段/附件是工资候选来源之一。本系统只读同步并保存副本、哈希和历史版本"],
        ["HR 受控 Excel", "OA 尚未结构化或由 HR 维护的工资档案、法定计算资料和其他工资输入", "过渡期生产入口；按模板、Sheet、行号、文件哈希和导入批次留痕，校验发布后方可参与计算"],
        ["本系统授权", "功能角色、敏感字段权限和例外授权", "最终权限为本地授权与当前致远组织范围交集；本地授权不得扩大致远范围"],
        ["本系统", "标准化映射、班次/工资规则、考勤与工资结果、异常、调整、假期台账、双月结快照、工资条和发布投影", "只对本系统派生事实负责；不向得力或致远写回"],
    ]
    add_table(doc, ["权威来源", "负责的事实", "系统边界"], source_truth_rows,
              [1.25, 2.15, 3.1], font_size=8.5, first_col_bold=True)
    add_callout(doc, "边界原则", "外部源事实、本系统规范化有效版本、派生结果和人工冲突处置/复核发布结论必须分表、分版本保存。OA 与 HR Excel 同一字段冲突时先执行已发布来源权威策略；仍无法唯一判断则硬阻断并由薪酬人员留痕选择，禁止最后写入者覆盖。", fill=LIGHT_BLUE, label_color=BLUE)

    add_heading(doc, "7.2 指标定义", 2)
    metrics = [
        ["应出勤分钟", "适用班次的工作区间分钟数，扣除休息段；不因请假而改变基线。"],
        ["请假/外出有效分钟", "最终有效单据时间与应工作区间的交集分钟数；被确认的提前出勤或出差覆盖时段不重复计算。"],
        ["实出勤分钟", "依据有效打卡和配置规则认定的工作分钟，不含已认定加班。"],
        ["有效应出勤分钟", "应出勤分钟 - 已批准请假分钟 - 按规则免打卡的外出分钟。"],
        ["出勤率", "实出勤分钟 / 有效应出勤分钟；分母为 0 时显示“不适用”，不显示 0%。"],
        ["迟到/早退分钟", "有效卡时间与班次时间比较，并应用宽限值后的正分钟数。"],
        ["加班申报有效", "首次提交时间 <= 实际加班结束时间 + 48 个自然小时，且最终审批通过、未撤回/作废。"],
        ["最终加班分钟", "申报有效时，默认取审批区间与实际打卡佐证区间交集，再扣除按地点/单次配置的休息；否则为 0。"],
        ["正式剩余年假", "已确认发放/导入额度 + 结转 + 人工调整 - 最终认定使用量 + 已确认销假/撤回返还；所有变化来自不可变台账分录。"],
    ]
    add_table(doc, ["指标", "定义"], metrics, [1.6, 4.9], first_col_bold=True)
    add_body(doc, "部门和公司汇总必须由员工日明细聚合，不能直接累加原始 OA 时长或原始打卡次数，以免重叠单据和重复卡导致口径不一致。")

    doc.add_page_break()
    add_heading(doc, "7.3 默认打卡与冲突规则", 2)
    default_rules = [
        ["多次打卡", "全部原始记录保留；固定班次在有效窗口内取最早上班卡、最晚下班卡作为默认有效卡。"],
        ["重复时间打卡", "来源记录不删除；核算层在配置的去重窗口内只选择一张有效卡。"],
        ["只有一次打卡", "先检查次日凌晨跨日窗口内是否存在可归属的下班卡；有则归入前一考勤日且不得复用，无则按缺卡并进入补签流程。"],
        ["补签时效", "业务时效值为一周；当前默认按 7 个自然日参数化，起算点、截止边界及超期特批待确认。审批后重算，月结后只产生差异。"],
        ["外出手续", "出差申请与外出登记须关联且有效；目的地、事由、起止时间缺失或区间不一致时进入“手续不完整”异常。"],
        ["跨夜加班", "凌晨下班卡保留真实时间并归属加班开始日；首次提交在 48 小时内且最终获批后重算，否则加班为 0。"],
        ["加班餐扣", "不得统一固定扣 0.5 小时；按单次人工确认 > 个人 > 地点 > 公司默认规则选择，大连晚间加班允许不扣。"],
        ["销假/提前出勤", "完整有效打卡覆盖的工作时段可先按实际出勤计算，并产生 DETECTED 提前返岗事件；AUTO_FULL_PUNCH 可自动确认，OA_CANCEL_ONLY 等待 OA 销假，HR_CONFIRM 等待 HR，HYBRID 按条件组合。只有 CONFIRMED 才写入返还台账，零散单卡始终复核。"],
        ["同类 OA 单据重叠", "标准化层合并相交时间段，不重复计算。"],
        ["不同 OA 类型重叠", "不得自动确定优先级，进入单据冲突异常，由考勤员复核。"],
        ["OA 单据撤销/作废", "开放月份自动冲销并重算；已月结月份生成差异。"],
        ["设备迟报", "按 check_time 所属考勤日重算；月结后不覆盖快照。"],
    ]
    add_table(doc, ["场景", "默认规则"], default_rules, [1.55, 4.95], font_size=8.9, first_col_bold=True)

    add_heading(doc, "7.4 工资核算与成本口径", 2)
    payroll_metric_rows = [
        ["工资分段", "同一员工在工资期间内、任职/组织/岗位/定薪/工资方案均稳定的最小半开区间；各分段无重叠且完整覆盖应计薪有效期。"],
        ["分段固定工资", "该段适用固定项目标准 × 已发布折算规则；部门工资和岗位工资分别计算、分别展示。"],
        ["考勤工资项", "考勤月结事件按实际发生日归入当天有效工资段；工资行引用考勤结果版本，不能重新解释原始打卡。"],
        ["月度其他项目", "奖金、补贴、一次性收入/扣减按项目策略计算；录入主体可配置为薪酬 HR、授权部门或两者，来源、凭证、发生日和版本必须可追溯。"],
        ["社保与公积金", "按员工参保/缴存地、基数、上下限、个人/单位比例和生效政策计算；个人额扣减实发，单位额仅计企业成本。"],
        ["个人所得税", "按已发布税务政策与员工年度累计台账计算；保存本期/累计收入、扣除、已预扣和本期税额，不得把缺数据当作 0。"],
        ["员工实发工资", "应发收入 - 员工承担社保公积金 - 个税 - 其他员工扣减 + 当期补发追扣；单位承担额不减少员工实发。"],
        ["成本分摊", "独立分摊行按责任组织/成本中心/项目归集；工资成本与单位承担社保公积金分别勾稽，不改变员工实发且不授予查看权。"],
        ["取整与勾稽", "中间值保留 Decimal 高精度；按项目配置的阶段和规则统一取整到分。分段合计、工资条合计、批次合计和成本分摊合计差异均为 0。"],
        ["调动生效日", "已确认生效日 00:00 起归新部门/岗位，当日计薪工作量进入新分段；OA 若有明确业务精确时刻则按来源时刻。"],
    ]
    add_table(doc, ["口径", "定义"], payroll_metric_rows, [1.6, 4.9], font_size=8.55, first_col_bold=True)
    add_callout(doc, "可解释性", "任一工资金额必须能展开到工资期间、员工分段、发生时组织/岗位、工资方案、项目、实际工作日/工时、单价/基数、公式与取整版本、考勤证据、来源文件/输入记录、社保公积金与个税政策版本、复核发布记录和成本分摊；页面不得只显示一个无法解释的月度总额。", fill=LIGHT_BLUE, label_color=BLUE)

    add_heading(doc, "7.5 核心领域模型与不变量", 2)
    domain_rows = [
        ["人员与组织", "Employee、EmployeeSourceBinding、OrganizationIdentity、OaOrganizationSourceVersion、OrganizationVersion、SourceOrganizationBinding、OrganizationHierarchySnapshot、OrganizationCurrentProjection、OrganizationSyncBatch、OrganizationLineage、OrganizationChangeReview、EmploymentAssignmentVersion、EmployeeNumberMapping、OaAuthorizationSnapshotItem", "致远为人员组织权威源；得力 ext_id 经显式来源绑定；内部身份稳定，来源 ID/编码/名称可变；历史版本与当前投影分离，复用不得继承历史或权限"],
        ["岗位与异动", "PositionIdentity、PositionVersion、SourcePositionBinding、EmploymentTransferDocumentVersion、CompensationChangeDocumentVersion", "岗位与调动按来源和有效期版本化；来源岗位 ID/编码复用不得污染工资段，只有已批准有效异动参与切段"],
        ["班次与规则", "ShiftDefinition、ShiftSegment、ShiftAssignment、ScheduleCycle、AttendanceDayAnchorStrategy、WorkCalendar、RuleVersion、LeavePolicyVersion", "规则不可原地覆盖；P0 仅启用 FIXED_ADMIN，但结构不假设同自然日或永久单班"],
        ["来源事实", "SourceInstance、OAQueryMappingProfile、PunchRecord、OADocumentVersion、DocumentStatusHistory、SyncRun、SyncPage", "最小化后的来源事实不可变；OA 逻辑业务键由映射版本定义，内容变化由本系统形成新版本"],
        ["薪资来源", "PayrollSourceBatch、PayrollSourceRecord、PayrollSourceAttachmentVersion、PayrollFieldMappingVersion、PayrollSourceAuthorityPolicyVersion、PayrollSourceConflict", "OA 字段/附件与 HR Excel 先暂存、校验、冲突处理后发布；原文件引用、哈希、Sheet/行号和处置证据不可丢失"],
        ["位置与地图", "PunchLocation、MapProviderConfig、LocationAccessAudit", "地点文字与坐标同为敏感信息；供应商/转换版本化；查看审计不得记录坐标值"],
        ["薪资配置", "PaySchemeVersion、PayComponentDefinitionVersion、PayRuleVersion、ProrationPolicyVersion、RoundingPolicyVersion、OrgPositionPayRateVersion、EmployeeCompensationAssignmentVersion、PayrollWorkMeasureSnapshot", "方案、项目、录入主体、工资标准、工作日/工时折算和取整均按有效期版本化；发布后不可原地覆盖"],
        ["法定扣缴", "SocialInsurancePolicyVersion、HousingFundPolicyVersion、EmployeeContributionProfileVersion、ContributionCalculationLine、TaxPolicyVersion、EmployeeTaxProfileVersion、TaxYtdLedgerSnapshot、TaxCalculationLine", "员工/单位承担分离；政策、基数和年度累计均版本化，关闭后不可重写，缺少必需输入不得以 0 代替"],
        ["薪资核算", "PayrollPeriod、PayrollRun、PayrollInputSnapshot、PayrollSegment、PayrollComponentLine、PayrollException、PayrollAdjustment、PayrollPostCloseDifference", "一次运行冻结输入；工资段互斥完整；相同输入和规则结果确定；关闭后只追加差异、补发追扣或新版本"],
        ["工资条与成本", "PayslipVersion、PayrollReviewRecord、PayrollCostAllocationLine", "已发布工资条不可删除；成本分摊与员工工资分离但必须勾稽；制表与独立复核发布职责分离"],
        ["授权与敏感域", "DataDomain、PayrollAuthorizationGrant、PayrollAuthorizationSnapshot、AuthorizationPolicyVersion、FieldPolicy、AccessDecisionLog", "PAYROLL 正式独立域；OA 范围只限制不授予；考勤、位置和薪资域不互相继承权限"],
        ["核算与异常", "AttendanceDayRevision、EvidenceLink、OvertimeResult、ExceptionCase、Adjustment", "一次核算固定输入指纹；过期任务不得更新当前版本；调整以追加/冲销表达"],
        ["假期台账", "LeavePolicyVersion、LeaveAccount、LeaveLedgerEntry、EarlyReturnCase", "发放与返岗使用受控策略；余额由不可变分录汇总，纠错使用冲销和新分录"],
        ["月结与差异", "MonthClose、MonthCloseSnapshot、PostCloseDifference", "已月结快照不可覆盖；反月结产生新版本，并保留旧版本和触发原因"],
        ["发布与互动", "RankingEntry、Comment、ModerationAction", "排行使用独立最小发布投影；评论和治理记录不得反向改变考勤事实"],
    ]
    add_table(doc, ["领域", "核心对象", "架构不变量"], domain_rows,
              [1.25, 2.7, 2.55], font_size=7.9, first_col_bold=True)
    for item in [
        "所有核心记录包含 company_id；每个外部连接另有 source_instance_id，凭据、组织或环境变化时不得复用来源唯一键空间。",
        "来源唯一键：得力为 source_instance_id + source_record_id；OA 由对应 OAQueryMappingProfile 定义可重复生成的逻辑业务键，可由单字段或复合字段组成。来源内容指纹变化时由本系统创建版本，不要求 OA 自建表提供固定 detail_id 或 source_version。",
        "组织来源定位键为 source_instance_id + source_org_id + 有效区间；source_org_id、org_code 和 org_name 都是可变、可复用的来源属性，不能代替 organization_identity_id。任何时点的有效绑定不得重叠。",
        "OrganizationVersion 为不可变事实，至少保存来源 ID、编码、名称、父级内部身份、完整路径、状态、valid_from、valid_to、source_updated_at、first_seen_at、last_seen_at、fingerprint 和 sync_batch_id；当前页面读取可重建的 OrganizationCurrentProjection。",
        "组织字段指纹没有变化时不得生成新版本；发生变化时必须先关闭旧版本再创建新版本。自动同步和手动同步遵循同一差异算法、事务边界与幂等键，手动按钮不得绕开校验或直接改写正式表。",
        "同一任职在组织身份或有效绑定切换点必须拆成两个 EmploymentAssignmentVersion，即使致远 assignment_id 未变化；开放月份重算旧/新区间并集，已月结月份只生成差异。",
        "工资期间将任职、组织、岗位、调动、调薪、定薪和工资配置的全部有效期边界取并集生成 PayrollSegment；任意应计薪时点必须且只能解析到一个有效分段，空档或重叠为硬阻断。",
        "工资项目行必须记录 earning_organization_identity_id 与 responsibility_organization_identity_id/cost_center_id；前者决定计算和历史展示，后者决定成本承担。两者可不同但都需有效分摊规则、来源凭证和复核记录。",
        "PayrollSourceBatch 以 OA_FIELD、OA_ATTACHMENT、HR_EXCEL 或 SYSTEM 标识来源；同一业务键多值不得按更新时间静默覆盖。未命中来源权威策略且未解决的 PayrollSourceConflict 必须阻断输入锁定。",
        "PayrollInputSnapshot 冻结 OA 异动/任职/岗位、OA 文件/字段、HR Excel 批次、工资配置、其他输入、法定政策、年度累计台账、考勤月结版本及哈希。相同输入快照和规则版本重复计算必须得到完全相同的分段、金额和舍入结果。",
        "调动生效日从 00:00 起归新部门/岗位，旧段为 [期间开始, 生效日)，新段为 [生效日, 期间结束)。每个项目独立绑定 ACTUAL_PAYABLE_WORKDAY 或 ACTUAL_PAYABLE_HOUR，PayrollWorkMeasureSnapshot 保存数量、分母及日历/排班/考勤版本并防止与缺勤项目重复扣减。",
        "社保、公积金分别保存个人与单位计算行；个人额进入员工扣减，单位额只进入企业成本。个税使用冻结的 TaxYtdLedgerSnapshot，已关闭期间不得回写累计值。",
        "PayrollComponentLine、PayrollAdjustment、PayrollCostAllocationLine 和已发布 PayslipVersion 均不可原地覆盖；纠错使用冲销和新记录，关闭后变化进入 PayrollPostCloseDifference。",
        "管理、财务和汇总角色访问历史工资时，按当前有效 PayrollAuthorizationGrant 与记录发生时组织/成本中心共同判断；员工本人按本人对象关系查看本人已发布历史工资条，不受后续组织变更影响。当前部门、来源组织 ID/编码复用、成本归属或 OA 穿透关系均不能自动授予他人工资权限。",
        "核算修订键至少由 company_id、employee_id、attendance_day、revision_no 组成；同一输入指纹重算不得生成不同正式结果。",
        "每个日结果记录 source_watermark、rule_version_id、organization_identity_id、organization_version_id、adjustment_version 和 evidence_hash；任务提交前必须确认输入指纹仍为最新，否则仅保留运行日志。",
        "历史报表按考勤发生日的 OrganizationVersion 归属；当前页面、接口和导出把最新致远授权节点解析到当前有效 organization_identity_id 后校验。同步失败或绑定待复核不得扩大权限，后续成功同步必须及时撤销离职、停用、移出范围或旧身份的访问。",
        "schedule_type 预留 FIXED_ADMIN、NIGHT、ROTATION、FLEX、TEMP；P0 仅允许发布 FIXED_ADMIN，其他类型返回“首版未启用”，但不得以自然日或单一班段作为数据库不变量。",
        "时间事实同时保存 source_time_raw、occurred_at、received_at、local_datetime 和 attendance_day；跨夜归属只改变 attendance_day，不改写原发生时间。",
        "源数据、规则或人员映射变化时，重算范围取旧影响区间与新影响区间并集；已月结期间只创建差异，不改快照。",
        "月结开始时生成 close_cutoff 并取得期间级锁；快照冻结截止点、源水位、OA 发布批次、规则/组织/假期台账水位、组织身份与版本、发生时编码/名称/父级路径、输入清单、行数与校验值。截止点之后的数据统一进入差异。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "7.6 核心状态机", 2)
    lifecycle_rows = [
        ["同步运行 SyncRun", "CREATED → RUNNING → SUCCEEDED / PARTIAL_FAILED / FAILED", "得力以 SyncPage 为原子单位：页面事实与 next_id 同事务提交，后续页失败保留已成功页；OA 暂存批次须整批校验后原子发布"],
        ["组织同步 OrganizationSyncBatch", "CREATED → EXTRACTING → STAGED → VALIDATING → DIFFING → READY_TO_PUBLISH / WAITING_REVIEW → PUBLISHING → PUBLISHED；可 FAILED/REJECTED", "AUTO 与 MANUAL 使用同一流程；无差异只记批次，有普通差异自动版本化，疑似复用/拆并/未知进入复核；同一来源实例只允许一个组织同步运行"],
        ["组织变更 OrganizationChangeReview", "OPEN → ASSIGNED → CONTINUITY_CONFIRMED / NEW_IDENTITY_CONFIRMED / SPLIT_MERGE_CONFIRMED / REJECTED", "待复核范围默认拒绝历史继承；复核须记录证据、操作者和生效时间，发布后触发授权快照重建、缓存失效和重算/差异"],
        ["人员来源绑定 EmployeeSourceBinding", "CANDIDATE → CONFIRMED；可 CONFLICT / CLOSED", "OA 工号与得力 employee_num 相同只生成候选；仅 ext_id/user_id 唯一校验及确认后可核算，禁止姓名自动匹配"],
        ["来源组织绑定 SourceOrganizationBinding", "CANDIDATE / PENDING_REVIEW → ACTIVE → MISSING_CANDIDATE → CLOSED；可 REJECTED", "ACTIVE 后不能原地更换 identity；关闭旧绑定再新建。单次缺失只进入 MISSING_CANDIDATE，可靠停用/删除语义或完整快照复核后才 CLOSED"],
        ["组织版本 OrganizationVersion", "DRAFT → PUBLISHED → SUPERSEDED", "PUBLISHED 不可编辑；新版本发布时关闭旧有效期并原子更新当前投影"],
        ["OA 授权快照 OaAuthorizationSnapshot", "STAGED → VALIDATED → PUBLISHED → REPLACED / STALE", "只能引用同一批次已发布的组织绑定与层级快照；禁止组织图和授权图一新一旧"],
        ["OA 标准单据", "PENDING / SUBMITTED / APPROVED / REJECTED / WITHDRAWN / VOIDED", "状态映射保留原始码；只有 APPROVED 且未撤回/作废参与核算"],
        ["日结果 AttendanceDayRevision", "DRAFT → CALCULATED → EXCEPTION / REVIEWED → CLOSED", "规则/证据变化生成不可变 revision；仅最新输入指纹可成为 current，CLOSED 只随反月结产生新版本"],
        ["异常 ExceptionCase", "OPEN → ASSIGNED → WAITING_EMPLOYEE / WAITING_OA → RESOLVED / IGNORED；可 REOPENED", "解决动作必须关联调整、源单据或明确忽略原因"],
        ["提前返岗 EarlyReturnCase", "DETECTED → CONFIRMED / REJECTED", "DETECTED 可影响出勤展示；当期 AUTO_FULL_PUNCH、OA_CANCEL_ONLY、HR_CONFIRM 或 HYBRID 策略决定确认触发，只有 CONFIRMED 可产生返还分录"],
        ["假期策略 LeavePolicyVersion", "DRAFT → VALIDATED → PUBLISHED → EXPIRED", "发布前校验并预览；PUBLISHED 不可编辑，变更创建新版本及生效日；已月结期间不回写"],
        ["月份 MonthClose", "OPEN → CALCULATING → READY_TO_CLOSE → CLOSED → REOPENED → CALCULATING", "CALCULATING 固定 close_cutoff 并等待截止点前任务完成；CLOSED 后任何变化只进入差异，反月结必须有权限和原因"],
        ["人工调整 Adjustment", "DRAFT → APPLIED → REVERSED", "APPLIED 后不可改写；更正先冲销原分录，再新增调整"],
        ["薪资来源批次 PayrollSourceBatch", "CREATED → STAGED → VALIDATING → VALIDATED → PUBLISHED；可 CONFLICT / FAILED / CANCELLED", "原文件/附件、映射和记录同批留痕；冲突未解决不得发布，同一内容哈希重放保持幂等"],
        ["工资期间 PayrollPeriod", "OPEN → COLLECTING → INPUT_LOCKED → CALCULATING → REVIEWING → READY_TO_PUBLISH → PUBLISHED → CLOSED", "REVIEWING 可退回 COLLECTING/CALCULATING；锁定后输入变化只标记过期。PUBLISHED/CLOSED 不反结覆盖，纠错创建补发、追扣或冲销运行"],
        ["工资运行 PayrollRun", "CREATED → SNAPSHOTTING → VALIDATING → READY → CALCULATING → CALCULATED → REVIEWING → REVIEWED → PUBLISHED → CLOSED；可 FAILED/RETURNED", "每次运行绑定唯一输入快照和规则集；制表与复核发布必须不同账号，复核人只可退回或发布，不能直接修数"],
        ["工资条 PayslipVersion", "DRAFT → PUBLISHED → SUPERSEDED", "只有独立复核发布完成的正式版本可发布；员工只见本人 PUBLISHED，纠错新建版本并保留旧版"],
        ["薪资调整 PayrollAdjustment", "DRAFT → VALIDATED → APPLIED → REVERSED", "人工补发追扣必须关联原期间/分段/项目和复核记录；更正使用冲销，不直接修改金额"],
        ["薪资授权 PayrollAuthorizationGrant", "DRAFT → ACTIVE → SUSPENDED / EXPIRED / REVOKED", "薪资权限管理员按预设模板直接配置，不设审批状态；自授权、兼任薪资数据角色或通配范围原子拒绝，撤销后会话/缓存/下载令牌及时失效"],
    ]
    add_table(doc, ["对象", "状态", "转换约束"], lifecycle_rows,
              [1.55, 2.35, 2.6], font_size=7.75, first_col_bold=True)
    add_callout(doc, "一致性要求", "得力单页提交、OA 整批发布、日结果 current 指针更新和月结快照各有明确事务边界。跨系统只要求最终一致，但必须可重放、可对账、可识别过期任务并告警；不得用一个笼统的“整批失败不部分提交”掩盖不同来源的原子单位。", fill=LIGHT_ORANGE, label_color=ORANGE)

    add_heading(doc, "8. 数据契约与外部依赖", 1)
    add_heading(doc, "8.1 外部依赖总览", 2)
    dep_rows = [
        ["得力云", "官方考勤数据 API", "上线后新增原始打卡及 GPS/外勤地点字段", "初始化前历史不纳入；需 App-Key/App-Secret"],
        ["致远 OA", "自建表只读查询 + 可替换标准对象/附件适配器", "员工、组织、岗位、任职、调动/调薪、账号状态、组织范围、考勤单据、薪资字段与关联文件", "人员组织与权限权威源；禁止写入；自建表物理字段后续映射，附件内容通道须单独确认"],
        ["HR Excel", "受控模板导入", "现有工资档案、法定计算资料与其他工资输入", "原文件、哈希、Sheet/行号、导入人和映射版本留存；校验发布后参与计算"],
        ["地图服务", "服务端受控适配器", "用户点击后绘制单次 GPS/外勤打卡位置", "坐标系须联调确认；供应商、Key 和转换策略可配置"],
        ["本系统配置", "管理后台维护或受控导入", "班次、日历、假别、工资方案/项目/标准/公式、税社保公积金政策、其他工资输入、制度和权限", "所有规则带有效期与发布版本；公式不允许任意脚本或 SQL"],
    ]
    add_table(doc, ["来源", "接入方式", "用途", "边界"], dep_rows, [1.0, 1.45, 2.45, 1.6], font_size=8.6, first_col_bold=True)

    add_heading(doc, "8.2 得力 API 产品约束", 2)
    deli_rules = [
        "生产基址为 https://v2-api.delicloud.com；所有请求使用 POST + JSON。考勤同步路径为 /v2.0/cloudappapi，并以 Api-Module: CHECKIN、Api-Cmd: checkin_query_init/checkin_query 区分指令。只允许 HTTPS。",
        "请求头包含 App-Key、13 位毫秒 App-Timestamp 与 App-Sig；得力外部协议要求对 path + timestamp + app_key + app_secret 计算小写的 Message-Digest Algorithm 5 摘要。该遗留算法只允许封装在得力适配器中用于协议兼容，禁止复用于密码、会话、文件完整性或任何内部安全控制。App-Secret 仅存服务端密钥设施，禁止进入前端、URL、日志、异常栈或导出。",
        "每套得力连接创建独立 source_instance_id；初始化状态、next_id、唯一键和重试均按该实例隔离。生产初始化需二次确认、变更记录和结果留档，测试环境不得使用生产凭据初始化。",
        "初始化成功后只能取得新增记录，不提供初始化前历史；本项目不补录历史。考勤同步使用服务端返回的 next_id，不能自行递增；单页最多 500 条，页面事实、批次明细和新 next_id 同事务提交，查询至空集合。",
        "得力只提供原始打卡事实，不提供排班、请假、加班审批或考勤计算结果；后者由致远自建单据和本系统规则产生。数据按设备上报顺序返回而非 check_time 排序，迟报须按 check_time 重算历史日期。",
        "得力以 ext_id 作为外部员工关联标识；employee_num 是独立工号字段。OA 工号与得力 employee_num 相同只能生成候选 EmployeeSourceBinding，ext_id/user_id 经过唯一性验证并确认后才可核算。",
        "check_type 兼容 fa、fp、pass、card、app_scan、gps、wifi、out_work、reissue、flexible；check_data 按多态载荷白名单解析，未知类型或解析失败只隔离该记录，不得导致整页失败。",
        "check_data 经纬度字段为 lat/lgt；仅 gps、out_work 且坐标合法时生成地图能力。文档未声明坐标系，须以真实样本确认坐标系、精度和 name/location/photo 语义后启用生产绘点。",
        "响应 code=0 为成功；网络错误、HTTP 5xx、109 超时和 110 限流采用指数退避与抖动，每次重试重新生成时间戳和签名；101—108 转配置/组织异常，不盲目重试。任何失败均不得推进 next_id。",
        "P0 使用读取与考勤同步接口白名单，禁止调用删除员工/部门、写员工/组织、人员特征和门禁抓拍接口；尤其不得调用会同时删除员工打卡记录的删除员工接口。",
        "系统保存来源记录 id、user_id、ext_id、terminal_id、check_type、check_time、最小化 check_data、同步批次和来源信封；坐标规范化到敏感列后从来源 JSON 移除，未经筛选的完整响应不得直接落库。",
    ]
    for item in deli_rules:
        add_bullet(doc, item)

    add_heading(doc, "8.3 致远接入与标准化逻辑契约", 2)
    add_callout(doc, "接入边界", "用户已确认请假、加班、外出、补签、调动、薪资等业务数据来自致远自建表，本次忽略物理表名与字段。V80 数据字典仅用于校准致远标准组织、岗位、人员、关系和附件元数据，不证明自建表具备任何固定字段。致远官方不建议直接操作数据库；本项目将只读自建表作为经业务确认的架构例外，并以最小授权和可替换适配层隔离升级风险。", fill=LIGHT_ORANGE, label_color=ORANGE)
    oa_rules = [
        "SeeyonCustomTableReader 使用独立只读账号读取白名单自建表/视图；不得执行 INSERT、UPDATE、DELETE、DDL、存储过程或创建对象。生产仅调用 DBA 审核的分页/时间窗口查询模板，禁止页面提交任意 SQL，并配置查询超时、最大行数和保护性限流。",
        "SeeyonStandardAdapter 负责标准员工、组织、岗位、任职、权限关系、流程与附件；可按实际许可选择只读标准视图、官方 REST 或 OA 侧受控服务，但统一输出同一内部逻辑契约。若 REST 已启用，标准组织、流程状态和附件内容优先使用官方接口；本系统始终不写回 OA。",
        "V80 字典显示标准对象可参考 ORG_MEMBER、ORG_UNIT、ORG_POST、ORG_LEVEL、ORG_RELATIONSHIP、CTP_ATTACHMENT、CTP_FILE 等；具体表结构仅作接入调研线索，不在 PRD 或领域模型中硬编码。ORG_MEMBER 的主组织/岗位与关系表中的兼职或次要任职须区分保存。",
        "所有致远单位、部门、人员、岗位、流程和附件来源 ID 均按字符串或精确 DECIMAL 保存并通过 API 以字符串返回，避免 19 位 Long 在 JavaScript/JSON Number 中丢失精度。",
        "每类自建对象通过版本化 OAQueryMappingProfile 定义逻辑业务键、人员、组织、起止时间、业务生效时间、提交/审批/最后变化时间、原始状态映射、关联单据、附件定位、时区、空值和能力标志；逻辑值可由单字段、复合字段或受控推导形成。",
        "映射配置可选择更新时间复合游标、变更序号或开放期间全量指纹扫描；删除/撤销可选择显式标志、可靠变更日志、完整快照差异或人工复核。不得要求来源必然存在 updated_at、source_detail_id、source_version 或 tombstone，也不得因某次未返回就推断删除。",
        "原始 OA 状态必须显式映射为标准状态；未知值一律隔离，绝不能默认视为已通过。只有映射为 APPROVED 且未撤回/作废的有效版本参与核算；内容变化由本系统生成 OADocumentVersion。",
        "人员、组织、岗位和任职须解析为内部稳定身份及有效期版本。来源明确提供业务有效时间时优先采用；否则保存 first_seen_at/last_seen_at 作为观测时间并标记 OBSERVED，不把来源更新时间当作历史业务生效时间。调动生效日的逻辑值无法可靠取得时，阻断对应工资分段。",
        "自动组织同步可覆盖当前投影，但每次均记录技术批次与差异；用户点击“同步组织架构”时执行全量比对并在有业务变化时形成组织版本。官方标准接口未承诺通用 updatedSince 分页，因此具体同步可采用分单位全量/分批对账，不预设通用增量能力。",
        "数据库可读附件元数据不代表可直接取得附件内容；致远附件可能加密。附件访问模式须配置为 OA_REST、OA_PLUGIN_SERVICE 或 MANUAL_IMPORT 等受控通道，禁止直接破解或读取 OA 文件目录。成功取得文件后由本系统计算内容哈希、执行白名单/大小/恶意文件扫描并保存受控副本。",
        "每次致远升级前执行字段指纹检查、脱敏样本试运行和集成回归；映射配置变更创建新版本并可回滚。字段缺失、类型漂移、业务键重复、未知状态、能力与声明不符或附件通道异常时整批隔离并告警。",
        "查询结果先进入暂存区；完成类型、行数、逻辑业务键、状态、必要值和批次完整性校验后发布。数据库或 REST 断连、超时、字段指纹变化时不清空既有正式数据；不同接入通道恢复后按幂等键补偿。",
    ]
    for item in oa_rules:
        add_bullet(doc, item)
    oa_objects = [
        ["EMPLOYEE", "来源人员定位、工号、姓名、状态", "形成致远权威人员身份；来源定位需稳定或由复合键可重复生成"],
        ["ORGANIZATION", "来源定位、编码、名称、父级、类型、状态", "形成内部组织身份与历史路径；缺业务有效时间时以观测版本保存"],
        ["ORG_CHANGE_LOG（如有）", "变更前后定位/编码/名称/父级及生效时间", "用于连续性判断；不存在时使用全量差异与人工复核，不假设必有变更日志"],
        ["EMPLOYMENT_ASSIGNMENT", "人员、主/兼职类型、组织、岗位、考勤地点、有效范围", "保存主任职与次要任职；缺有效范围时以观测时间标记待确认"],
        ["POSITION", "来源定位、编码、名称、组织/序列/职级、状态", "形成内部岗位身份和版本；物理来源由标准适配器决定"],
        ["HR_TRANSFER", "员工、旧/新组织/岗位/职级、业务生效、原始状态", "生效日与最终有效状态必须可映射，否则阻断工资分段"],
        ["SALARY_CHANGE（如有）", "员工、旧/新工资方案或项目、生效、原始状态", "若 OA 不保存金额，只同步调薪事实并由本系统定薪版本承接"],
        ["PAYROLL_RECORD / ATTACHMENT", "员工、期间、工资逻辑值、附件定位与元数据", "附件内容须经受控通道取得；内容哈希在本系统读取成功后计算"],
        ["ACCOUNT / ORG_DATA_SCOPE", "登录状态、授权节点、穿透关系、批次完整性", "由标准表/接口关系计算权限快照；无法完整解析时默认拒绝"],
        ["LEAVE / LEAVE_CANCEL", "假别、起止、原申请关联、实际返岗", "关联无法自动形成时隔离或人工复核，不臆造关联 ID"],
        ["OVERTIME", "起止、首次提交逻辑时间、原始状态", "首次提交无法可靠推导时不得自动判定 48 小时时效"],
        ["BUSINESS_TRIP / OUTING_REGISTER", "申请与登记关联、起止、目的地、事由", "关联策略由映射配置定义；缺失时不自动豁免"],
        ["REISSUE", "缺卡日期/卡点、首次提交逻辑时间、原始状态", "首次提交无法可靠推导时进入人工复核"],
    ]
    add_table(doc, ["逻辑对象", "业务必要信息", "缺失/适配处理"], oa_objects,
              [1.55, 2.8, 2.15], font_size=8.1, first_col_bold=True)
    add_callout(doc, "映射冻结点", "本次无需提供自建表名与字段。后续在《致远接入映射说明》中按对象冻结 QueryMappingProfile、脱敏样本和能力清单；无法取得补签首次提交、调动生效、撤销/作废、双表关联或附件内容时，必须选择明确降级/人工复核，不能用猜测值继续自动核算。", fill=LIGHT_ORANGE, label_color=ORANGE)

    add_heading(doc, "8.4 重算触发", 2)
    recalc_rows = [
        ["新增/补传打卡", "该员工 check_time 对应考勤日；跨日班次同时检查相邻日"],
        ["OA 单据新增、提交、通过、撤回、作废", "起止时间覆盖的全部考勤日；补签重算原缺卡日"],
        ["OA 修改人员或时间", "旧范围与新范围的并集"],
        ["班次、日历、规则变化", "生效范围内员工与日期；已月结月份仅产生差异"],
        ["销假、提前出勤确认或出差覆盖", "原申请范围、新有效范围和额度台账"],
        ["跨日卡或迟到加班单", "加班开始日及相邻自然日；同一卡只归属一次"],
        ["工号映射修正", "该员工未匹配原始记录及受影响日期"],
        ["组织名称/编码/来源 ID/父级/有效期/绑定变化", "旧组织版本与新组织版本覆盖日期的并集，以及受影响任职员工；开放月份重算，已月结月份只生成差异"],
        ["OA 调动/调薪新增、修改、撤回或生效时间变化", "工资期间内旧分段与新分段范围并集；工资关闭后生成 PayrollPostCloseDifference"],
        ["岗位/工资方案/项目/标准/定薪/折算规则变化", "生效范围内全部工资分段；已关闭工资不覆盖，只生成差异或补发追扣建议"],
        ["OA 薪资字段/附件或 HR Excel 新增、替换、删除", "命中业务键的员工、工资项目和期间；输入锁定后标记运行过期，关闭后只生成差异"],
        ["社保公积金/个税政策、员工基数或累计台账变化", "生效范围内法定计算行、实发与单位成本；关闭后不改原累计台账，只生成更正运行"],
        ["考勤月结新版本或工资选择的考勤快照变化", "引用该考勤期间的工资项目与员工分段；工资输入已锁定时标记过期并阻断发布"],
        ["人工调整", "对应日结果及月度、部门、公司汇总"],
    ]
    add_table(doc, ["触发事件", "重算范围"], recalc_rows, [2.1, 4.4], first_col_bold=True)

    add_heading(doc, "8.5 时间语义与数据最小化", 2)
    contract_rows = [
        ["时间字段", "source_time_raw 保留源值；occurred_at 保存绝对发生时刻；received_at 保存接收时刻；local_datetime 保存业务时区显示值；attendance_day 保存考勤归属日。跨夜只改变 attendance_day。"],
        ["来源载荷", "raw_payload 指完成字段最小化后的来源信封，不是未经处理的完整响应；只保存核算、审计和故障定位所必需字段。"],
        ["组织完整副本", "保存数据契约批准的组织节点、父子边、编码/名称/状态、任职、授权范围、源版本、同步批次及全部历史版本；不等于无选择复制致远整个数据库。"],
        ["薪资事实与快照", "OA/Excel 来源批次与文件引用、规范化有效版本、输入快照、分段、项目行、社保公积金、个税累计、差异、复核发布、工资条和成本分摊保存在 PAYROLL 独立表/schema；高度敏感值不得复制到普通 HR/考勤表、搜索索引、埋点或共享缓存。"],
        ["敏感字段", "外勤照片、手机型号、Wi-Fi MAC/SSID 默认在入库前移除；精确坐标规范化到数据库敏感列并加密/受控保存，不得在普通 JSON、日志、索引或缓存中重复出现。"],
        ["在线留存", "P0 的来源事实、版本、考勤/工资核算结果、假期台账、双月结快照、工资条、审计及规范化位置字段均保存在业务数据库；不建设独立归档层。薪资导出临时文件最长 24 小时删除，不等同于删除正式工资事实；调整保留策略须覆盖副本、缓存、导出与备份生命周期。"],
    ]
    add_table(doc, ["主题", "契约要求"], contract_rows, [1.35, 5.15], font_size=8.45, first_col_bold=True)

    add_heading(doc, "9. 交互与响应式要求", 1)
    add_heading(doc, "9.1 多端布局", 2)
    responsive = [
        ["PC（1440px 基准）", "按角色提供工作台；固定侧边导航与顶部上下文；列表/详情分栏。薪资核算使用类表格多 Sheet、冻结关键列、差异标记和计算轨迹侧滑详情；页面显隐不替代服务端授权。"],
        ["平板（768-1199px）", "侧边导航折叠；看板两列；筛选收纳为抽屉；详情使用全宽页或侧滑层；保留主要审批、查询和异常处理。"],
        ["手机（360-767px）", "移动优先的员工自助；底部/折叠导航、拇指可达主操作、指标卡片和渐进披露；表格转关键字段卡片。工资条使用当前安全会话和 no-store，按分段折叠展示，不横向缩小 PC 核算表。"],
        ["数据大屏（1920×1080）", "16:9 管理驾驶舱只展示考勤聚合；顶部核心 KPI、中部趋势/组织对比、两侧异常与新鲜度。不得请求或展示工资、奖金、个人成本、工资批次或可推算个人薪资的小组数据。"],
    ]
    add_table(doc, ["断点", "交互要求"], responsive, [1.55, 4.95], first_col_bold=True)
    for item in [
        "所有页面显示数据截至时间；同步过期时以醒目标识提示，但仍允许查看最近成功数据。",
        "列表加载、空数据、无权限、同步失败、计算中和已月结均提供明确状态与下一步动作。",
        "重要操作（人工调整、批量导出、管理者查看地图、月结、反月结）需要二次确认。",
        "移动端图表优先展示结论和可点击数据点；详细明细通过下钻查看，不缩小到无法阅读。",
        "三端使用同一指标、版本和授权服务，但分别使用最小响应 DTO；大屏和移动端不得复用包含额外敏感字段的 PC 详情响应。",
        "账号、角色、组织范围或查看对象切换时，先清除上一范围的页面状态、预取数据和客户端缓存，再按新授权重新请求；不得短暂闪现旧范围内容。",
        "薪资页进入后台、切换标签、系统任务预览或重新唤醒时遮罩；动态水印包含查看人、时间和请求标识。水印用于追责和降低泄露风险，不宣称能够完全阻止截屏。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "9.2 关键页面交互", 2)
    interaction_rows = [
        ["工作台", "默认显示当前开放月份；异常卡片可直接下钻到已带筛选条件的列表。"],
        ["员工月历", "日期格显示正常/异常/请假/外出等标签；点击打开证据链与反馈入口。"],
        ["班次配置", "按总部季节、地点、个人维护有效期；发布前校验日期空档和同层级重叠，预览指定员工指定日的命中结果。"],
        ["假别/特殊工时", "按假别维护额度来源、发放周期、公式、结转/失效、最小单位、日历口径、返岗模式和关系选项；发布前校验、预览并形成不可变版本。"],
        ["异常工作台", "左侧筛选、中央列表、右侧详情；保存调整后只重算受影响日期。"],
        ["组织镜像", "展示致远当前节点/层级、最新批次、数据截至时间和来源字段；提供“同步组织架构”按钮，运行中禁用重复触发并显示进度、结果与请求编号。"],
        ["组织历史/变更复核", "并排展示变更前后来源 ID、编码、名称、父级、有效时间和证据，以及受影响员工、月份、报表和权限；授权人员确认同一身份、新身份、拆分或合并，提交前二次确认并记录原因。"],
        ["工资数据准备", "分来源展示 OA 字段/附件、HR Excel、本系统工资档案、法定政策、年度累计、考勤月结和其他输入的新鲜度/缺口；下钻到文件、Sheet/行号和冲突，锁定后任何变化显示过期提示。"],
        ["调动影响/分段预览", "按员工时间轴并排展示调动前后组织、岗位、工资标准、生效日和工资分段；空档、重叠、OA 状态未生效或缺标准醒目标记并阻断。"],
        ["工资核算工作台", "类表格多 Sheet 展示员工、分段、工资项目、异常和本期/上期差异；冻结姓名/工号列，金额列按权限返回，点击金额查看公式轨迹但不暴露无权字段。"],
        ["工资复核发布", "复核发布人只读查看批次总额、税社保公积金、个税、异常、差异、组织/成本汇总和抽样明细，可退回或发布；与制表账号冲突时服务端禁止提交。"],
        ["我的工资", "员工通过正常有效会话查看本人已发布工资条；按销售/采购等工资分段折叠展示部门工资、岗位工资、考勤工资项、社保公积金、个税、其他增减项和实发合计，不显示内部公式或他人数据。"],
        ["成本分摊", "财务或显式授权角色查看责任组织/成本中心/项目分摊与勾稽；负责人默认只看满足小样本与互补抑制的本范围汇总，不提供个人总工资。"],
        ["看板", "公司→部门→员工→日期逐级下钻；无明细权限时停止在授权层级。"],
        ["签到排行", "默认进入本人所属地点/班次当日榜单；显示发布字段、数据截至时间和规则说明，留言支持举报。"],
        ["月结", "先展示硬阻断、软提醒和源数据水位，再允许确认；过程显示阶段状态，完成后提供版本号。"],
        ["定位查看", "GPS/外勤且 lat/lgt 合法时展示“查看地图”；点击后服务端鉴权，管理者记录原因，在弹层显示单点标记、地点文字、打卡时间和方式；地图失败回退为地点文字。"],
    ]
    add_table(doc, ["页面", "主要交互"], interaction_rows, [1.35, 5.15], first_col_bold=True)

    add_heading(doc, "9.3 视觉设计系统与三端体验", 2)
    add_body(doc, "体验参考北森一体化 HR SaaS 的多角色工作台、人员异动/考勤联动分段计薪、类 Excel 多 Sheet 核对、差异提示和 PC/移动工资条信息组织，借鉴其清爽、实用的呈现原则，不复制品牌、页面、算法或素材。本系统不设置业务审批和 MFA，以考勤/工资核对效率、独立复核发布和高敏数据安全为第一优先。")
    visual_rows = [
        ["整体风格", "清爽、可信、克制；PC/手机使用浅灰底、白色卡片和单一主品牌蓝，大屏可使用低眩光深色主题；语义色仅用于状态，避免装饰渐变、无意义动效和伪 3D 图表。"],
        ["信息层级", "每页先回答“当前状态、需要处理什么、下一步是什么”；首屏 3-5 个核心指标，异常/待办优先于装饰性统计，二级信息渐进展开。"],
        ["栅格与密度", "采用 8px 基础间距、统一圆角/阴影/字号/图标令牌；PC 支持舒适/紧凑密度，手机触控目标不小于 44×44px，大屏关键数字远距离可读。"],
        ["表格与筛选", "PC 表格支持列显隐、冻结、排序、筛选摘要、保存视图和批量操作；手机只显示关键字段并进入详情；大屏只展示聚合，不放可编辑表格。"],
        ["图表", "优先趋势、排名、结构和目标差异；统一指标定义、时间范围、单位、图例和数据截至时间。颜色不是唯一编码，悬停/点击可解释且下钻受权限控制。"],
        ["反馈状态", "骨架加载、空数据、无权限、同步过期、部分失败、计算中、已月结、地图不可用均有专用状态与下一步，不使用空白页或只显示错误码。"],
        ["可访问性", "正文与背景对比度至少 4.5:1；支持键盘焦点、语义标签、200% 缩放、系统字号和减少动态效果；异常不能只靠红绿色区分。"],
        ["敏感界面", "位置和薪资不出现在大屏、通知金额、浏览器标题、最近访问缩略图或第三方埋点；薪资页面强制动态水印、离开/后台遮罩、no-store；发布、导出和权限变更不满足最近登录条件时要求重新登录。"],
    ]
    add_table(doc, ["设计维度", "强制规范"], visual_rows, [1.35, 5.15], font_size=8.35, first_col_bold=True)
    add_callout(doc, "角色化首页", "员工首页突出今日状态、月历、异常、假期和已发布工资条入口；考勤员首页突出同步健康、异常队列和考勤月结；薪酬制表/复核角色使用独立工资工作台；老板/高管大屏不展示薪资。角色切换不改变底层权限，服务端仍逐请求鉴权。", fill=LIGHT_BLUE, label_color=BLUE)

    add_heading(doc, "10. 非功能与安全要求", 1)
    add_heading(doc, "10.1 容量、性能与可用性", 2)
    nfr_rows = [
        ["容量基线", "P0 只按现网规模设计，不采用未证实的人员增长或规划上界。架构设计开始时采集当前在职人数、实际日均/峰值打卡、月末在线并发、平均载荷和数据库现存量，形成容量基线。"],
        ["存储验证", "核心数据、规范化位置字段、组织完整副本、来源版本、身份绑定、层级/授权快照和变更审计全部在线保存在业务数据库，P0 不建设归档层。以实测现网峰值和 36 个月时间累积数据完成存储、索引、备份与恢复 POC；该累积仅考虑时间增长，不推定业务规模增长。"],
        ["查询性能", "50 并发下常用列表 P95 <=2 秒；个人/部门看板 P95 <=3 秒；单条详情 P95 <=1 秒。"],
        ["三端呈现", "PC/手机核心页面首个可用内容 <=2.5 秒；大屏 1920×1080 连续运行 8 小时无明显错位或内存持续增长，刷新失败保留上次成功数据并标识时间。"],
        ["导出性能", "50,000 行以内 <=60 秒；超过阈值转异步任务，不阻塞页面。"],
        ["同步时效", "默认每 5 分钟执行；上游正常时 95% 新数据 10 分钟内可见，99% 在 30 分钟内可见。"],
        ["核算/月结", "在现网容量基线 POC 下，全月重算与月结暂定 <=30 分钟；实测基线取得后冻结最终指标。"],
        ["工资核算/关闭", "在现网在职人数和工资项目基线 POC 下，正式工资运行暂定 <=30 分钟；员工本人已发布工资条 P95 <=1 秒，薪资工作台常用筛选 P95 <=3 秒。"],
        ["积压恢复", "以现网观测到的最大日峰值作为积压样本；在上游限流允许范围内，4 小时内完成入库、受影响日期重算和异常生成。"],
        ["可用性", "月度可用性 >=99.5%（计划维护除外）；同步故障不影响既有结果查询。"],
        ["一致性", "同源幂等率 100%；汇总与下钻差异为 0；得力单页、OA 发布批次、组织图+任职+授权快照、日结果 current 更新和月结快照分别按其事务边界保证原子性。"],
        ["灾备", "RPO <=15 分钟，RTO <=4 小时；须具备持续增量日志或等效时间点恢复能力，每日全量备份作为基线，季度完成数据库、规则、月结清单与水位一致恢复演练。"],
        ["浏览器", "最近两个主版本 Chrome/Edge、iOS Safari、Android Chrome；支持 360px 及以上宽度。"],
        ["可观测性", "记录同步水位、耗时、数量、错误和重试；连续失败 3 次或超过 15 分钟无成功触发告警。"],
    ]
    add_table(doc, ["维度", "验收基线"], nfr_rows, [1.35, 5.15], font_size=8.7, first_col_bold=True)
    add_callout(doc, "容量说明", "用户已确认按现有规模建设、数据保留在数据库。架构团队不得以历史规划上界采购或分区；应先采集现网基线，再用现网峰值和 36 个月时间累积量验证在线查询、备份与恢复。基线未取得不阻断总体架构设计，但会阻断生产资源规格冻结。", fill=LIGHT_ORANGE, label_color=ORANGE)

    add_heading(doc, "10.2 安全与个人信息边界", 2)
    security = [
        "全链路使用 TLS；App-Secret、数据库密码等密钥加密保存，页面只掩码展示，日志不得输出明文。",
        "服务端对每次请求重新校验功能、组织、字段和月份状态权限，不依赖前端隐藏。",
        "所有接口默认拒绝并使用统一授权中间件；对象查询、字段投影、聚合、附件、导出、地图和后台任务均在数据访问前校验。任何客户端传入的用户、角色或组织范围都不能覆盖会话与致远授权快照。",
        "致远原始授权组织 ID 必须通过与组织图同一已发布批次的 SourceOrganizationBinding 解析到内部组织身份；禁止用来源 ID、编码或名称直接等值关联业务数据。疑似复用或连续性待复核时，受影响身份和 scope 立即 fail closed；组织图与授权快照禁止半批发布。",
        "采用安全会话：Cookie 场景使用 __Host- 前缀、Secure、HttpOnly、SameSite 与 CSRF token/Origin 校验；令牌不进入 URL、日志或 Web Storage。注销、撤权和切换账号时服务端失效会话并清除客户端敏感缓存。",
        "CORS 仅允许精确受控源；配置 CSP（地图域仅加入必要 connect/img 白名单）、HSTS、frame-ancestors、nosniff、Referrer-Policy、输出编码、参数化 SQL 和上传白名单，禁止任意查询模板、脚本或外部 URL。",
        "连续 5 次登录失败锁定 15 分钟；普通会话闲置 30 分钟、绝对时长 8 小时且可配置。本系统不建设 MFA。薪资后台默认仅允许公司内网、VPN 或配置的可信网络访问（员工查看本人已发布工资条除外），闲置 15 分钟锁定、绝对会话 8 小时；发布、批量导出和薪资权限变更要求最近 15 分钟内完成过登录，否则使用现有 OA/本地账号重新登录。",
        "登录、对象访问、批量查询、导出和地图接口按账号、IP、设备和资源设置分级限流；连续枚举 ID、跨组织拒绝或大量 403/404 触发安全告警和会话风险处置。",
        "不采集或保存人脸模板、指纹模板、身份证特征；得力人员特征接口不在范围内。",
        "P0 将规范化精确坐标在线保存在业务数据库，不执行自动归档或删除；位置字段继续按高敏感数据独立授权、加密并审计。未来启用删除/降精度策略须版本化并覆盖副本与备份生命周期。",
        "外勤照片、手机型号、Wi-Fi MAC/SSID 默认不落库；地图第三方服务需完成供应商与个人信息评估，坐标仅在用户点击查看时发送，地图 Key 按密钥管理。",
        "raw_payload 入库前执行字段白名单/删除策略；精确坐标不得在来源 JSON、日志、搜索索引或缓存中形成不受控副本。",
        "签到排行使用独立发布数据模型，只含经确认的名次、展示名和时间粒度；不得复用原始打卡详情接口。",
        "留言支持敏感词、举报、屏蔽和审计；哺乳、孕检、病假原因等敏感信息不得出现在排行、留言提示或聚合看板。",
        "权限变更、位置查看、批量导出、人工调整、规则变更、月结和反月结日志至少保存 3 年。",
        "PAYROLL 使用独立表/schema、数据库账号、权限命名空间、字段加密密钥和 DTO；证件、税务标识和社保公积金标识采用字段级信封加密，数据库管理员不得同时取得解密密钥。普通 HR、考勤、BI 和报表服务不得直连薪资明细库；P0 入库白名单不接收银行账号。",
        "薪资授权使用独立 PayrollAuthorizationGrant，明确账号、预设角色、法人/薪资组/成本中心/内部组织身份、动作、字段、历史期间、用途和有效期。致远 OA 范围只能作为上限；考勤主管、HR 普通、负责人、管理层和系统管理员均默认无工资明细。",
        "同一工资批次的制表人与复核发布人必须不同。薪资权限管理员按预设模板直接配置，不设置权限审批；授权人与被授权人相同、权限管理员持有薪资数据角色、系统管理员授予薪资角色或通配范围时，服务端和数据库约束均原子拒绝。变更立即审计、通知固定 HR/安全联系人并进入每日权限差异清单。",
        "员工本人工资接口从会话推导 employee_id，不接受客户端参数决定本人对象；管理接口按对象、期间、字段和用途逐次授权。列表、总数、搜索、聚合、分页、批量、导出和后台任务均先限定授权范围再查询；混入一个越权对象的批量请求整体拒绝。",
        "工资、奖金、个税、社保公积金基数和个人成本不得进入大屏、普通 HR/部门看板、签到排行或通知预览。工资通知只显示“工资条已发布”。负责人/高管的薪资汇总使用独立页面，默认最小群组 5 人并执行互补抑制，防止通过筛选差值推算个人。",
        "所有薪资响应使用 private, no-store；禁止共享缓存、Service Worker、浏览器持久缓存、离线包和页面预取。服务端不缓存个人工资明细；受控汇总缓存键包含用户、授权版本、字段策略、工资批次和查询条件。",
        "薪资批量导出默认关闭并使用独立权限；创建、执行、下载三次重新鉴权，冻结人员/期间/字段/用途。链接绑定申请人和会话，默认 10 分钟、一次性；临时文件最长 24 小时删除。P0 不存在银行付款文件生成或下载能力。",
        "应用日志、网关、APM、Trace、错误堆栈、埋点和搜索索引不得记录工资金额、工资条、证件、个税或社保公积金数据。审计仅记录主体/对象哈希、动作、字段策略、用途、工资批次、授权版本、决定/原因、风险和 request_id。",
        "生产数据库访问必须关联工单、双人批准、限时 JIT、命令审计和会话记录；开发、运维、DBA、薪资业务人员职责分离。测试环境只使用合成或不可逆脱敏数据，备份恢复到非生产环境前再次脱敏。",
        "紧急薪资访问使用 break-glass 流程，限定对象、字段和时长，自动通知数据负责人并事后复核。未来财务、银行、税务/社保直接申报接口使用独立服务身份、PAYROLL scope/audience、mTLS 或等效认证、短时令牌、字段白名单和出站审计。",
        "调岗、离职、账号停用、权限撤销或组织/岗位绑定变化后，薪资会话、缓存、下载令牌、搜索权限投影和实时订阅须在 5 分钟内失效。管理、财务和汇总角色的历史访问按当前有效薪资授权与发生时组织/成本中心共同判断；员工本人仍可按本人对象关系查看本人已发布历史工资条。",
        "发布门禁包含权限矩阵单元/集成测试、改 URL/抓包对象替换测试、批量与导出越权测试、SAST/SCA/密钥扫描、依赖清单及经授权的 API 安全测试；高危问题未关闭不得上线。",
        "上线前向员工告知处理的数据种类、用途、访问角色、保存期限和查询/更正渠道。",
    ]
    for item in security:
        add_bullet(doc, item)
    add_callout(doc, "隐私原则", "位置用于核验一次打卡事实，不用于持续追踪员工；薪资用于本人知情、核算、法定扣缴、复核发布、成本与审计的明确目的。精确坐标和个人薪资均默认不可见、不可进入普通导出/大屏/通用看板；动态水印只能降低泄露风险，不能替代服务端授权。", fill=LIGHT_RED, label_color=RED)

    add_heading(doc, "11. 数据指标与埋点", 1)
    add_heading(doc, "11.1 产品运行指标", 2)
    kpis = [
        ["自动核算率", "无需人工调整即可形成日结果的员工日数 / 总可核算员工日数", ">=90%［试运行目标］", "月"],
        ["异常处理时长", "异常产生至首次处理的中位时长", "首版建立基线，后续下降 50%", "周/月"],
        ["同步完整率", "本地源记录数 / 上游金标准记录数", "100%", "联调/上线"],
        ["同步新鲜度", "新源数据到本系统可见的延迟", "P95 <=10 分钟", "实时"],
        ["组织变更留痕率", "已发布组织字段变化中存在对应版本与批次的记录数 / 全部变化记录数", "100%", "每批/月"],
        ["组织镜像完整率", "已发布节点、父子边、任职与 scope 数 / 致远金标准数量", "100%", "每次全量"],
        ["未决高风险组织变更", "影响当前权限或待月结期间的 ID 复用/连续性未知数量", "发布与月结前为 0", "实时/月结"],
        ["金标准准确率", "系统结果与业务确认样本一致数 / 样本总数", "100%", "发布前"],
        ["汇总勾稽差异", "汇总指标 - 下钻明细重新聚合结果", "0", "月结"],
        ["员工自助覆盖", "查看个人考勤或年假的活跃员工 / 在职员工", "上线后建立基线", "月"],
        ["工资分段准确率", "金标准员工期间中组织/岗位/工资方案/边界完全一致的分段数 / 全部分段", "100%", "工资发布前"],
        ["工资自动核算率", "无需人工改金额即可通过校验的员工数 / 本期应发薪员工数", ">=90%［试运行目标］", "工资期"],
        ["工资勾稽差异", "工资项目合计、工资条合计、批次合计与成本分摊合计之间的差额", "0", "工资关闭"],
        ["薪资高风险异常", "任职空档/重叠、OA/Excel 冲突、法定政策/年度期初累计缺失、职责冲突或授权过期数量", "发布前为 0", "实时/工资期"],
        ["法定计算完整率", "已生成社保、公积金和个税有效结果的应发薪员工 / 全部应发薪员工", "发布前 100%", "工资期"],
        ["薪资敏感访问审计率", "薪资查看/拒绝/导出/复核/发布/授权变更有审计的次数 / 全部次数", "100%", "实时/月"],
    ]
    add_table(doc, ["指标", "定义", "目标", "周期"], kpis, [1.35, 2.9, 1.25, 1.0], font_size=8.3, first_col_bold=True)

    add_heading(doc, "11.2 关键事件", 2)
    events = [
        ["sync_completed", "双源同步批次结束", "source, batch_id, duration_ms, pulled, inserted, updated, failed, watermark", "监控时效与完整性"],
        ["organization_sync_completed", "自动或手动组织同步结束", "trigger_type, batch_id, fetched, unchanged, created, renamed, recoded, moved, disabled, review_required, result", "监控组织副本、版本与人工按钮使用"],
        ["organization_change_reviewed", "疑似复用/拆分/合并变更被复核", "change_id, change_type, decision, effective_at, reviewer_role, evidence_type", "审计组织身份连续性结论；不得记录敏感业务值"],
        ["calculation_completed", "日/月核算任务结束", "period, employee_count, day_count, exception_count, duration_ms, rule_version", "监控核算效率"],
        ["exception_resolved", "异常被调整、忽略或退回", "exception_type, resolution, handler_role, elapsed_minutes", "分析人工工作量"],
        ["month_closed", "月结成功", "period, version, deli_watermark, oa_batch, blocker_count, warning_count", "追踪正式版本"],
        ["report_exported", "导出完成", "report_type, period, row_count, org_scope, includes_sensitive_fields", "审计与使用分析"],
        ["feedback_submitted", "员工提交反馈", "feedback_type, linked_date, linked_exception, device_type", "分析员工异议"],
        ["reissue_evaluated", "补签同步或状态变化", "attendance_date, submitted_at, deadline, in_time, approval_status", "监控补签时效"],
        ["overtime_evaluated", "加班同步或重算", "start_at, end_at, submitted_at, within_48h, evidence_minutes, final_minutes", "核验加班规则"],
        ["ranking_interacted", "查看排行、留言或举报", "scope_type, shift_id, action, content_status", "使用分析与内容审计"],
        ["location_map_opened", "点击查看地图", "viewer_role, employee_id_hash, record_id, reason_code, check_type, provider, source_coordinate_system, result", "敏感访问与地图可用性审计"],
        ["authorization_decided", "敏感允许/拒绝或连续越权", "request_id, subject_hash, data_domain, action, object_type, policy_version, oa_scope_version, decision, reason_code", "权限解释、异常枚举检测和安全审计；不得记录敏感值"],
        ["payroll_input_locked", "工资输入快照锁定", "payroll_period_id, run_id, employee_count, assignment_version, attendance_close_version, config_hash", "证明工资输入和版本稳定；不得记录金额"],
        ["payroll_source_imported", "OA/Excel 薪资批次校验或发布", "source_type, batch_id, template_version, records, conflicts, file_hash, result", "监控来源质量与冲突；不记录工资值"],
        ["payroll_calculation_completed", "工资运行完成", "period, run_id, employee_count, segment_count, exception_count, duration_ms, input_hash, rule_set_version", "监控分段与核算质量；不得上报金额"],
        ["statutory_calculation_completed", "税社保公积金计算结束", "period, run_id, employee_count, policy_versions, missing_count, result", "监控法定计算完整性；不得上报基数或金额"],
        ["payroll_review_completed", "工资批次复核退回或发布", "period, run_id, decision, reviewer_role, sod_result, reason_code", "审计制表与复核发布职责分离"],
        ["payslip_published", "工资条发布", "period, payslip_version, employee_count, template_version", "追踪发布完整性；不含个人金额"],
        ["payroll_sensitive_accessed", "工资明文、批量或导出访问", "request_id, subject_hash, action, object_hash, field_policy, purpose, grant_version, decision, risk_result", "薪资高敏审计与越权告警；不记录敏感值"],
        ["payroll_post_close_difference_created", "关闭后人事/考勤/规则变化", "original_period, difference_type, affected_employee_count, source_version, target_run_type", "追踪补发追扣来源；不含金额"],
    ]
    add_table(doc, ["事件", "触发", "关键字段", "用途"], events, [1.35, 1.45, 2.55, 1.15], font_size=7.9, first_col_bold=True)

    add_heading(doc, "12. 验收标准", 1)
    add_heading(doc, "12.1 数据接入与一致性", 2)
    ac_data = [
        ["AC-D01", "正式上线初始化成功后记录时间；初始化前历史不采集；普通用户不能再次初始化。"],
        ["AC-D02", "得力同步从持久化 next_id 开始，每页 <=500 条；该页全部入库后才推进游标。"],
        ["AC-D03", "同一得力记录 ID 重放 3 次，原始数据和核算结果均不重复。"],
        ["AC-D04", "设备补传昨天打卡时按 check_time 归入昨天；开放月份自动重算。"],
        ["AC-D05", "OA 工号与得力 employee_num 相同只生成 CANDIDATE；ext_id/user_id 唯一且 EmployeeSourceBinding=CONFIRMED 后才核算。ext_id 为空、重复或冲突时进入异常，不按姓名猜测。"],
        ["AC-D06", "所有核心数据均带 company_id；跨公司 ID、缓存键、导出任务和对象 URL 不得互相访问或碰撞。"],
        ["AC-D07", "得力原始卡、OA 原始载荷及其历史版本不可由业务操作覆盖；人工调整只能追加并可追溯到原证据。"],
        ["AC-D08", "切换得力连接或凭据实例后使用新的 source_instance_id；即使来源记录 ID 相同，也不得与旧实例数据碰撞。"],
        ["AC-D09", "得力生产请求仅走 HTTPS，签名路径、13 位时间戳、App-Key/App-Sig 与 Api-Module/Api-Cmd 符合契约；App-Secret、签名材料不得出现在前端、URL、日志或导出。"],
        ["AC-D10", "模拟网络/5xx/109/110 时指数退避并重新签名且游标不前移；101—108 不盲目重试。故障恢复后重放不重不漏。"],
        ["AC-D11", "应用运行身份只获批准的读取/考勤同步能力；删除员工/部门、写员工/组织、人员特征和门禁抓拍调用全部被配置及测试阻断。"],
        ["AC-O01", "每个致远自建对象通过已发布 OAQueryMappingProfile 形成稳定逻辑业务键；键为空或批内重复时整批隔离，不要求固定物理单据/明细 ID。"],
        ["AC-O02", "只有已审批且未撤销/作废单据参与核算；其他状态可查但不抵扣考勤。"],
        ["AC-O03", "OA 单据修改或撤销时，开放月份自动重算；已月结月份只生成差异。"],
        ["AC-O04", "数据库断连、超时或字段变化时，本批整体失败，既有数据不被清空。"],
        ["AC-O05", "补签/加班首次提交、调动生效、出差/外出关联等逻辑值由映射或受控推导取得；不能可靠取得时进入数据异常或配置的人工复核，不伪造来源值。"],
        ["AC-O06", "OA 查询仍为只读；员工办理补签、加班、出差或销假时跳转/指引至 OA，不由本系统写回。"],
        ["AC-O07", "每个 OA 对象按映射配置采用更新时间复合游标、变更序号或开放期间全量指纹扫描；内容变化形成本系统内部版本，不要求来源提供 source_version。"],
        ["AC-O08", "OA 查询先入暂存区，字段/类型/行数/唯一键/状态校验通过后整批发布；校验失败时正式版本和水位不变。"],
        ["AC-O09", "员工、组织、任职、账号状态和组织数据范围均从致远只读同步；本系统不得由得力、姓名匹配或本地维护覆盖其权威值。"],
        ["AC-O10", "调岗按生效日期形成历史 OrganizationVersion；当前访问按最新致远授权快照解析后的内部组织身份校验。同步失败不扩大访问范围，账号停用/移出范围在下一成功批次后失去权限。"],
        ["AC-O11", "首次成功同步后，本系统当前组织数量、父子层级、编码、名称和状态与致远金标准一致；每个组织均存在内部 organization_identity_id、有效 OrganizationVersion 和 SourceOrganizationBinding。"],
        ["AC-O12", "连续自动同步或点击“同步组织架构”3次且源数据未变化时，只增加同步批次与 last_seen_at，不增加 OrganizationVersion；同一来源实例同时只能运行一个组织同步任务。"],
        ["AC-O13", "致远组织仅改名、改编码或调整上级时，旧版本 valid_to 正确关闭，新版本按生效时间启用，内部身份保持不变；变更前考勤仍显示原编码、名称和父级路径。"],
        ["AC-O14", "同一致远组织 ID 或编码先属于采购部、后被销售部复用时，系统创建两个内部组织身份；跨期正式报表不合并，当前销售部负责人访问旧采购部列表、详情、数量、导出、缓存和地图全部拒绝。"],
        ["AC-O15", "来源 ID 变化但被确认是同一组织时，新旧来源绑定按有效期指向同一内部身份；名称、编码和来源 ID 同时变化且无可靠证据时进入 PENDING_REVIEW，受影响历史访问默认拒绝。"],
        ["AC-O16", "具备 organization.sync.execute 权限的用户点击手动按钮后可查看全量差异和请求编号；无权限用户、重复提交或并发运行不能创建第二个任务。组织连续性结论只能由具备 organization.change.review 权限的用户提交并完整审计。"],
        ["AC-O17", "组织绑定变化时，即使致远 assignment_id 未变，也按变化时点拆分本地任职版本；开放月份重算旧/新区间并集，已月结月份仅生成 PostCloseDifference。"],
        ["AC-O18", "组织全量批次存在重复来源 ID、孤儿、层级环、悬空任职/scope、数量或总哈希不符时不得发布；上次组织图和授权快照保持可用且不新增权限。"],
        ["AC-O19", "组织单次查询缺失只进入 MISSING_CANDIDATE；具备可靠停用/删除语义，或完整全量快照确认并完成复核后才关闭。系统不得用名称相似或成员重叠自动合并组织。"],
        ["AC-O20", "映射配置发布前用脱敏样本试运行，验证逻辑业务键、字段类型、状态、时间和业务必要值；配置变更形成新版本且可回滚。"],
        ["AC-O21", "自建表字段缺失、类型变化、业务键冲突或状态出现未知值时，批次不得静默发布，并产生结构漂移告警；未知状态不得默认视为已通过。"],
        ["AC-O22", "致远单位、部门、人员、岗位、流程与附件 19 位来源 ID 经采集、数据库、API 和前端往返后字符值完全一致，不发生浮点精度损失。"],
        ["AC-O23", "OA 数据库账号执行 INSERT、UPDATE、DELETE、DDL 和未授权查询均失败；查询超时、最大行数、白名单和升级字段指纹检查可验证。"],
        ["AC-O24", "附件必须分别验证元数据可读、内容可取、哈希可计算；仅能读元数据时，自动薪资附件接入不得标记完成。附件只能经授权后端通道取得。"],
        ["AC-O25", "组织权限范围无法完整解析、同步失败或快照过期时默认拒绝，不得沿用可能已过期的扩大权限；恢复后按新快照授权。"],
    ]
    add_table(doc, ["编号", "验收标准"], ac_data, [1.0, 5.5], font_size=8.9, first_col_bold=True)

    add_heading(doc, "12.2 核算、月结与报表", 2)
    ac_calc = [
        ["AC-C01", "正常、迟到、早退、缺上班卡、缺下班卡、缺勤、请假、外出和加班金标准案例结果 100% 一致。"],
        ["AC-C02", "每条日结果可追溯到班次/日历版本、原始打卡、OA 单据、规则版本和人工调整。"],
        ["AC-C03", "不同类型单据时间重叠时不重复计算，并生成单据冲突异常。"],
        ["AC-C04", "人工调整必须填写原因，调整前后值、操作人和时间不可覆盖。"],
        ["AC-C05", "补签在配置的一周截止边界前提交且最终获批时替代对应缺卡；超时不自动认可；撤回后恢复缺卡并重算。"],
        ["AC-C06", "当日上班卡与次日凌晨下班卡命中跨日窗口时，凌晨卡归入前一考勤日且不再作为次日上班卡；边界冲突进入复核。"],
        ["AC-C07", "加班单在实际结束后 48 小时内首次提交并最终获批时参与认定；截止后 1 秒提交则加班为 0 或进入已配置特批流程。"],
        ["AC-C08", "无加班单时原始晚间/凌晨卡仍可追溯，最终加班为 0；系统不得把原始时间改写为标准下班时间。"],
        ["AC-C09", "出差申请与外出登记均有效且目的地、事由、起止时间完整时才自动豁免；任一缺失生成手续不完整异常。"],
        ["AC-C10", "大连晚间加班的 30 分钟晚餐扣减由规则或单次人工确认决定，不得无条件扣除。"],
        ["AC-C11", "假期内完整有效打卡区间可按实际出勤计算并生成提前返岗事件；只有 CONFIRMED 写入返还台账。AUTO_FULL_PUNCH、OA_CANCEL_ONLY、HR_CONFIRM、HYBRID 四种模式分别按配置触发，零散单卡始终复核；有效出差覆盖年假时不得重复计假。"],
        ["AC-M01", "存在同步失败、人员未匹配、班次/日历缺失、核算失败、影响当月归属/权限的组织变更待复核，或按策略要求确认但仍未确认的提前返岗事件时禁止月结；自动确认模式不产生无意义阻断。"],
        ["AC-M02", "月结是原子操作；故障时不出现部分人员已锁定，重复并发只允许一个成功。"],
        ["AC-M03", "月结后迟报数据只进入差异；反月结需专门权限和原因，旧版本永久保留。"],
        ["AC-R01", "公司汇总可下钻至部门、员工和日期；下钻数据重新汇总与上级差异为 0。"],
        ["AC-R02", "导出严格遵守筛选条件、组织范围和字段权限，并包含生成时间和月结版本。"],
        ["AC-R03", "查询日期跨越同一来源 ID 从采购部变为销售部时，正式报表按发生时组织输出两行且公司合计一致；不得因来源 ID、编码或当前名称相同而合并。"],
        ["AC-R04", "月结后致远改名、改编码、改来源 ID 或调父级，旧月结中的组织身份、版本、来源 ID、编码、名称、路径和哈希均不变；变化只进入差异或反月结新版本。"],
    ]
    add_table(doc, ["编号", "验收标准"], ac_calc, [1.0, 5.5], font_size=8.8, first_col_bold=True)

    add_heading(doc, "12.3 班次与假期规则", 2)
    ac_policy = [
        ["AC-S01", "总部规则在 04-30/05-01、09-30/10-01 分别正确切换；12-31 在未确认前进入规则空档，不静默沿用。"],
        ["AC-S02", "大连班次识别 07:30-12:00 与 13:00-16:30 两段，12:00-13:00 不计入应出勤分钟。"],
        ["AC-S03", "成都指定人员命中 09:00-12:00、13:00-18:00；上海指定人员仍命中总部季节规则，证明个人规则优先于地点。"],
        ["AC-S04", "同一对象同层级有效期重叠时禁止发布；调岗或个人规则到期前后按日期分别使用正确版本。"],
        ["AC-S05", "P0 只允许发布 FIXED_ADMIN；NIGHT、ROTATION、FLEX、TEMP 配置返回“首版未启用”。模型仍可保存 schedule_type、周期组、跨日锚点、弹性窗口和临时分配，不以同自然日单班作为不变量。"],
        ["AC-L01", "入职满 1 年后年假基线为 5 天；年假和病假 0.5 天按当日所选半日实际时段换算，不固定为 4 小时。"],
        ["AC-L02", "丧假关系下拉选择配偶/子女/父母时校验 3 天，选择兄弟姐妹/祖父母/外祖父母时校验 1 天。"],
        ["AC-L03", "护理假以 15 个连续自然日计算，包含周末和法定节假日；起算条件按确认后的规则执行。"],
        ["AC-L04", "事假允许最小 0.5 小时；病假允许最小 0.5 天；不符合最小/递增规则时提示，不静默取整。"],
        ["AC-L05", "哺乳/孕检特殊时段按员工及有效期叠加到基础班次；到期恢复下一优先级班次，重叠豁免不重复扣减。"],
        ["AC-L06", "假期策略必须经过 DRAFT、VALIDATED、PUBLISHED、EXPIRED 状态；发布前可预览，发布后不可编辑，变更创建新版本和生效日且不重写已月结月份。"],
        ["AC-L07", "额度来源可选自动规则、导入或人工；自然年、入职周年、自定义发放及结转/失效参数均可按适用范围配置，所有余额变化形成不可变台账。"],
    ]
    add_table(doc, ["编号", "验收标准"], ac_policy, [1.0, 5.5], font_size=8.7, first_col_bold=True)

    add_heading(doc, "12.4 员工端、权限与多端", 2)
    ac_ui = [
        ["AC-U01", "员工只能查看本人考勤、OA 单据、年假、制度和反馈；直接访问他人工号 URL 返回无权限。"],
        ["AC-U02", "员工可从具体日期发起反馈，状态至少包含待处理、处理中、已解决和已驳回，回复全程留痕。"],
        ["AC-U03", "公司/部门看板在无个人明细权限时停止下钻，不泄露姓名、请假类型或位置。"],
        ["AC-U04", "排行只返回本人所属发布范围的名次、展示名和允许时间粒度；接口不得返回他人的原始卡、OA 单据、设备、位置或敏感假别。"],
        ["AC-U05", "不同班次不按绝对签到时间直接混排；补签、人工调整和外出卡是否入榜按配置执行，默认不参与。"],
        ["AC-U06", "留言创建、回复、删除、举报和屏蔽均校验榜单范围并留痕；屏蔽后普通员工不可见，审计记录仍保留。"],
        ["AC-P01", "无 attendance.location.map 权限或不在当前致远组织范围时，通过页面、URL、接口、导出或缓存均不能取得地图或精确坐标；员工仅可查看本人记录。"],
        ["AC-P02", "check_type 为 gps/out_work 且 lat、lgt 合法时显示地图入口；缺失、非数值、越界或其他打卡方式不显示入口。来源 name/location 文字与坐标使用同一位置权限，未经授权不得返回。"],
        ["AC-P03", "将 check_data.lgt 正确解析为 longitude；坐标系未用真实样本确认前不得在生产地图绘制未经验证点位。"],
        ["AC-P04", "点击查看地图产生包含查看人、对象、原因、来源记录、供应商、坐标系和结果的审计；普通报表与公司/部门看板不包含个人位置或坐标。"],
        ["AC-P05", "地图供应商超时或失败时页面显示地点名称/地址及“地图暂不可用”，不影响考勤详情；坐标只在点击后发送给已配置供应商。"],
        ["AC-W01", "PC、平板和 360px 手机端均可完成查询、查看详情和反馈，核心操作不依赖横向滚动。"],
        ["AC-W02", "常用列表 P95 <=2 秒，看板 P95 <=3 秒；50,000 行导出 <=60 秒或转异步。"],
        ["AC-W03", "员工、考勤、HR、管理者和系统管理员登录后只获得对应工作台模块与最小字段；未授权模块不出现在 HTML、接口响应、预取、缓存或前端状态中。"],
        ["AC-W04", "在 360、768、1440 和 1920×1080 下完成对应核心流程，无内容遮挡、关键文字截断或非必要横向滚动；手机触控目标不小于 44×44px。"],
        ["AC-W05", "员工首页到当日详情、从异常日发起反馈均不超过 2 次操作；考勤员从工作台进入带条件异常清单不超过 2 次。"],
        ["AC-W06", "列表、指标、图表和地图覆盖加载、无数据、无匹配、无权限、失败、过期和部分降级；错误页面不暴露 SQL、堆栈、接口地址或内部对象标识。"],
        ["AC-W07", "键盘可完成查询、筛选、详情和反馈；状态不只依赖颜色，关键文字对比度至少 4.5:1，页面放大 200% 后仍可操作。"],
        ["AC-W08", "大屏不返回姓名、工号、敏感假别、坐标、单据原因、留言、工资、奖金、个人成本或工资批次；个人下钻必须转到重新鉴权的 PC 端。"],
        ["AC-W09", "切换账号、角色、组织或员工对象后，旧范围页面数据、预取与客户端缓存立即清除；抓包中不存在上一范围业务载荷。"],
    ]
    add_table(doc, ["编号", "验收标准"], ac_ui, [1.0, 5.5], font_size=8.8, first_col_bold=True)

    add_heading(doc, "12.5 防越权与敏感域安全", 2)
    ac_security = [
        ["AC-SEC01", "普通员工将 URL、查询参数、请求体或 Header 中的 employee_id 改为他人值，列表、详情、统计、附件和地图均不得返回他人数据；响应不泄露对象是否存在。"],
        ["AC-SEC02", "部门负责人把对象 ID 改为当前致远授权范围外员工，或把 organization_id 改为上级/兄弟部门，页面和 API 均拒绝；系统管理员也不会因技术角色自动取得业务明细。"],
        ["AC-SEC03", "列表、总数、搜索建议、聚合图表和分页游标在同一授权范围内计算；无明细权时不能通过 count、筛选差异、小样本或错误信息反推出个人。"],
        ["AC-SEC04", "API 使用字段白名单；抓包检查所有角色响应均不包含未展示的坐标、敏感假别、内部主键或薪资字段，写接口不能通过附加 JSON 属性修改无权字段。"],
        ["AC-SEC05", "权限缩小、调岗、账号停用或角色撤销后，下一次成功同步/策略发布立即使旧 URL、浏览器缓存、服务端缓存、下载令牌和 WebSocket 订阅失效。"],
        ["AC-SEC06", "导出创建、执行和下载均重新鉴权；混入越权对象的批量请求整体拒绝且不暴露哪条 ID 存在。下载令牌绑定用户、短时、一次性，复制给他人不可用。"],
        ["AC-SEC07", "连续枚举 opaque ID、跨组织访问、批量 403/404 和异常抓包重放命中限流与告警；审计可关联 request_id、策略版本和 OA 快照，但不记录敏感字段值。"],
        ["AC-SEC08", "自动化权限矩阵覆盖员工、部门负责人、考勤、HR、管理层、系统管理员与审计员的横向/纵向越权；改 URL、抓包替换、缓存、导出、批量和推送测试均为发布阻断门禁。"],
        ["AC-SEC09", "PAYROLL 正式启用后，除员工本人和显式薪酬岗位外所有现有角色均默认拒绝；D/E 级响应 no-store，不进入大屏、通知预览和第三方埋点。"],
        ["AC-SEC10", "OA 或 PAYROLL 授权快照超过新鲜度阈值时，组织明细、位置、批量导出、工资详情、复核和发布均 fail closed；同步恢复后旧范围不自动恢复。"],
        ["AC-SEC11", "无 CSRF token、伪造 Origin、通配凭据 CORS、GET 修改状态或脚本读取会话 Cookie 均失败；令牌不出现在 URL、Web Storage、日志和错误信息。"],
        ["AC-SEC12", "若启用 WebSocket/SSE，伪造来源、订阅范围外主题和篡改消息动作均拒绝；注销、停用或撤权后连接立即关闭，消息载荷只含授权字段。"],
        ["AC-SEC13", "上线前经授权的越权安全测试中，URL 篡改、抓包重放、对象枚举、缓存/导出/实时通道越权成功数为 0；Critical/High 安全发现为 0 方可发布。"],
        ["AC-SEC14", "组织来源 ID、编码、内部身份或版本绑定变化发布后，旧范围的服务端缓存、浏览器缓存、下载令牌、搜索索引权限投影和 WebSocket/SSE 订阅立即失效；客户端提交旧 OA ID/编码不能恢复旧权限。"],
        ["AC-SEC15", "经可靠证据确认来源 ID 变化但内部组织身份连续后，当前 OA scope 可按既定历史访问策略查看该身份的历史；授权日志必须记录 source binding、organization hierarchy 与 OA authorization snapshot 版本。"],
    ]
    add_table(doc, ["编号", "验收标准"], ac_security, [1.0, 5.5], font_size=8.45, first_col_bold=True)

    add_heading(doc, "12.6 工资核算、分段与工资条", 2)
    ac_payroll = [
        ["AC-PAY01", "员工 7 月 16 日调动时，系统生成 [07-01,07-16) 销售部与 [07-16,08-01) 采购部两个无重叠、无空档的工资分段；7 月 16 日整日及其计薪工作量归采购部。"],
        ["AC-PAY02", "两个分段分别引用发生时销售/采购组织身份与版本、岗位身份与版本、工资方案、定薪和折算规则；后续改名或 ID 复用不改变已关闭工资。"],
        ["AC-PAY03", "部门工资、岗位工资和基本工资分别计算、分别展示；员工可展开查看数量、基数、折算和取整依据，不出现无法解释的合并金额。"],
        ["AC-PAY04", "加班、请假、缺勤、迟到早退等工资项目按考勤事实发生日归入当天有效工资分段，并引用正式考勤月结版本。"],
        ["AC-PAY05", "工资总额等于全部分段项目、月度项目及补发追扣之和；工资条合计、批次合计和项目重新聚合差异为 0。"],
        ["AC-PAY06", "成本分摊比例之和为 100%，分摊金额与可分摊工资金额差异为 0；调整成本责任组织不改变员工应发工资。"],
        ["AC-PAY07", "相同输入快照和规则版本连续计算 3 次，分段、项目金额、舍入差额和结果哈希完全一致。"],
        ["AC-PAY08", "开放工资期内调动单撤回、修改生效日或调薪变化时，旧、新影响区间并集全部重算；锁定后先标记输入过期，不静默改结果。"],
        ["AC-PAY09", "工资关闭后到达的 OA 文件/Excel、调动、调薪、考勤或规则变化只生成 PayrollPostCloseDifference，不覆盖原工资条、分段或关闭/发布版本链。"],
        ["AC-PAY10", "补发追扣可追溯到原工资期间、原分段、原工资项目、原发生/责任组织、原输入版本、调整原因和来源证据。"],
        ["AC-PAY11", "任职/工资段空档或重叠、调动缺生效时间、工资方案/标准多匹配或无匹配、公式循环/除零、无效/重复/来源冲突输入、法定政策或个税累计缺失时禁止关闭工资。"],
        ["AC-PAY12", "员工工资条清楚展示月内各部门/岗位分段、部门工资、岗位工资、考勤工资项、其他增减项、员工承担社保公积金、个税、分段小计与实发合计。"],
        ["AC-PAY13", "部门变化但工资标准相同仍保留两个可解释分段；同一部门内岗位或定薪变化也按有效期切段。"],
        ["AC-PAY14", "考勤月结版本未选择、已过期或工资锁定后发生变化时不得发布正式工资条；页面显示影响员工和差异来源。"],
        ["AC-PAY15", "同一账号不能完成同一工资运行的制表和复核发布；复核发布人不能修改输入、公式或金额，并发、重放或直接调用状态接口均不能绕过。"],
        ["AC-PAY16", "工资条只有在运行已计算、校验通过、独立复核发布完成且模板已发布时生成；纠错发布新版本，员工历史版本可追溯但默认只展示最新有效版。"],
        ["AC-PAY17", "已发布/关闭工资不可原位修改；纠错只能创建补发、追扣或冲销运行，并引用原期间、原分段和原项目。"],
        ["AC-PAY18", "P0 对全部应发薪员工计算社保、公积金和个税；任一必要输入缺失时显示明确异常并阻止发布，不得以 0 表示已计算。"],
        ["AC-PAY19", "系统不存在银行付款文件生成、下载、付款状态或银企直连接口；工资核算和工资条发布不依赖银行状态。"],
        ["AC-PAY20", "任一工资输入可追溯到 OA 字段、OA 单据/附件版本、HR Excel 批次/Sheet/行号或本系统录入记录。"],
        ["AC-PAY21", "OA 附件或 Excel 后续修改/删除时，已关闭工资仍可凭系统保存的原文件引用/受控副本、哈希、规范化记录和输入快照重现。"],
        ["AC-PAY22", "相同业务键出现多个来源值时不得静默覆盖；未被来源权威策略唯一判定且未人工解决的冲突阻止输入锁定。"],
        ["AC-PAY23", "工资项目可分别配置按实际计薪工作日或实际计薪工时折算；同一输入快照重复计算 3 次，工作量、分段、金额、舍入差额和结果哈希一致。"],
        ["AC-PAY24", "分段实际工作日/工时数量之和等于月度对应数量，每个数量可追溯到日历、排班、考勤月结及 PayrollWorkMeasureSnapshot；缺勤不得重复扣减。"],
        ["AC-PAY25", "社保、公积金按项目展示基数、上下限处理、个人比例/金额、单位比例/金额和政策版本；单位承担额不减少员工实发。"],
        ["AC-PAY26", "个税明细可还原本期收入、年度累计收入、各类扣除、已预扣税、本期税额和政策版本；跨月累计连续且关闭后不可原地改写。"],
        ["AC-PAY27", "工资条满足应发收入 − 员工承担社保公积金 − 个税 − 其他员工扣减 + 补发追扣 = 实发，批次重新聚合差异为 0。"],
        ["AC-PAY28", "同一其他工资项目可配置 HR_CENTRAL、DEPARTMENT_SCOPED 或两者录入；部门录入只能选择授权员工/部门，抓包修改归属后服务端拒绝。"],
        ["AC-PAY29", "其他工资项目定义变更只影响生效日后的工资期间；已关闭期间仅生成补发、追扣或冲销差异。"],
        ["AC-PAY30", "工资输入锁定后 OA、附件、Excel、考勤或法定政策变化只将运行标记过期，不静默进入当前结果。"],
        ["AC-PAY31", "P0 不采集或保存银行付款账号字段；OA/Excel 入库白名单遇到此类字段时拒绝或删除并记录数据质量结果。"],
    ]
    add_table(doc, ["编号", "验收标准"], ac_payroll, [1.0, 5.5], font_size=8.45, first_col_bold=True)

    add_heading(doc, "12.7 薪资高敏权限与安全", 2)
    ac_payroll_security = [
        ["AC-PSEC01", "所有考勤、HR 普通、部门负责人、管理层和系统管理员账号在没有独立 PayrollAuthorizationGrant 时访问任一薪资接口或字段均被拒绝。"],
        ["AC-PSEC02", "员工只能通过本人接口查看本人已发布工资条；替换 URL、请求体、Header 或抓包对象 ID 后不能获得他人工资，也不能判断对象是否存在。"],
        ["AC-PSEC03", "销售转采购案例中，员工可见两个工资段和合计；两个部门负责人即使有成本汇总权，也只能获得各自部门满足小样本规则的汇总。"],
        ["AC-PSEC04", "员工按本人对象关系可查看本人已发布历史工资条，后续组织变化不影响本人访问；管理、财务和汇总角色查看历史工资时，同时校验当前薪资授权与记录发生时组织/成本中心。当前部门、OA 组织 ID/编码/名称变化或复用不得授予他人工资访问权。"],
        ["AC-PSEC05", "HR 普通角色只看到调动事实和资料状态；页面与抓包响应中不存在工资、薪级、税费、社保公积金基数或计算中间值。"],
        ["AC-PSEC06", "财务接口只返回已发布批次的法定申报汇总、单位承担额或会计成本分摊必需字段，不返回假勤原因、位置、工资公式和无关项目。"],
        ["AC-PSEC07", "系统管理员可维护技术配置和任务状态，但不能读取工资明文、取得解密密钥、生产直库查询或给自己授予薪资角色。"],
        ["AC-PSEC08", "列表、总数、搜索建议、分页、聚合和批量接口均在授权范围内计算；混入任一越权 ID 的批量请求整体拒绝且不指出具体对象。"],
        ["AC-PSEC09", "所有薪资响应使用角色/用途字段白名单；抓包、错误响应、前端源码、接口描述和客户端状态中均无未授权字段。"],
        ["AC-PSEC10", "薪资后台不在可信网络、会话闲置超过 15 分钟、敏感操作最近登录超过 15 分钟或授权快照过期时，工资明文、复核、发布、导出和权限变更均拒绝；重新登录使用现有单因子账号，不要求 MFA。"],
        ["AC-PSEC11", "工资页面不进入共享缓存、Service Worker、离线包、浏览器持久缓存或预取；注销、调岗和撤权后旧页面无法恢复内容。"],
        ["AC-PSEC12", "薪资导出在创建、执行、下载三次鉴权；链接复制给其他账号、超过 10 分钟、已使用一次或权限撤销后均无法下载。"],
        ["AC-PSEC13", "人数少于配置阈值（默认 5 人）的薪资汇总不展示精确值；组合筛选、相邻月份和父子组织差值测试不能推算个人工资。"],
        ["AC-PSEC14", "工资金额、税务/社保公积金字段和工资条内容在应用日志、网关、APM、Trace、埋点、搜索索引和普通错误信息中的命中数为 0。"],
        ["AC-PSEC15", "薪资查看/拒绝、批量、导出、未脱敏、复核、发布和授权变更审计覆盖率为 100%，审计不包含敏感明文。"],
        ["AC-PSEC16", "调岗、离职、账号停用或授权撤销后 5 分钟内，旧会话、服务端缓存、下载令牌、权限投影和实时连接全部失效。"],
        ["AC-PSEC17", "生产 DBA 在没有应用解密服务和独立运维工单授权时不能还原税务标识或薪资密文；非生产环境检索不到真实工资明细。"],
        ["AC-PSEC18", "工资、奖金、个人成本和工资批次不出现在数据大屏、通知金额、浏览器标题、普通 HR/部门看板、签到排行或第三方分析工具。"],
        ["AC-PSEC19", "P0 路由清单不存在银行付款文件接口；未来财务/银行/税务接口使用无 PAYROLL scope 的服务身份访问全部拒绝，Webhook、消息和普通 API 日志不含薪资明文。"],
        ["AC-PSEC21", "薪资权限管理员给本人授权、持有薪资数据角色、授予通配范围，或系统管理员授予薪资角色时，API 和数据库均原子拒绝；权限变更审计并通知固定联系人。"],
        ["AC-PSEC20", "上线前 URL/抓包、BOLA/IDOR、批量枚举、权限提升、缓存、导出、日志、数据库和内部人员越权测试成功数为 0；Critical/High 安全发现为 0。"],
    ]
    add_table(doc, ["编号", "验收标准"], ac_payroll_security, [1.0, 5.5], font_size=8.25, first_col_bold=True)

    add_heading(doc, "12.8 架构不变量验收", 2)
    ac_architecture = [
        ["AC-A01", "得力每页事实、批次明细和 next_id 原子提交，后续页故障不回滚已成功页；OA 暂存批次仅在整批校验通过后发布。两类故障重放均不重复。"],
        ["AC-A02", "相同员工、考勤日和输入指纹连续核算 3 次，只保留一个等价 current 结果；旧任务在新证据到达后完成时不得覆盖新 revision。"],
        ["AC-A03", "OA 单据或人员映射修改时间范围后，旧范围与新范围并集全部被重算；已月结日期只产生差异。"],
        ["AC-A04", "月结取得 close_cutoff 后，截止点前任务全部完成才发布；之后到达的数据只进差异。快照可还原源水位、输入清单、规则/组织/台账版本、行数与校验值。"],
        ["AC-A05", "APPLIED 人工调整和已生效假期分录不可直接编辑或删除；纠错以 REVERSED/冲销加新记录表达。"],
        ["AC-A06", "状态机拒绝非法跳转：失败的同步事务单元不得推进其水位，开放月不得越过待月结直接关闭，已月结结果不得原地修改。"],
        ["AC-A07", "raw_payload 不含未批准保存的照片、手机/Wi-Fi 标识或重复坐标；规范化坐标只存在于受控敏感列，抽查日志、索引、缓存和普通导出均无法恢复这些字段。"],
        ["AC-A08", "灾备演练证明 15 分钟内增量可恢复，并同时恢复规则版本、月结清单和同步水位；仅完成每日全量备份不视为达标。"],
        ["AC-A09", "现网基线与 36 个月时间累积数据 POC 证明在线数据库查询、备份和恢复达到指标；P0 不依赖独立归档层。"],
        ["AC-A10", "OrganizationVersion 与 SourceOrganizationBinding 的有效区间不可重叠；任意历史日期均至多解析到一个组织身份和一个版本，当前投影可从版本表完整重建。"],
        ["AC-A11", "自动与手动组织同步对同一源快照计算出完全相同的差异和版本结果；同步中途失败时正式版本、当前投影、授权快照和水位均保持原值。"],
        ["AC-A12", "月结快照可还原每条结果的 organization_identity_id、organization_version_id、发生时编码/名称/父级路径和组织批次校验值；后续同步不得改写该快照。"],
        ["AC-A13", "PayrollSegment 在工资期间内无重叠、无未解释空档；任意时点至多解析到一个任职、组织、岗位、定薪和工资方案版本。"],
        ["AC-A14", "PayrollInputSnapshot 可还原 OA 调动/任职/岗位与薪资字段/附件、HR Excel 批次、工资配置、其他输入、税社保公积金政策、个税累计、考勤月结版本和校验值；锁定后变化不得进入同一运行。"],
        ["AC-A15", "PayrollComponentLine 与 PayrollCostAllocationLine 独立保存且分别勾稽；修改成本责任不能改变员工工资金额或薪资查看授权。"],
        ["AC-A16", "已关闭工资、已发布工资条和已应用薪资调整不可原地修改；关闭后变化只生成差异、补发追扣、冲销或经授权的新期间版本。"],
        ["AC-A17", "PAYROLL 明细库、服务账号、权限命名空间、加密密钥和 DTO 与普通 HR/考勤隔离；测试环境、备份恢复和生产直库均通过安全门禁。"],
    ]
    add_table(doc, ["编号", "验收标准"], ac_architecture, [1.0, 5.5], font_size=8.7, first_col_bold=True)

    add_heading(doc, "12.9 最小金标准案例", 2)
    golden = [
        ["正常行政班", "上下班各一张有效卡", "正常，证据链完整"],
        ["迟到/早退", "超过配置宽限", "分钟数与公式一致"],
        ["单张打卡", "仅上午或仅下午一张卡", "明确缺上班卡或缺下班卡"],
        ["设备迟报", "昨天离线卡今天上报", "归属昨天并触发重算"],
        ["班中请假", "已审批请假与工作时段部分重叠", "只计算交集，休息段不计"],
        ["全天外出", "已审批外出且无设备卡", "按配置免打卡或产生差异"],
        ["加班差异", "审批 3 小时、实际佐证 2 小时", "同时展示两者和最终认定值"],
        ["跨夜加班", "当日上班卡 + 次日凌晨下班卡", "凌晨卡归属前日且不误报缺下班卡"],
        ["加班超时", "实际结束后 48 小时 1 秒提交", "原始卡保留，认定加班为 0/进入特批"],
        ["补签边界", "一周截止前 1 秒与后 1 秒提交", "分别允许参与与标记超时"],
        ["总部换季", "04-30/05-01、09-30/10-01", "分别命中冬令、夏令、夏令、冬令"],
        ["大连半天假", "上午半天与下午半天", "分别按 4.5 小时与 3.5 小时处理"],
        ["提前返岗四模式", "年假未结束产生完整有效打卡", "按 AUTO_FULL_PUNCH/OA_CANCEL_ONLY/HR_CONFIRM/HYBRID 分别验证确认与返还"],
        ["OA 撤销", "已生效单据后撤销", "开放月冲销；已结月产生差异"],
        ["组织改名/改编码", "采购部 CG01 改为采购中心 CG02，业务身份连续", "新建版本但内部身份不变；历史日期仍显示 CG01/采购部"],
        ["组织 ID/编码复用", "D001/CG01 原采购部停用后被销售部使用", "创建新内部身份；销售部权限不能读取旧采购部数据"],
        ["手动组织同步", "连续点击同步按钮且第二次源数据无变化", "第一次按差异发布；第二次只留批次，不产生空版本或并发任务"],
        ["组织变更待复核", "来源 ID、编码、名称同时变化且无变更日志", "进入 PENDING_REVIEW，受影响历史访问默认拒绝，复核后再确定延续或新身份"],
        ["月结并发", "两名用户同时确认", "仅一个成功且版本唯一"],
        ["地图位置", "GPS/外勤记录含合法 lat/lgt", "显示单点地图；非法/缺失坐标无入口；越权访问拒绝且不返回坐标"],
        ["URL/抓包越权", "员工替换他人 ID；负责人替换范围外组织；复用他人下载链接", "详情、统计、缓存与导出全部拒绝，不泄露对象是否存在并产生安全审计"],
        ["月中调动分段", "7 月 1-15 日销售部，16-31 日采购部，部门/岗位工资不同", "两个无重叠分段分别计算固定项和对应日期考勤项，工资条逐段展示并正确合计"],
        ["工资关闭后异动", "已关闭工资后补录 7 月 16 日生效的调动/调薪", "原工资不变，生成引用原月份/分段/项目的补发追扣差异"],
        ["成本与查看权分离", "销售/采购分别承担各自分段成本，负责人有汇总权", "分摊合计正确；负责人只见满足阈值的本部门汇总，不能见个人总工资或另一分段"],
        ["薪资 URL/抓包越权", "员工替换他人工资条 ID；HR 普通或系统管理员直接调用薪资接口", "全部拒绝且响应不含对象存在性或工资字段，形成不含金额的安全审计"],
        ["来源冲突", "同一员工/项目/有效期在 OA 附件与 HR Excel 出现不同值", "不按更新时间覆盖；命中来源策略或留痕解决前阻止锁定"],
        ["工作日/工时折算", "两个项目分别配置实际计薪工作日与实际计薪工时", "按各自 WorkMeasureSnapshot 计算，分段量等于月度量且不与缺勤重复扣减"],
        ["税社保公积金", "基数触及上下限、个人/单位比例不同且存在年中个税期初累计", "个人/单位金额、累计预扣和实发勾稽正确，单位承担不扣员工实发"],
        ["职责分离", "制表员直接发布本人运行或复用复核发布请求", "状态机拒绝，复核发布人不能修数，工资条不得越过复核发布"],
    ]
    add_table(doc, ["案例", "输入", "预期"], golden, [1.4, 2.6, 2.5], font_size=8.5, first_col_bold=True)

    add_heading(doc, "13. 项目计划与风险", 1)
    add_heading(doc, "13.1 建议里程碑", 2)
    milestones = [
        ["第1-3周", "需求、数据契约与规则评审", "PRD、OA/附件/Excel 字段映射、考勤/工资项目字典、法定政策与个税期初累计、折算/取整确认单、金标准案例"],
        ["第3-5周", "交互、安全与技术评审", "手机/PC/大屏与工资工作台原型、PAYROLL 权限矩阵、制表/复核发布互斥、单因子补偿控制、威胁建模和接入方案"],
        ["第4-8周", "数据底座与来源接入", "得力增量、致远组织/岗位/调动/薪资附件只读接入、HR Excel 受控导入、有效期模型、冲突和监控"],
        ["第6-11周", "考勤核算、异常、月结与报表", "考勤规则引擎、异常工作台、考勤月结版本和清单汇总"],
        ["第8-15周", "工资与法定扣缴核算", "工资方案/项目/标准、工作日/工时分段、税社保公积金、个税累计、差异、复核发布、工资条和成本分摊"],
        ["第10-15周", "管理后台与三端前台", "角色工作台、手机/PC/考勤大屏、我的工资、假期、制度、反馈、权限和地图"],
        ["第15-17周", "联调、安全与金标准测试", "考勤/工资/税社保准确性、勾稽、职责分离、URL/抓包、缓存/导出、日志/数据库、三端和性能报告"],
        ["第18周", "试运行与双月结演练", "试运行问题清单、考勤月结与工资关闭演练、个税累计连续性、工资条确认、培训和上线评审"],
    ]
    add_table(doc, ["时间", "里程碑", "主要交付物"], milestones, [1.1, 2.1, 3.3], font_size=8.7, first_col_bold=True)
    add_callout(doc, "估算前提", "社保、公积金、个税及 OA/Excel 双来源进入 P0 后，当前 18 周为重叠开发的初步建议。团队至少需产品 1、UI/交互 1、前端 2、后端 3、QA 2、运维/安全与薪酬业务代表兼职；取得数据字典、工资/法定政策、年度期初累计、项目折算和现网容量后重新估算。")

    add_heading(doc, "13.2 主要风险", 2)
    risks = [
        ["致远自建表结构漂移或映射不完整", "中", "高", "V80 字典仅校准标准对象；自建对象使用版本化 OAQueryMappingProfile、脱敏样本试运行、字段指纹、未知状态隔离与可回滚发布"],
        ["组织变化被自动覆盖而丢失历史", "中", "高", "当前投影与不可变版本分表；任何字段变化自动版本化，手动按钮只触发全量比对，不作为唯一留痕入口"],
        ["组织 ID/编码复用导致权限串用", "中", "高", "使用内部组织身份和有效期来源绑定；疑似复用默认新身份/待复核，旧缓存、导出和实时订阅立即失效"],
        ["得力 ext_id 未配置或与候选人员冲突", "高", "高", "工号一致只生成候选 EmployeeSourceBinding；ext_id/user_id 唯一验证并确认后核算，禁止按姓名自动合并"],
        ["设备离线迟报", "高", "中", "按 check_time 重算；月结后差异和反月结版本"],
        ["OA 直读数据库受升级影响", "中", "高", "将只读自建表作为显式例外；白名单视图/模板、分页限流、查询超时、字段指纹、升级回归及可替换适配器隔离"],
        ["OA 无统一更新时间/删除标志", "中", "高", "映射配置选择复合游标、变更序号或开放期间全量指纹扫描；删除/撤销采用可靠语义、完整快照差异或人工复核"],
        ["OA 附件元数据可读但内容加密/不可取", "中", "高", "上线前确认 OA_REST、OA_PLUGIN_SERVICE 或 MANUAL_IMPORT；内容取得并计算哈希前不得完成自动薪资附件接入"],
        ["致远 19 位 ID 被前端截断", "中", "高", "数据库精确保存、服务端与前端统一字符串传输，加入端到端精度验收"],
        ["班次日期/午休存在空档", "高", "高", "发布前校验有效期；12-31 和总部午休确认前列为配置阻断"],
        ["OA 缺少补签/销假/双表外出字段", "高", "高", "扩充只读视图；无法提供时进入受审计的人工复核，不伪造源单据"],
        ["跨日卡误归属", "中", "高", "可配置窗口、同一卡唯一归属、冲突清单和边界金标准"],
        ["假期参数尚未逐项冻结", "高", "中", "使用受控版本化策略承载额度来源、发放、结转/失效与返岗模式；测试前完成配置"],
        ["坐标系或地图供应商不匹配", "中", "高", "用真实 gps/out_work 样本确认坐标系；服务端地图适配器按配置转换，失败降级为地点文字"],
        ["位置数据越权", "中", "高", "致远组织范围与独立地图权限取交集，二次点击、原因、审计和导出限制"],
        ["URL/抓包导致对象级越权", "中", "高", "统一服务端授权、默认拒绝、先授权后查询、opaque ID、字段白名单、缓存/导出三次鉴权和专项回归"],
        ["OA 与 HR Excel 来源冲突", "高", "高", "来源已确认但实际字段、附件/模板和优先级未冻结；按业务键识别冲突，未命中来源权威策略时阻止锁定并人工留痕处置"],
        ["折算重复扣减", "高", "高", "已确认实际工作日/工时和生效日归新部门；逐项目冻结数量来源、分子分母、缺勤处理和取整，金标准覆盖重复扣款"],
        ["其他工资项目持续扩张", "高", "高", "先形成工资项目字典，明确来源、HR/部门录入模式、组织范围、分段/发生日/月度归属、成本和计税属性；新增项目走版本化发布"],
        ["员工工资与成本分摊混淆", "中", "高", "PayrollComponentLine 与 PayrollCostAllocationLine 分离，分别勾稽；成本责任不等于个人工资查看权限"],
        ["薪资复用考勤/OA权限", "中", "高", "PAYROLL 独立授权对象，OA 范围只限制不授予；短会话、可信网络、职责分离、字段白名单和专项安全准入"],
        ["单因子账号被盗", "中", "高", "已确认不建设 MFA；以可信网络、短会话、最近登录重验、限流/风险告警、no-store、水印和全量审计补偿，并接受剩余风险"],
        ["工资数据经日志/导出/运维泄露", "中", "高", "独立 schema/账号/密钥，no-store，导出三次鉴权，日志零敏感值，JIT 生产访问和非生产脱敏"],
        ["法定政策与期初累计缺失", "高", "高", "P0 必算税社保公积金；取得参保地、基数/上下限、个人/单位比例、个税累计规则、专项扣除和年中期初累计后方可冻结物理模型与金标准"],
        ["个税更正级联后续月份", "中", "高", "TaxYtdLedgerSnapshot 按期不可变；前月更正生成后续受影响期间差异与更正运行，不静默回写已关闭累计"],
        ["签到排行引发隐私或内容风险", "中", "中", "最小发布字段、按班次分组、后台开关、举报屏蔽和审计"],
        ["现网容量基线未测量", "中", "中", "采集当前人数、峰值、并发、载荷和数据库现存量；按现网峰值与 36 个月累积数据完成 POC 后冻结资源"],
        ["未来新增夜班/轮班改变核算", "低", "高", "P0 禁止启用非固定行政班；领域模型保留 schedule_type、周期、跨日锚点与临时分配策略"],
        ["报表范围持续膨胀", "中", "中", "以本 PRD 清单为 P0；新模板通过变更评审进入后续版本"],
    ]
    add_table(doc, ["风险", "概率", "影响", "应对"], risks, [2.15, 0.7, 0.7, 2.95], font_size=8.3,
              alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER,
                          WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT], first_col_bold=True)

    add_heading(doc, "14. 架构决策确认与剩余数据契约", 1)
    add_heading(doc, "14.1 本轮已确认决策", 2)
    decision_rows = [
        ["D-01", "致远数据字典与自建表", "部分确认", "V80 数据字典已提供并用于标准平台对象校准；自建业务表物理表名/字段本次按用户要求忽略，后续由版本化映射说明冻结"],
        ["D-02", "人员、组织与组织权限", "已确认", "致远 OA 为权威源；报表按历史快照，访问按当前致远组织范围穿透"],
        ["D-03", "假期发放与返岗", "已确认", "额度来源、发放、折算、结转/失效及四种返岗模式均采用受控可配置策略"],
        ["D-04", "容量与留存", "已确认", "按现网规模设计；核心数据与规范化位置字段在线保存在数据库，P0 不建归档层"],
        ["D-05", "非行政班型", "已确认", "首批无夜班、轮班、弹性或临时排班；P0 不启用，领域模型预留扩展"],
        ["D-06", "经纬度地图", "已确认", "gps/out_work 的 lat/lgt 合法时可点击地图；服务端授权、单点展示并审计"],
        ["D-07", "三端体验与防越权", "已确认", "手机/PC/大屏按角色提供最小视图；统一服务端授权防 URL、抓包、枚举、缓存和导出越权，启用 PAYROLL 独立域"],
        ["D-08", "组织架构同步与历史", "已确认", "致远为权威源；自动同步当前副本并对变化自动版本化；手动按钮执行全量比对/复核；内部身份与来源 ID、编码、名称分离"],
        ["D-09", "月内异动分段算薪", "原则已确认", "OA 调动/岗位/生效时间与考勤月结驱动有效期分段；部门工资、岗位工资、考勤工资项和其他项目分别计算并展示"],
        ["D-10", "工资与成本分摊", "原则已确认", "员工工资项目与责任组织/成本中心分摊分模型保存，分别勾稽；成本归属不授予个人工资查看权"],
        ["D-11", "PAYROLL 高敏权限", "已确认", "独立数据域、预设角色授权、字段白名单、短会话/可信网络/最近登录重验、制表与复核发布分离、加密/no-store、三次导出鉴权和运维隔离；不建设 MFA"],
        ["D-12", "工资数据来源", "已确认", "当前来自致远 OA 字段/关联文件和 HR 受控 Excel；暂存校验、冲突处理并发布为本系统有效期版本，原文件引用、哈希与批次留存"],
        ["D-13", "折算与调动日", "原则已确认", "P0 支持实际计薪工作日与实际计薪工时的项目级策略；调动生效日 00:00 起归新部门/岗位"],
        ["D-14", "法定计算与银行边界", "已确认", "P0 计算社保、公积金、个税并区分个人/单位承担；不采集银行账号、不生成银行付款文件，银企直连和直接申报列 P2"],
        ["D-15", "工资流程", "已确认", "不设置业务审批与审批状态；保留制表—独立复核发布控制，复核人只可退回/发布且不能修数"],
        ["D-16", "致远/得力接入校准", "已确认", "自建表只读接入保留；标准对象/附件使用可替换适配器；得力 ext_id 通过 EmployeeSourceBinding 确认，不再等同于工号"],
    ]
    add_table(doc, ["编号", "决策主题", "状态", "架构落点"], decision_rows,
              [0.65, 1.55, 1.0, 3.3], font_size=7.9, first_col_bold=True)
    add_callout(doc, "决策结果", "组织历史、考勤、薪资与权限边界继续成立，可以进入总体架构。自建表物理名称和字段本次不作为架构输入；它们通过版本化 OAQueryMappingProfile 收口，在接入详细设计/联调前冻结。得力人员绑定、OA 附件内容通道、工资项目/来源优先级、法定政策及年中个税期初累计仍分别阻断对应接入或公式上线，不阻断总体模块边界。", fill=LIGHT_BLUE, label_color=BLUE)

    add_heading(doc, "14.2 剩余数据契约与可带假设事项", 2)
    contract_open_rows = [
        ["DC-01", "黄色", "致远 V8.0 实际补丁、数据库类型、标准组织接入通道及 REST 服务/资源授权是否启用；自建表物理名称和字段本次忽略", "OA 接入详细设计前"],
        ["DC-02", "黄色", "得力真实 gps/out_work 样本中的 check_data 类型、lat/lgt 精度、坐标系及 name/location 语义", "地图联调与 UAT 前"],
        ["DC-03", "黄色", "当前在职人数、日均/峰值打卡、月末并发、平均载荷和数据库现存量", "生产资源规格冻结前"],
        ["DC-04", "黄色", "总部午休、12-31、打卡窗口/宽限、跨日截止、补签边界、餐扣及各假别正式参数", "对应模块金标准冻结前"],
        ["DC-05", "黄色", "地图供应商、应用 Key、坐标转换实现和个人信息评估", "地图生产启用前"],
        ["DC-06", "红色（详细设计）", "按对象冻结 OAQueryMappingProfile 与脱敏样本：逻辑业务键、人员/组织、调动生效、状态、时间、增量/撤销策略和能力标志；不要求本 PRD 写物理表名字段", "OA 接入详细设计与联调前"],
        ["DC-07", "红色", "OA 与 HR Excel 的实际字段/模板、员工与项目键、期间/有效期、金额/基数/费率、来源优先级与切换日；其他项目名称、方向、HR/部门录入模式、范围、计税/分段/成本属性", "工资配置与金标准冻结前"],
        ["DC-08", "红色（部分关闭）", "已确认实际工作日/实际工时与调动日归新部门；仍需逐项目确认策略、工作量事实来源、分子分母、请假缺勤防重复扣减、金额精度与取整阶段", "工资公式发布前"],
        ["DC-09", "黄色", "按默认矩阵配置实际薪酬制表、复核发布、薪资权限、财务成本和审计人员及范围；确认可信网络、会话时长和导出开关。不支持 MFA、不设置业务审批", "PAYROLL UAT 与上线前"],
        ["DC-10", "红色", "参保/缴存地、险种、基数与上下限、个人/单位比例、政策生效期；个税累计预扣、专项扣除/附加扣除、年度期初累计收入/扣除/已缴税、精度取整和追溯更正规则", "法定计算物理模型与金标准前"],
        ["DC-11", "黄色", "部门成本是否首批按实名员工分段核对，或仅提供满足小样本保护的汇总；成本中心/项目主数据来源", "成本分摊原型与 UAT 前"],
        ["DC-12", "红色（得力接入）", "现网员工 ext_id、user_id、employee_num 与 OA 工号的真实关系；App-Key/Secret 环境隔离、额度/限流及时间戳容差", "得力生产初始化与人员绑定前"],
        ["DC-13", "红色（薪资附件）", "OA 薪资附件内容采用 OA_REST、OA_PLUGIN_SERVICE 还是 MANUAL_IMPORT；验证文件内容可取、哈希可计算及授权范围", "自动薪资附件接入前"],
    ]
    add_table(doc, ["编号", "等级", "待补数据/配置", "最晚关闭"], contract_open_rows,
              [0.65, 0.7, 4.05, 1.1], font_size=7.8, first_col_bold=True)

    add_heading(doc, "15. 架构准入评审", 1)
    add_heading(doc, "15.1 准入结论", 2)
    add_callout(doc, "结论：有条件通过（可进入总体架构设计）", "建议现在进入总体架构设计。考勤、组织历史、OA 自建表适配边界、工资来源、分段与法定计算逻辑模型、工资/成本分离、PAYROLL 独立权限域、复核发布和双月结均已定义。自建表物理表名/字段不作为总体架构阻断项；映射样本、得力人员绑定、附件内容通道、工资项目/折算及法定政策分别在对应详细设计或上线门前关闭。", fill=LIGHT_BLUE, label_color=BLUE)
    admission_rows = [
        ["现在可以开展", "总体架构与模块边界；组织/岗位有效期模型；OA/Excel 暂存与冲突域；考勤月结到工资输入快照；工资分段/项目/税社保公积金/个税/成本逻辑模型；PAYROLL 独立权限与复核发布；双月结/追溯状态机；容量压测方案。"],
        ["暂不可冻结", "致远自建表查询/映射、标准组织接入通道、附件内容读取；得力正式人员绑定与限流；OA/Excel 冲突优先级；项目级工作日/工时分母、缺勤去重与取整；税社保政策和个税期初累计；地图坐标转换与资源规格。"],
        ["详细设计门槛", "逻辑领域与安全框架可立即开展；OA 映射须关闭 DC-01/DC-06，得力生产接入关闭 DC-12，自动附件接入关闭 DC-13，工资公式/金标准关闭 DC-07/DC-08/DC-10，权限 UAT 关闭 DC-09，地图关闭 DC-02/DC-05。"],
    ]
    add_table(doc, ["阶段判断", "边界"], admission_rows, [1.45, 5.05], font_size=8.5, first_col_bold=True)

    add_heading(doc, "15.2 红色：架构准入阻断项", 2)
    red_rows = [
        ["R-01", "致远自建对象映射未冻结", "总体架构已关闭", "V80 标准字典已校准；物理表名/字段本次忽略，后续 OAQueryMappingProfile 与样本阻断接入详细设计、生产 SQL 和状态映射"],
        ["R-02", "人员/组织/账号权威源与生命周期", "V1.4 已关闭", "致远为权威源；内部组织身份、不可变有效期版本与来源绑定分离，历史按发生时版本，访问按当前范围解析"],
        ["R-03", "假期发放与提前返岗结算", "V1.3 已关闭", "使用受控策略、四种返岗模式、不可变台账及版本状态机"],
        ["R-04", "容量与数据留存原则", "V1.3 已关闭", "按现网规模；全部核心数据在线保存于数据库，实测后冻结资源规格"],
        ["R-05", "首批夜班/轮班/弹性排班", "V1.3 已关闭", "P0 无此类班型且禁止启用，领域模型预留扩展"],
        ["R-06", "地图位置契约与权限边界", "V1.3 已关闭", "限定 gps/out_work 合法 lat/lgt、当前 OA 范围与独立权限、单点地图及审计"],
        ["R-07", "权威事实与写入边界", "V1.2 已关闭", "得力管打卡、OA 管主数据/审批、本系统管派生结果，且不上游写回"],
        ["R-08", "领域对象、状态机与并发不变量", "V1.2 已关闭", "稳定 ID、不可变版本、输入指纹、close_cutoff、期间锁和差异机制已定义"],
        ["R-09", "最小化来源载荷与灾备", "V1.3 已关闭", "坐标规范化到敏感列且不重复；RPO 使用持续增量日志/时间点恢复"],
        ["R-10", "薪资对象/字段越权与内部人员边界", "V1.6 原则关闭", "接受不建设 MFA 的剩余风险；以独立授权、职责互斥、可信网络、短会话、最近登录重验、加密/no-store、导出/运维/数据库安全补偿，上线前仍须专项验收"],
        ["R-11", "组织 ID/编码变化与复用", "V1.4 已关闭", "自动变化留痕、手动全量比对、疑似复用复核、默认不继承历史权限，组织绑定变化使旧缓存和令牌失效"],
        ["R-12", "工资来源、项目和有效期契约", "部分关闭", "来源类型已确认为 OA+HR Excel；实际字段/模板、附件读取、项目字典和冲突优先级仍影响物理模型、公式与追溯，须关闭 DC-06/DC-07"],
        ["R-13", "月中折算与调动生效日口径", "部分关闭", "实际工作日/工时及生效日归新部门已确认；项目映射、事实来源、分子分母、缺勤去重与取整仍须关闭 DC-08"],
        ["R-14", "税社保公积金政策与个税期初累计", "开放", "P0 法定计算依赖参保地、政策、基数/比例和年中累计输入；DC-10 未关闭前阻断物理模型、金标准和正式工资发布"],
        ["R-15", "得力 ext_id 人员绑定", "开放", "工号一致不能证明 ext_id 一致；DC-12 未关闭前阻断得力记录正式归人和考勤核算发布"],
        ["R-16", "OA 薪资附件内容读取", "开放", "数据库元数据可读不代表文件可用；DC-13 未关闭前阻断自动附件入薪，允许受控 HR Excel/人工导入兜底"],
        ["R-17", "OA 直读数据库边界", "V1.7 原则关闭", "保留用户确认的自建表只读方案，以白名单、最小 SELECT、分页限流、超时、字段指纹、升级回归和适配器隔离风险"],
    ]
    add_table(doc, ["编号", "问题", "状态", "架构影响/处理"], red_rows,
              [0.65, 1.85, 1.1, 2.9], font_size=7.6, first_col_bold=True)

    add_heading(doc, "15.3 黄色：带明确假设进入", 2)
    yellow_rows = [
        ["Y-01", "总部午休、12-31、打卡窗口、宽限与取整", "全部版本化配置；缺值即异常/阻断，不写死"],
        ["Y-02", "补签一周、跨日窗口、加班 48 小时与餐扣", "用策略参数表达；最终边界在金标准冻结前确认"],
        ["Y-03", "出差/外出与各假别详细参数", "保留通用对象、关系和规则版本；资料不完整不自动豁免"],
        ["Y-04", "得力坐标系、name/location 和 check_data 真实类型", "按官方 lat/lgt 建模，真实样本确认前不绘制生产点位"],
        ["Y-05", "地图供应商与转换实现", "经服务端适配器隔离；供应商、Key、坐标转换与版本可配置"],
        ["Y-06", "现网容量具体数字", "不推定增长；取得实测基线后完成 36 个月累积数据 POC 并冻结资源"],
        ["Y-07", "SSO/本地账号与部署拆分", "认证适配器隔离；不建设 MFA，默认模块化单体和独立异步任务边界"],
        ["Y-08", "会话时长、限流、小样本阈值与大屏视觉细节", "首版采用本文默认值，安全/交互详细设计时参数化并通过回归验证"],
        ["Y-09", "其他工资项目录入主体", "项目级配置 HR_CENTRAL、DEPARTMENT_SCOPED 或两者；具体首批项目归属可在项目字典阶段确认"],
        ["Y-10", "负责人/高管和财务薪资可见边界", "采用默认权限：负责人/高管仅看满足 5 人阈值与互补抑制的汇总；财务仅看已发布法定/会计成本必需字段，个人明细另行显式授权"],
        ["Y-11", "实际工作日/工时细节", "类型与调动日已确认；项目映射、事实来源、分母、防重复扣减和取整在 DC-08 关闭前均参数化并阻断缺值"],
        ["Y-12", "致远自建表物理表名与字段", "本 PRD 不固化；通过版本化 QueryMappingProfile 和脱敏样本在接入详细设计冻结，不改变领域模型"],
        ["Y-13", "致远标准对象接入通道", "使用可替换 SeeyonStandardAdapter；REST 可用时优先，否则使用受控只读视图/服务，不改变内部逻辑契约"],
    ]
    add_table(doc, ["编号", "事项", "进入架构设计时采用的假设"], yellow_rows,
              [0.65, 2.2, 3.65], font_size=8.0, first_col_bold=True)

    add_heading(doc, "15.4 绿色：不影响架构，可后续迭代", 2)
    green_rows = [
        ["G-01", "Excel 清单/汇总最终列名、顺序、单位和样式", "报表开发前确认，数据口径仍以日明细聚合为准"],
        ["G-02", "签到排行展示名、时间粒度、排序细节和留言字数/文案", "独立发布投影和治理接口已预留"],
        ["G-03", "看板图表类型、颜色、默认筛选和移动端微交互", "不改变指标、权限或数据版本"],
        ["G-04", "制度正文、附件样式、阅读确认与通知频率", "阅读确认和通知可进入 P1"],
        ["G-05", "地图弹层样式、默认缩放级别和图标", "不改变坐标契约、授权和审计"],
        ["G-06", "空状态、提示语、帮助说明和培训材料", "验收前按原型与试运行反馈完善"],
        ["G-07", "工资条颜色、折叠方式、温馨文案和打印样式", "不改变工资项目、金额、权限、版本或发布状态"],
    ]
    add_table(doc, ["编号", "可延后事项", "延后边界"], green_rows,
              [0.65, 2.65, 3.2], font_size=8.2, first_col_bold=True)

    add_heading(doc, "15.5 必须补充、可以延后与本次直接修订", 2)
    must_add = [
        ["必须补充", "进入对应详细设计/上线门前补充：各自建对象 OAQueryMappingProfile 与脱敏样本；得力 ext_id/user_id/employee_num 现网对账；OA 附件内容通道；HR Excel 模板与来源优先级；工资项目、折算/取整；税社保政策与个税期初累计。自建表物理字段不阻断总体架构。"],
        ["设计中补测", "现网容量数字；得力 gps/out_work 真实样本、坐标系、name/location 语义；地图供应商和转换方案。均可在总体架构阶段并行完成。"],
        ["可以延后", "银企直连、税务/社保直接申报、自动政策更新、多地区高级政策库、成本中心细节和黄色规则可延后；绿色的报表列、排行、看板/地图/工资条样式、制度增强和提示文案可后续迭代。"],
        ["V1.7 已直接修订", "根据 V80 字典及致远/得力官方文档校准：自建表采用版本化逻辑映射；得力 ext_id 独立绑定；OA 19 位 ID 字符串化；附件内容受控通道；API 鉴权、重试、只读边界与验收补齐。"],
    ]
    add_table(doc, ["类别", "内容"], must_add, [1.35, 5.15], font_size=8.3, first_col_bold=True)

    doc.add_page_break()
    add_heading(doc, "附录A. 字段字典与参考资料", 1)
    add_heading(doc, "A.1 得力标准打卡字段", 2)
    deli_fields = [
        ["company_id", "本系统上下文", "公司隔离键；P0 单公司仍必填"],
        ["source_instance_id", "数据源配置", "得力连接实例；与 source_record_id 组成唯一键"],
        ["source_record_id", "id", "来源记录标识，按字符串保存；不可脱离来源实例使用"],
        ["deli_user_id", "user_id", "得力内部人员标识；按字符串保存，不直接作为 OA 工号"],
        ["deli_ext_id", "ext_id", "得力官方外部员工关联标识；经 EmployeeSourceBinding 确认后归属内部员工"],
        ["deli_employee_num", "check_data.employee_num（如有）", "得力工号；可与 OA 工号对账并生成候选，不能替代 ext_id 绑定"],
        ["employee_source_binding_id", "本系统生成", "指向 CONFIRMED 人员来源绑定；未确认记录不得进入正式核算"],
        ["source_time_raw/occurred_at", "check_time", "同时保留源值与绝对发生时刻，再按 Asia/Shanghai 计算本地时间和考勤日"],
        ["received_at", "本系统生成", "实际接收时刻，用于识别设备迟报"],
        ["terminal_id", "terminal_id", "考勤机 SN、手机设备 ID 或 apply"],
        ["check_type", "check_type", "fa/fp/pass/card/app_scan/gps/wifi/out_work/reissue/flexible；未知值隔离"],
        ["check_data", "check_data", "仅保存白名单清洗后的补充字段；解析失败记录状态与摘要，不保存可形成坐标副本的原文"],
        ["latitude/longitude", "check_data.lat/lgt", "仅 gps/out_work 合法坐标规范化进入敏感列；lgt 为经度，解析后从 JSON 移除"],
        ["location_name/address", "check_data.name/location", "GPS/外勤地点名称和地址；与坐标同属 D 级敏感信息"],
        ["source_coordinate_system", "联调确认", "WGS84/GCJ-02/BD-09/UNKNOWN；UNKNOWN 不允许生产绘点"],
        ["coordinate_status", "本系统生成", "MISSING/INVALID/UNVERIFIED/VERIFIED/CONVERSION_FAILED"],
        ["map_provider/config_version", "本系统配置", "记录实际使用的地图供应商和转换配置版本，不记录坐标到审计事件"],
        ["synced_at", "本系统生成", "实际接收时间"],
        ["raw_payload", "最小化来源信封", "仅保留必要字段；不得包含未批准照片、手机/Wi-Fi 标识或重复坐标"],
    ]
    add_table(doc, ["标准字段", "来源字段", "说明"], deli_fields, [1.55, 1.8, 3.15], font_size=8.5, first_col_bold=True)

    add_heading(doc, "A.2 致远 OA 逻辑对象与标准化字段", 2)
    add_body(doc, "下表是 HR 系统核算所需的逻辑字段，不代表致远自建表中存在同名物理列。每个逻辑值须在版本化 OAQueryMappingProfile 中映射、组合或受控推导；实际表名、字段名和 SQL 在后续《致远接入映射说明》中冻结。V80 数据字典仅用于标准平台对象参考。")
    oa_fields = [
        ["company_id / source_instance_id", "系统必需", "由数据源配置生成；缺失则整批拒绝，不从名称推断"],
        ["mapping_profile_id/version", "系统必需", "标识查询、字段/状态/时间映射版本；未发布或未知版本不得生产运行"],
        ["logical_business_key", "对象必需", "由单字段或复合字段稳定生成；为空或重复时整批隔离"],
        ["source_employee_reference / employee_no", "人员对象必需", "至少一种可稳定定位人员的来源引用；经有效期映射为内部 employee_id，禁止按姓名自动匹配"],
        ["source_org_reference / code / name / parent", "组织对象必需", "定位、编码、名称、父级均按来源版本保存；任一物理字段缺失时按映射组合或将对象隔离"],
        ["organization_type / status / sort", "组织业务必要", "保留原始码并显式映射；未知状态不得默认启用"],
        ["organization_effective_range", "历史归属建议", "有来源业务有效时间则使用；否则以 first_seen/last_seen 记录 OBSERVED 版本并允许受控更正"],
        ["assignment_reference / assignment_kind", "任职业务必要", "区分主任职与兼职/次要任职；无稳定单键时可使用复合业务键"],
        ["source_position_reference / code / name", "岗位业务必要", "解析为内部岗位身份与版本；来源定位和编码均可变、可复用"],
        ["document_type / subtype", "单据对象必需", "标准化为 LEAVE、LEAVE_CANCEL、REISSUE、OVERTIME、BUSINESS_TRIP、OUTING_REGISTER、HR_TRANSFER 等"],
        ["business_start / business_end", "时段单据必需", "无法取得有效区间时不得自动覆盖考勤或工资"],
        ["source_status_raw / normalized_status", "单据对象必需", "映射为标准状态并保留原始值；未知状态隔离，只有最终有效版本参与核算"],
        ["submitted / approved / changed time", "按规则条件必需", "补签/加班时效需要首次提交；若无法可靠取得则进入人工复核。增量可改用全量指纹扫描"],
        ["linked_document_reference", "关联规则条件必需", "销假—请假、外出登记—出差等关系可由来源键或受控组合规则形成；无法形成时不自动豁免"],
        ["transfer_old_new / effective_at", "工资分段必需", "旧/新组织岗位职级与生效日无法可靠映射时，阻断对应工资分段"],
        ["payroll_period / component / value", "薪资来源条件必需", "按来源能力映射员工、期间、项目及金额/基数/费率；缺失项目进入冲突/异常"],
        ["attachment_reference / metadata", "附件条件必需", "附件定位、名称、类型、大小可先从元数据取得；不假设来源提供版本或内容哈希"],
        ["attachment_content_status / computed_hash", "自动附件入薪必需", "仅受控通道取得内容后计算哈希；只有元数据时标记 METADATA_ONLY 并阻断自动入薪"],
        ["authorization_nodes / penetration", "权限业务必需", "由致远标准组织/关系/授权数据或服务计算完整快照；无法完整解析时默认拒绝"],
        ["change / withdrawal / deletion semantics", "版本正确性必需", "可用显式标志、变更日志、完整快照差异或人工复核；单次未返回不能推断删除"],
        ["source_fingerprint / raw_payload", "系统生成", "保存最小化来源信封与内容指纹；内容变化由本系统创建 OADocumentVersion"],
    ]
    add_table(doc, ["逻辑字段", "业务必要性", "缺失处理/说明"], oa_fields, [2.05, 1.1, 3.35], font_size=8.3, first_col_bold=True)

    add_heading(doc, "A.2.1 HR 薪资 Excel 受控导入字段", 3)
    excel_fields = [
        ["template_version", "是", "模板与字段映射版本；未知版本整批进入校验失败，不猜测列含义"],
        ["import_batch_id/file_hash", "是", "导入批次与原文件内容哈希；同文件重放保持幂等，原文件进入受控存储"],
        ["sheet_name/row_number", "是", "定位每条来源行与错误；不得只保存规范化结果"],
        ["company_id/employee_no", "是", "公司与工号；工号按有效期映射内部员工，禁止按姓名自动匹配"],
        ["payroll_period/effective_from/to", "是", "工资期间或业务有效期；定义项目进入哪个输入快照和工资分段"],
        ["component_code", "是", "命中已发布工资项目定义；未知项目阻断该行并进入数据质量清单"],
        ["amount/quantity/base/rate", "条件必填", "按项目类型提供金额、数量、基数或费率；保留源精度和单位"],
        ["source_org/cost_center", "条件必填", "部门范围录入或成本分摊所需来源组织/成本中心；解析为内部身份"],
        ["business_key/source_record_id", "是", "员工+项目+期间/有效期+来源行唯一键，用于重复和跨来源冲突识别"],
        ["evidence_ref/remark", "条件建议", "项目要求的凭证引用和备注；附件沿用 PAYROLL 高敏存储与授权"],
        ["input_owner_mode", "系统配置", "HR_CENTRAL、DEPARTMENT_SCOPED 或 BOTH；服务端校验录入人组织范围"],
        ["validation/conflict_status", "系统生成", "VALID/REJECTED/CONFLICT/RESOLVED/PUBLISHED；未解决冲突不得锁定工资输入"],
    ]
    add_table(doc, ["标准字段", "必填", "说明"], excel_fields, [1.9, 0.8, 3.8], font_size=8.35, first_col_bold=True)
    add_callout(doc, "最小化要求", "P0 Excel 入库白名单不接收银行账号；文件包含未批准列时拒绝或在可验证删除后导入，并在数据质量报告记录。", fill=LIGHT_ORANGE, label_color=ORANGE)

    add_heading(doc, "A.3 核心内部记录字段", 2)
    internal_fields = [
        ["company_id", "所有核心记录", "公司隔离、唯一键、授权与分区的首要维度"],
        ["source_instance_id", "所有来源事实", "隔离不同外部连接、环境与凭据生命周期的来源键空间"],
        ["employee_id", "人员关联记录", "本系统稳定内部 ID；工号变化不改变该 ID"],
        ["employee_number_mapping_id", "来源事实/映射", "指向带有效期的工号映射，避免工号复用污染历史"],
        ["employee_source_binding_id", "得力打卡/人员对账", "指向 CONFIRMED EmployeeSourceBinding；保存得力 ext_id/user_id/employee_num 与 OA 人员引用的有效期关系"],
        ["oa_query_mapping_profile_id", "OA 来源批次/单据", "固定查询模板、逻辑业务键、字段/状态/时间、增量/撤销、附件与能力清单版本"],
        ["external_id_string", "致远标准来源对象", "所有可能为 64 位 Long 的人员、组织、岗位、流程、附件 ID 按字符或精确十进制保存并以字符串输出"],
        ["attendance_day", "日结果/异常/证据", "考勤业务日期，可与自然日不同"],
        ["organization_identity_id", "组织/任职/日结果/月结", "本系统生成的永久组织身份；名称、编码或 OA ID 变化不直接改变该 ID"],
        ["organization_version_id", "任职/日结果/月结", "固定发生日组织版本，包含当时来源 ID、编码、名称、父级、路径和状态"],
        ["source_organization_binding_id", "组织版本/授权", "指向来源实例+来源组织 ID 与内部身份的有效期绑定；同一时点不得重叠"],
        ["organization_hierarchy_snapshot_id", "组织版本/授权/月结", "同批发布的祖先后代闭包或路径快照；历史路径与当前穿透权限均引用明确版本"],
        ["organization_sync_batch_id", "组织/任职/授权/月结", "记录产生该源版本、绑定、组织图和授权快照的同步批次"],
        ["trigger_type/sync_mode/idempotency_key", "OrganizationSyncBatch", "区分 SCHEDULED/MANUAL 与 FULL/INCREMENTAL；同一请求重放不得重复发布"],
        ["node/edge/assignment/scope_count + source_hash", "OrganizationSyncBatch", "证明批次完整性并支持源端对账、恢复和审计"],
        ["observed_from/to + source_effective_from/to", "来源版本/绑定", "同时保存系统观测时间与来源业务有效时间；缺少来源有效时间时显式标记未验证"],
        ["change_type/decision/evidence/reviewer", "OrganizationChangeReview", "记录疑似复用、连续性、拆分合并的证据、结论、操作者、原因和时间"],
        ["organization_attribution_basis", "日结果/报表", "默认 OCCURRENCE_TIME；未来 CURRENT_RESTATEMENT 只能生成非正式独立口径"],
        ["oa_authorization_snapshot_id", "访问审计", "记录鉴权使用的最新致远范围、内部组织身份、绑定和层级快照版本"],
        ["rule_version_id", "日结果/假期/加班", "记录本次计算使用的有效规则版本"],
        ["input_fingerprint/evidence_hash", "日结果/月结", "覆盖来源版本、任职、规则和调整版本；提交前用于识别过期任务"],
        ["source_watermark/close_cutoff", "同步/月结", "标识已发布来源位置与月结截止点，截止后变化只进入差异"],
        ["calculation_version", "日/月结果", "输入或规则变化时递增；相同版本重算结果等价"],
        ["month_close_version", "快照/报表/导出", "正式结果版本；反月结后生成新版本，旧版本保留"],
        ["created_by/reason", "调整/反月结/敏感访问", "记录操作者、原因与时间，禁止为空"],
        ["position_identity_id/position_version_id", "任职/工资分段/月结", "发生时内部岗位身份和版本；来源岗位 ID/编码/名称变化不改写历史"],
        ["payroll_period_id/payroll_run_id", "工资输入/分段/项目/工资条", "工资期间与一次计算运行；正式、补发、追扣或纠正运行分别版本化"],
        ["payroll_input_snapshot_id", "工资运行", "冻结 OA 调动/任职/岗位/薪资字段与附件、HR Excel、工资配置、其他输入、法定政策、个税累计、考勤月结版本及总哈希"],
        ["payroll_source_batch/type/hash", "薪资来源", "标识 OA_FIELD/OA_ATTACHMENT/HR_EXCEL/SYSTEM、原文件/单据、导入人、时间、模板/映射版本与内容哈希"],
        ["source_sheet/row/business_key", "PayrollSourceRecord", "Excel Sheet/行号或 OA 记录定位、员工/项目/期间业务键；支持重复与冲突追溯"],
        ["source_authority_policy/conflict_id", "薪资来源冲突", "记录来源优先级/切换生效日及多候选冲突处置；未解决不得锁定"],
        ["payroll_segment_id/valid_from/to", "工资分段", "员工在期间内组织、岗位、定薪与工资规则稳定的半开区间；不得重叠或空档"],
        ["pay_scheme/rate/proration/rounding_version_id", "工资分段/项目", "引用已发布工资方案、标准、折算和取整版本；公式不可原地覆盖"],
        ["component_id/type/quantity/rate/base/result", "工资项目行", "收入/扣减项目的实际工作日/工时数量、分母、单价、基数和结果；保存来源版本与公式轨迹哈希"],
        ["work_measure_snapshot_id", "工资分段/项目", "冻结 ACTUAL_PAYABLE_WORKDAY/ACTUAL_PAYABLE_HOUR、分子分母、日历/排班/考勤版本和缺勤去重策略"],
        ["contribution_policy/profile/line_id", "社保公积金", "冻结参保/缴存地、基数上下限、个人/单位比例与金额、政策版本和取整轨迹"],
        ["tax_policy/profile/ytd_snapshot/line_id", "个人所得税", "冻结本期/累计收入、扣除、已预扣、本期税额、年度期初和政策版本"],
        ["earning_org/responsibility_org/cost_center", "工资项目/成本分摊", "分别记录计算归属与成本承担；可不同但必须有配置/来源凭证并分别勾稽"],
        ["payroll_authorization_grant/snapshot_id", "薪资访问/复核/导出", "冻结薪资范围、字段、用途、期间、有效期、配置人与策略版本；OA scope 只作为上限"],
        ["payslip_version/status/hash", "员工工资条", "DRAFT/PUBLISHED/SUPERSEDED；发布后不可删除，员工只见本人已发布版本"],
    ]
    add_table(doc, ["字段", "适用对象", "说明"], internal_fields,
              [1.9, 1.55, 3.05], font_size=8.2, first_col_bold=True)

    add_heading(doc, "A.4 名词解释", 2)
    glossary = [
        ["考勤日", "用于归属班次和核算结果的业务日期，不一定等同于自然日。"],
        ["开放期间", "尚未月结、允许自动重算和人工调整的考勤月份。"],
        ["金标准案例", "由考勤负责人确认输入与预期结果，用于开发联调和 QA 验收的脱敏样本。"],
        ["月结后差异", "已冻结月份收到迟报打卡、补单、撤单或规则变化后产生的差异记录。"],
        ["数据水位", "表示数据源已成功同步到的位置，如得力 next_id 或 OA 更新时间/批次。"],
        ["证据链", "支撑日考勤结果的班次、日历、打卡、OA 单据、规则和人工调整。"],
        ["特殊工时", "在基础班次上按员工和有效期叠加的免考勤时段，如哺乳、孕检。"],
        ["排行发布字段", "为签到榜单单独生成、允许员工相互查看的最小字段集，不等于原始考勤明细。"],
        ["数据域", "权限与数据生命周期相互隔离的业务边界，如 ATTENDANCE、LOCATION、PAYROLL；一个域的角色不自动继承另一域权限。"],
        ["Policy Enforcement Point", "所有 API、导出和实时消息访问数据前调用的统一服务端授权执行点，默认拒绝且逐请求判断。"],
        ["Opaque ID", "对外不可由工号或自增主键推导的随机对象标识；用于降低枚举，但不能替代对象级授权。"],
        ["组织身份（OrganizationIdentity）", "本系统生成、永久稳定的业务组织身份；不等于致远组织 ID、编码或名称。"],
        ["组织版本（OrganizationVersion）", "组织在一个有效期内的不可变状态，保存当时来源 ID、编码、名称、父级路径、状态和同步批次。"],
        ["来源组织绑定", "将某来源实例的组织 ID 在指定有效期内绑定到内部组织身份；允许来源 ID 分时复用，也允许同一身份更换来源 ID。"],
        ["当前组织投影", "由最新已发布组织版本重建、供页面与授权快速读取的可覆盖视图；不承担历史事实保存。"],
        ["组织变更复核", "当名称、编码、来源 ID 同时变化或疑似复用、拆分合并时，由授权人员确认身份是否连续的审计流程。"],
        ["工资分段（PayrollSegment）", "工资期间内任职、组织、岗位、定薪和工资方案均稳定的最小半开区间；月中调动会产生多个分段。"],
        ["工资项目（PayrollComponentLine）", "一个分段或月份下可解释的收入/扣减行，保存数量、基数、费率、公式/取整版本、来源和结果。"],
        ["成本分摊", "把工资相关金额归集到责任组织、成本中心或项目的独立结果；不改变员工工资，也不自动授予个人工资查看权。"],
        ["工资输入快照", "一次工资运行冻结的 OA 异动/任职/岗位、工资配置、人工输入、考勤月结版本和校验值。"],
        ["补发追扣", "工资关闭后发现历史差异时，在新工资运行中追加的正/负调整；必须引用原期间、原分段和原项目。"],
        ["敏感操作重新登录", "本系统不建设 MFA；发布、批量导出或薪资权限变更时，若最近登录超过 15 分钟，使用现有 OA/本地账号重新登录，并配合可信网络与短会话。"],
    ]
    add_table(doc, ["术语", "解释"], glossary, [1.4, 5.1], first_col_bold=True)

    add_heading(doc, "A.5 参考资料", 2)
    refs = [
        ("得力云开放文档中心：《考勤数据对接协议》", "https://doc.delicloud.com/v3/integration/oa.html"),
        ("致远开放平台开发文档", "https://open.seeyon.com/book/"),
        ("致远官方：关于数据字典与数据库直读提示", "https://open.seeyon.com/book/ctp/dd.html"),
        ("致远官方：组织模型管理 REST 接口", "https://open.seeyon.com/book/ctp/restjie-kou/zu-zhi-mo-xing-guan-li.html"),
        ("致远官方：表单流程集成", "https://open.seeyon.com/book/ctp/restjie-kou/biao-dan-liu-cheng-ji-cheng.html"),
        ("致远官方：文档与附件服务", "https://open.seeyon.com/book/ctp/restjie-kou/wen-dang-fu-wu-guan-li.html"),
        ("致远官方：数据库开发规范", "https://open.seeyon.com/book/dbSpec.html"),
        ("致远互联官方场景：考勤管理", "https://www.seeyon.com/tiyan/iframe/changjing/renli/renli2.html"),
        ("致远互联官方说明：OA 系统人员考勤管理", "https://www.seeyon.com/News/desc/id/6251.html"),
        ("北森官方：一体化 HR SaaS（多角色工作台与员工自助参考）", "https://www.beisen.com/product/hrsaas/"),
        ("北森官方：People Analytics（移动端、PC 与数据大屏参考）", "https://www.beisen.com/people-analytics/"),
        ("北森官方：薪酬管理系统（入转调离、假勤、薪酬体系与成本分摊参考）", "https://www.beisen.com/product/xcgl/"),
        ("北森官方案例：神州租车复杂算薪（按人员异动和考勤分段计薪、差异提示与多 Sheet 参考）", "https://www.beisen.com/customer/230.html"),
        ("北森官方：个税通与工资单（分段计薪、个税与 PC/移动工资条参考）", "https://www.beisen.com/solution/gst/"),
        ("北森官方：安全门户（数据隔离、字段加密、脱敏、水印与职责分离参考）", "https://trust.beisen.com/"),
        ("OWASP：Authorization Cheat Sheet（默认拒绝与每请求鉴权）", "https://cheatsheetseries.owasp.org/cheatsheets/Authorization_Cheat_Sheet.html"),
        ("OWASP API Security：Broken Object Level Authorization", "https://owasp.org/API-Security/editions/2023/en/0xa1-broken-object-level-authorization/"),
        ("OWASP WSTG：Testing for Excessive Data Exposure", "https://owasp.org/www-project-web-security-testing-guide/latest/4-Web_Application_Security_Testing/12-API_Testing/03-Testing_for_Excessive_Data_Exposure"),
    ]
    for label, url in refs:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        add_hyperlink(p, label, url)
    add_bullet(doc, "用户提供的原始系统流程图与补充制度问答截图（2026-07-17）。")
    add_bullet(doc, "用户提供：《V80数据字典.pdf》（433 页；用于致远标准平台对象校准，自建业务表物理字段不据此推定）。")

    if SOURCE_IMAGE is not None and SOURCE_IMAGE.exists():
        add_heading(doc, "A.6 用户原始流程图", 2)
        add_figure(doc, SOURCE_IMAGE, "图A-1  用户提供的原始流程图，仅作为需求来源留档", width=6.25,
                   alt_text="用户提供的流程草图：得力考勤机和OA数据库汇入考勤数据库，再提供管理后台、权限定位和响应式前端看板、年假、制度及反馈功能。")

    if (
        SOURCE_RULE_IMAGE_1 is not None and SOURCE_RULE_IMAGE_1.exists()
    ) or (
        SOURCE_RULE_IMAGE_2 is not None and SOURCE_RULE_IMAGE_2.exists()
    ):
        add_heading(doc, "A.7 用户补充制度问答", 2)
        if SOURCE_RULE_IMAGE_1 is not None and SOURCE_RULE_IMAGE_1.exists():
            add_figure(doc, SOURCE_RULE_IMAGE_1, "图A-2  补签、外出、跨夜加班、权限及季节班次补充", width=6.25,
                       alt_text="用户补充的考勤问题和规则：补签一周、外出双表、跨夜加班、48小时申报、权限及总部夏冬令。")
        if SOURCE_RULE_IMAGE_2 is not None and SOURCE_RULE_IMAGE_2.exists():
            add_figure(doc, SOURCE_RULE_IMAGE_2, "图A-3  地区班次、假别、特殊工时及销假补充", width=6.25,
                       alt_text="用户补充的地区和假期规则：大连成都上海班次、年假丧假护理假病假事假、哺乳孕检及提前返岗。")

    add_heading(doc, "A.8 PRD 自检结论", 2)
    review_rows = [
        ["目标与价值", "通过", "产品定位、目标指标和不做范围已明确"],
        ["角色与权限", "有条件通过", "PAYROLL 独立授权、制表/复核发布互斥、无 MFA 补偿控制及字段/对象/用途权限已定义；实际人员与范围上线前配置"],
        ["领域与数据模型", "通过", "组织/岗位有效期、OA/Excel 来源冲突、调动分段、税社保公积金、输入快照、工资条、成本分摊和追溯差异均已定义"],
        ["正常/异常流程", "通过", "覆盖考勤/工资来源、分段、法定计算、异常、复核发布、双月结、工资条和补发追扣状态迁移，不含业务审批或银行状态"],
        ["集成可开发性", "有条件通过", "得力官方契约、人员来源绑定、OA 自建表版本化映射与标准对象适配边界已定义；实际映射样本、附件内容通道和来源优先级在详细设计关闭"],
        ["容量可验证性", "通过", "按现网实测峰值与 36 个月时间累积完成数据库 POC，不依赖归档层"],
        ["验收可测试性", "通过", "考勤、组织变更、月中调动工资分段、成本勾稽、工资追溯、职责分离、权限和架构不变量均有 AC"],
        ["三端体验", "通过", "薪资核算工作台与员工 PC/手机工资条已定义；大屏明确禁止请求薪资数据"],
        ["安全与隐私", "有条件通过", "PAYROLL 独立域覆盖 URL/抓包、缓存、导出、日志、数据库和内部人员威胁；已明确无 MFA 并采用可信网络、短会话、最近登录、职责互斥和审计补偿"],
    ]
    add_table(doc, ["检查项", "结论", "说明"], review_rows, [1.45, 1.05, 4.0], font_size=8.8, first_col_bold=True)
    add_callout(doc, "评审结论", "PRD V1.7 建议“有条件通过，可以进入总体架构设计”。本轮根据 V80 数据字典及致远/得力官方文档完成集成校准：自建表不固化物理字段，得力 ext_id 使用显式人员来源绑定，OA 19 位 ID 字符串化，附件内容采用受控通道，并补齐只读、鉴权、重试与验收边界。实际映射样本、附件通道、工资来源优先级、项目折算和法定政策只阻断对应详细设计或上线门。", fill=LIGHT_BLUE, label_color=BLUE)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    QA_DIR.mkdir(parents=True, exist_ok=True)
    draw_development_flow()
    draw_system_data_flow()
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_contents(doc)

    # Content is appended below in the second half of this builder.
    build_prd_body(doc)
    doc.core_properties.title = "神州HR考勤与薪资核算系统 PRD V1.7"
    doc.core_properties.subject = "得力云与致远OA双源考勤、人事异动分段算薪、工资条、报表、员工自助和高敏权限"
    doc.core_properties.author = "Codex"
    doc.core_properties.keywords = "PRD, 考勤, 薪资, 工资条, 得力云, 致远OA, HR"
    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    main()
